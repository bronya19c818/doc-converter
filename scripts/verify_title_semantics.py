"""Independent title/subtitle/H1 metrics and actual DOCX, with whole-source baseline."""

from __future__ import annotations

import argparse
import io
import json
import os
import statistics
import subprocess
import sys
import tarfile
import tempfile
import time
from collections import Counter
from pathlib import Path

from docx import Document

from doc_converter.docx_direct.analyzer import analyze_docx
from doc_converter.docx_direct.classifier import classify_document
from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.docx_direct.ooxml_patch import parse_xml_payload, read_zip_member
from doc_converter.docx_styles import ROLE_TO_STYLE_ID
from doc_converter.normalizer import normalize_markdown
from doc_converter.renderer import render_docx
from doc_converter.style_library import create_style_pack

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures/title_semantics"
TARGETS = ("document-title", "document-subtitle", "heading-1")


def canonical(text):
    return "".join(text.split()).removeprefix("作者：").removeprefix("单位：")


def measure(rows):
    metrics = {}
    for role in TARGETS:
        tp = sum(r["expected"] == role and r["actual"] == role for r in rows)
        fp = sum(r["expected"] != role and r["actual"] == role for r in rows)
        fn = sum(r["expected"] == role and r["actual"] != role for r in rows)
        metrics[role] = {
            "tp": tp,
            "fp": fp,
            "fn": fn,
            "precision": tp / (tp + fp) if tp + fp else 0,
            "recall": tp / (tp + fn) if tp + fn else 0,
            "f1": 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else 0,
        }
    return {
        "roles": metrics,
        "items": len(rows),
        "errors": sum(r["expected"] != r["actual"] for r in rows),
        "output_checks": sum(r["output_correct"] is not None for r in rows),
        "output_errors": sum(r["output_correct"] is False for r in rows),
    }


def evaluate(output, render):
    pack = create_style_pack(name="标题语义验收", format_id="docx")
    records, timings = [], []
    for split in ("cases", "holdout"):
        cases = json.loads((FIXTURES / (split + ".json")).read_text(encoding="utf-8"))
        for case in cases:
            for route in ("markdown", "docx"):
                destination = output / split / route / case["name"]
                destination.mkdir(parents=True, exist_ok=True)
                expected = case["paragraphs"]
                source_docx = Document()
                for text, style, _ in expected:
                    source_docx.add_paragraph(text, style=style.replace("Heading1", "Heading 1"))
                source = destination / "source.docx"
                source_docx.save(source)
                before = time.perf_counter()
                roles = {}
                if route == "markdown":
                    text = "\n\n".join(
                        ("# " if style in {"Heading1", "Title"} else "") + t
                        for t, style, _ in expected
                    )
                    (destination / "source.md").write_text(text, encoding="utf-8")
                    normalized = normalize_markdown(text, filename=case["name"] + ".md")
                    for d in normalized.report["decisions"]:
                        role = (
                            "organization-department"
                            if d["role"] in {"organization", "department"}
                            else d["role"]
                        )
                        for part in str(d["before"]).removeprefix("# ").split("\n\n"):
                            roles[canonical(part)] = role
                    (destination / "normalized.md").write_text(
                        normalized.markdown, encoding="utf-8"
                    )
                    (destination / "audit.json").write_text(
                        json.dumps(normalized.report, ensure_ascii=False, indent=2),
                        encoding="utf-8",
                    )
                else:
                    analysis = classify_document(analyze_docx(source))
                    roles = {canonical(e.text): e.selected_decision.role for e in analysis.elements}
                timings.append(
                    {"case": case["name"], "route": route, "seconds": time.perf_counter() - before}
                )
                written = None
                if render:
                    path = destination / "output.docx"
                    if route == "markdown":
                        render_docx(
                            markdown=normalized.markdown,
                            style_pack=pack,
                            output_path=path,
                            work_dir=destination / "render",
                            role_decisions=normalized.report["decisions"],
                        )
                    else:
                        normalize_docx_direct(source, pack, path, audit_dir=destination / "audit")
                    written = Document(path)
                    resolver = EffectiveFormatResolver(
                        parse_xml_payload(read_zip_member(path, "word/styles.xml"))
                    )
                    assert [canonical(p.text) for p in written.paragraphs] == [
                        canonical(t) for t, _, _ in expected
                    ], (case["name"], route, "content/order")
                for i, (text, _, role) in enumerate(expected):
                    check, style, outline = None, None, None
                    if written is not None:
                        p = written.paragraphs[i]
                        style = p.style.style_id
                        outline = resolver.resolve_paragraph(p._p, paragraph_style_id=None).get(
                            "outline_level"
                        )
                        check = style == ROLE_TO_STYLE_ID[role] and (
                            outline == 0 if role == "heading-1" else outline in {None, 9}
                        )
                    records.append(
                        {
                            "split": split,
                            "case": case["name"],
                            "route": route,
                            "text": text,
                            "expected": role,
                            "actual": roles.get(canonical(text), "missing"),
                            "output_correct": check,
                            "actual_style": style,
                            "actual_outline": outline,
                        }
                    )
    return {
        "schema": "title-semantics-quality/v1",
        "records": records,
        "timings": timings,
        "groups": {
            split: {
                route: measure([r for r in records if r["split"] == split and r["route"] == route])
                for route in ("markdown", "docx")
            }
            for split in ("cases", "holdout")
        },
        "scope": "20 pre-implementation paired development cases, 8 separate validation cases; not a random business corpus",
    }


def benchmark(output):
    """Same 602 paragraphs, warmed engines, five runs; exclude rendering/audit IO."""
    paragraphs = [
        ("面向应急响应的区域资源共享", "Heading 1"),
        ("以城乡服务网络为研究对象", "Heading 1"),
    ]
    for i in range(1, 101):
        paragraphs.append((f"{i} 工作要求", "Heading 1"))
        paragraphs.extend(
            (f"这是第{i}部分第{j}段的实际业务说明，保留原文内容并记录执行结果。", "Normal")
            for j in range(5)
        )
    text = "\n\n".join(("# " if style == "Heading 1" else "") + t for t, style in paragraphs)
    doc = Document()
    for t, style in paragraphs:
        doc.add_paragraph(t, style=style)
    path = output / "performance-source.docx"
    doc.save(path)
    result = {}
    for route in ("markdown", "docx"):
        durations = []
        for i in range(6):
            start = time.perf_counter()
            if route == "markdown":
                normalized = normalize_markdown(text, filename="long-document.md")
            else:
                analysis = classify_document(analyze_docx(path))
            if i:
                durations.append(time.perf_counter() - start)
        result[route] = {
            "paragraphs": len(paragraphs),
            "characters": len(text),
            "runs": durations,
            "median_seconds": statistics.median(durations),
            "final_role_counts": dict(Counter(d["role"] for d in normalized.report["decisions"]))
            if route == "markdown"
            else dict(Counter(e.selected_decision.role for e in analysis.elements)),
        }
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--baseline-ref")
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--benchmark", action="store_true")
    parser.add_argument(
        "--check",
        action="store_true",
        help="Require development contracts and bounded validation errors",
    )
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.baseline_ref:
        with tempfile.TemporaryDirectory() as tmp:
            payload = subprocess.check_output(
                ["git", "archive", args.baseline_ref, "src"], cwd=ROOT
            )
            with tarfile.open(fileobj=io.BytesIO(payload)) as archive:
                archive.extractall(tmp, filter="data")
            env = {**os.environ, "PYTHONPATH": str(Path(tmp) / "src")}
            subprocess.run(
                [
                    sys.executable,
                    __file__,
                    "--output",
                    str(args.output / "baseline"),
                    *(["--render"] if args.render else []),
                    *(["--benchmark"] if args.benchmark else []),
                ],
                env=env,
                check=True,
            )
    report = evaluate(args.output / "current", args.render)
    if args.benchmark:
        report["performance"] = benchmark(args.output)
    (args.output / "quality.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report["groups"], ensure_ascii=False, indent=2))
    if args.check:
        for route in ("markdown", "docx"):
            assert report["groups"]["cases"][route]["errors"] == 0
            assert report["groups"]["cases"][route]["output_errors"] == 0
            validation = report["groups"]["holdout"][route]
            assert validation["errors"] <= 1
            for role in ("document-title", "document-subtitle"):
                assert validation["roles"][role]["fp"] == 0
            if args.benchmark:
                assert report["performance"][route]["final_role_counts"] == {
                    "document-title": 1,
                    "document-subtitle": 1,
                    "heading-1": 100,
                    "body-paragraph": 500,
                }


if __name__ == "__main__":
    main()
