"""Repeatable DOCX -> API -> DOCX/audit acceptance with a pre-change baseline.

Capture before implementation, then verify with the same fixtures/style pack:
uv run python scripts/verify_docx_shadow.py --capture-baseline artifacts/shadow-baseline
uv run python scripts/verify_docx_shadow.py --baseline artifacts/shadow-baseline --output artifacts/shadow-verified
"""

from __future__ import annotations

import argparse
import json
import os
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi.testclient import TestClient
from lxml import etree

from doc_converter.api import create_app
from doc_converter.style_library import StyleLibrary, StylePack, create_style_pack

ROOT = Path(__file__).resolve().parents[1]


def write_json(path: Path, payload: object) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def make_fixtures(folder: Path) -> dict[str, dict[str, str]]:
    expected: dict[str, dict[str, str]] = {}
    for name, headings in (
        (
            "chinese",
            ["一、实施背景", "（一）现状分析", "1. 业务情况", "（1）具体问题", "①详细措施"],
        ),
        (
            "decimal",
            [
                "1 范围",
                "1.1 适用范围",
                "1.1.1 具体要求",
                "1.1.1.1 方法",
                "1.1.1.1.1 细节",
                "1.1.1.1.1.1 参数",
            ],
        ),
        ("native", []),
    ):
        doc = Document()
        doc.styles["Normal"].font.size = Pt(12)
        roles: dict[str, str] = {}
        for text, size, role in (
            ("仓储业务实施方案", 24, "document-title"),
            ("年度重点任务", 16, "document-subtitle"),
            ("单位：示例公司", 12, "organization-department"),
            ("作者：张三", 12, "author"),
            ("日期：2026年10月3日", 12, "date"),
        ):
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(size)
            roles[text] = role
        for level, text in enumerate(headings, 1):
            doc.add_paragraph(text)
            roles[text] = f"heading-{level}"
            doc.add_paragraph("本段说明业务背景、实施措施和运行结果，应作为普通正文保留。")
        if name == "native":
            for level in range(1, 7):
                text = f"原生层级{level}"
                doc.add_heading(text, level=level)
                roles[text] = f"heading-{level}"
                doc.add_paragraph("这一段为完整正文，用于确认正文格式与结构识别的竞争。")
        for text in ("1. 这是一条完整的编号正文，不应升格为标题。", "二、后续工作"):
            doc.add_paragraph(text)
        if name == "chinese":
            roles["二、后续工作"] = "heading-1"
        roles["1. 这是一条完整的编号正文，不应升格为标题。"] = "body-paragraph"
        p = doc.add_paragraph("正文包含局部")
        r = p.add_run("大字号粗体")
        r.font.size = Pt(28)
        r.bold = True
        p.add_run("，其余文字介绍业务情况及相关措施，应当保持正文身份。")
        roles[p.text] = "body-paragraph"
        for text in ("表1 统计结果", "图1 工作流程"):
            doc.add_paragraph(text)
            roles[text] = "table-caption" if text.startswith("表") else "figure-caption"
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "一、表内内容"
        table.cell(0, 1).text = "参考文献"
        table.cell(1, 0).merge(table.cell(1, 1)).text = "合并单元格正文"
        for text in ("一、表内内容", "合并单元格正文"):
            roles[text] = "body-table"
        doc.add_paragraph("自动列表内容", style="List Number")
        roles["自动列表内容"] = "unknown"
        custom = doc.styles.add_style("UnexplainedCustom", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("孤立短语", style=custom)
        roles["孤立短语"] = "unknown"
        doc.add_paragraph("")
        doc.add_paragraph("参考文献")
        doc.add_paragraph("[1] 示例参考资料。")
        roles["[1] 示例参考资料。"] = "reference-entry"
        doc.save(folder / f"{name}.docx")
        expected[name] = roles
    # Heading-less body, and a deliberate hierarchy conflict, exercise context safety.
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_paragraph("没有显式标题的文章仍然包含连续完整正文，这一段介绍实施背景。")
    doc.add_paragraph("第二段介绍业务措施和预期结果，应归入正文区域。")
    doc.save(folder / "plain.docx")
    expected["plain"] = {p.text: "body-paragraph" for p in doc.paragraphs}
    doc = Document()
    doc.add_heading("父标题", level=1)
    doc.add_heading("异常跳级", level=4)
    doc.save(folder / "jump.docx")
    expected["jump"] = {}
    pressure = ROOT / "tests/fixtures/DOCX转Markdown压力测试_V1_结构兼容性.docx"
    (folder / "pressure.docx").write_bytes(pressure.read_bytes())
    expected["pressure"] = {}
    return expected


def part_hashes(path: Path) -> dict[str, str]:
    with ZipFile(path) as package:
        assert package.testzip() is None
        return {name: sha256(package.read(name)).hexdigest() for name in package.namelist()}


def convert(source: Path, output: Path, pack: StylePack) -> dict[str, object]:
    output.mkdir(parents=True, exist_ok=True)
    StyleLibrary(output / "app-data").save(pack.model_copy(update={"version": 0}))
    with TestClient(create_app(data_root=output / "app-data")) as client:
        assert client.get("/health").status_code == 200
        response = client.post(
            "/api/jobs",
            data={"style_pack_id": pack.id, "style_pack_version": "1"},
            files={"source": (source.name, source.read_bytes(), "application/octet-stream")},
        )
        assert response.status_code == 201, response.text
        job = response.json()
        assert job["conversion_mode"] == "docx-direct"
        response = client.post(f"/api/jobs/{job['id']}/normalize-docx-direct", json={})
        assert response.status_code == 200, response.text
        completed = response.json()
        for artifact in completed["artifacts"]:
            downloaded = client.get(f"/api/jobs/{job['id']}/artifacts/{artifact}")
            assert downloaded.status_code == 200
            (output / artifact).write_bytes(downloaded.content)
        job_path = output / "app-data/jobs" / job["id"]
        assert (job_path / "source/original.docx").read_bytes() == source.read_bytes()
        assert not (job_path / "source/imported.md").exists()
    analysis = json.loads((output / "docx-analysis.json").read_text(encoding="utf-8"))
    return {
        "source_sha256": sha256(source.read_bytes()).hexdigest(),
        "output_parts": part_hashes(output / completed["output_name"]),
        "legacy": [
            {k: e[k] for k in ("element_id", "role", "confidence", "action", "reasons")}
            for e in analysis["elements"]
        ],
    }


def check_shadow(output: Path, expected: dict[str, str]) -> int:
    payload = json.loads((output / "docx-analysis.json").read_text(encoding="utf-8"))
    roles = {
        "document-title",
        "document-subtitle",
        "organization-department",
        "author",
        "date",
        "body-paragraph",
        "body-table",
        "table-caption",
        "figure-caption",
        "reference-heading",
        "reference-entry",
        "unknown",
    } | {f"heading-{i}" for i in range(1, 7)}
    for e in payload["elements"]:
        assert e.get("scored_decision"), f"缺少影子识别：{e['text']}"
        assert {c["role"] for c in e["candidates"]} == roles
        assert e["features"]
        for c in e["candidates"]:
            assert c["evidence"]
            raw = sum(v["contribution"] for v in c["evidence"] if v["stage"] == "raw")
            assert raw == c["raw_score"]
            assert sum(v["contribution"] for v in c["evidence"]) == c["resolved_score"]
        ranked = sorted(e["candidates"], key=lambda c: (-c["resolved_score"], c["role"]))
        decision = e["scored_decision"]
        assert decision["score"] == ranked[0]["resolved_score"]
        assert decision["margin"] == ranked[0]["resolved_score"] - ranked[1]["resolved_score"]
        assert decision["runner_up_role"] == ranked[1]["role"]
        assert decision["reasons"]
        assert e["legacy_decision"]["role"] == e["role"]
        if e["text"] in expected:
            assert decision["role"] == expected[e["text"]], (e["text"], decision, ranked[:2])
        if e["text"] == "异常跳级":
            assert decision["action"] != "mapped"
        if e["text"] == "参考文献" and not e["table_context"]:
            assert decision["role"] == "reference-heading"
    report = json.loads((output / "docx-normalization-report.json").read_text(encoding="utf-8"))
    assert report["classification_mode"] == "legacy-with-shadow"
    assert all(r["legacy_decision"] and r["scored_decision"] for r in report["elements"])
    return len(payload["elements"])


def structure_signature(node):
    """Compare content/structure while ignoring only permitted visual formatting."""
    local = etree.QName(node).localname
    if local == "rPr" or local in {"pStyle", "jc", "ind", "spacing", "keepNext", "outlineLvl"}:
        return None
    children = tuple(value for child in node if (value := structure_signature(child)) is not None)
    # A pPr consisting solely of permitted formatting can appear/disappear.
    if local == "pPr" and not children and not node.attrib:
        return None
    return node.tag, tuple(sorted(node.attrib.items())), node.text, children


def check_policy_baseline(actual: dict, baseline: dict, previous: Path, output: Path) -> dict:
    """Phase two changes visual output, but cannot change recognition or content."""
    assert actual["source_sha256"] == baseline["source_sha256"]
    assert actual["legacy"] == baseline["legacy"]
    before, after = baseline["output_parts"], actual["output_parts"]
    assert before.keys() == after.keys()
    changed = sorted(k for k in before if before[k] != after[k])
    assert set(changed) <= {"word/document.xml", "word/styles.xml"}, changed
    previous_analysis = json.loads((previous / "docx-analysis.json").read_text(encoding="utf-8"))
    analysis = json.loads((output / "docx-analysis.json").read_text(encoding="utf-8"))
    fields = ("features", "candidates", "scored_decision", "legacy_decision")
    assert [{k: e[k] for k in fields} for e in previous_analysis["elements"]] == [
        {k: e[k] for k in fields} for e in analysis["elements"]
    ], "策略改变了影子评分或识别证据"
    previous_docx = next(previous.glob("*.docx"))
    output_docx = next(output.glob("*.docx"))
    with ZipFile(previous_docx) as old, ZipFile(output_docx) as new:
        roots = [etree.fromstring(z.read("word/document.xml")) for z in (old, new)]
        assert structure_signature(roots[0]) == structure_signature(roots[1]), "内容结构改变"
        style_roots = [etree.fromstring(z.read("word/styles.xml")) for z in (old, new)]
        styles = [
            {
                s.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}styleId"): s
                for s in root
                if etree.QName(s).localname == "style"
            }
            for root in style_roots
        ]
        for style_id, definition in styles[0].items():
            assert etree.tostring(definition, method="c14n", exclusive=True) == etree.tostring(
                styles[1][style_id], method="c14n", exclusive=True
            ), style_id
    report = json.loads((output / "docx-normalization-report.json").read_text(encoding="utf-8"))
    assert all(
        e["policy_name"] and e["policy_family"] and e["transform_safety"]
        for e in report["elements"]
    )
    return {
        "legacy_unchanged": True,
        "scored_unchanged": True,
        "structure_unchanged": True,
        "changed_parts": changed,
    }


def check_edges(folder: Path, pack: StylePack) -> dict[str, int]:
    """Adversarial full conversions, with conservative-action assertions."""
    folder.mkdir(parents=True, exist_ok=False)
    cases = {
        "isolated-number": ["1. 日常业务已经完成", "这是后续普通正文，介绍处理结果。"],
        "manual-steps": [
            "1. 打开设备",
            "2. 检查电源",
            "3. 点击启动",
            "后续记录操作结果并完成审核。",
        ],
        "same-size-preface": ["备注", "这是普通正文，说明具体业务事项。"],
        "large-left-title": ["正式实施方案", "这是普通正文，说明具体业务事项。"],
        "plain-short-phrases": ["适用范围", "工作要求", "执行安排"],
        "metadata-before-subtitle": [
            "正式方案",
            "作者：张三",
            "年度任务",
            "一、背景",
            "这是完整正文，说明实施背景。",
        ],
    }
    results = {}
    for name, texts in cases.items():
        doc = Document()
        doc.styles["Normal"].font.size = Pt(12)
        for i, text in enumerate(texts):
            p = doc.add_paragraph(text)
            if name == "same-size-preface" and i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if name == "large-left-title" and i == 0:
                p.runs[0].font.size = Pt(24)
            if name == "metadata-before-subtitle" and i < 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].font.size = Pt(24 if i == 0 else 16 if i == 2 else 12)
        source = folder / f"{name}.docx"
        doc.save(source)
        output = folder / name
        convert(source, output, pack)
        check_shadow(output, {texts[0]: "document-title"} if name == "large-left-title" else {})
        data = json.loads((output / "docx-analysis.json").read_text(encoding="utf-8"))
        if name == "large-left-title":
            assert data["elements"][0]["scored_decision"]["action"] == "mapped"
            results[name] = len(data["elements"])
            continue
        if name == "plain-short-phrases":
            assert all(e["features"]["is_body"] for e in data["elements"])
            assert all(
                e["scored_decision"]["role"] not in {"document-title", "document-subtitle"}
                for e in data["elements"]
            )
        limit = 3 if name == "manual-steps" else 1
        targets = (
            data["elements"][:limit]
            if name != "metadata-before-subtitle"
            else data["elements"][2:3]
        )
        assert all(e["scored_decision"]["action"] != "mapped" for e in targets), (
            name,
            [e["scored_decision"] for e in targets],
        )
        results[name] = len(data["elements"])
    return results


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--capture-baseline", type=Path)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--with-preview", action="store_true")
    parser.add_argument("--edge-cases", action="store_true")
    parser.add_argument(
        "--shadow-reference",
        type=Path,
        help="第一阶段已通过影子验证的输出目录，用于核对 scored 证据",
    )
    parser.add_argument(
        "--policy-output",
        action="store_true",
        help="第二阶段：允许 document/styles 视觉变化，严格核对识别与结构",
    )
    args = parser.parse_args()
    if not args.with_preview:
        # Preview is outside this acceptance's scope; use the supported missing-engine path.
        os.environ["DOC_CONVERTER_LIBREOFFICE"] = str(ROOT / "artifacts/no-preview-engine")
    if args.capture_baseline:
        folder = args.capture_baseline.resolve()
        folder.mkdir(parents=True, exist_ok=False)
        expected = make_fixtures(folder)
        pack = create_style_pack(name="影子识别验收", format_id="docx")
        (folder / "pack.json").write_text(pack.model_dump_json(indent=2), encoding="utf-8")
        write_json(folder / "expected.json", expected)
        results = {name: convert(folder / f"{name}.docx", folder / name, pack) for name in expected}
        write_json(folder / "baseline.json", results)
        print(f"BASELINE captured: {folder.resolve()}")
        return
    assert args.baseline and args.output
    args.baseline = args.baseline.resolve()
    args.output = args.output.resolve()
    args.output.mkdir(parents=True, exist_ok=False)
    pack = StylePack.model_validate_json((args.baseline / "pack.json").read_text(encoding="utf-8"))
    expected = json.loads((args.baseline / "expected.json").read_text(encoding="utf-8"))
    baseline = json.loads((args.baseline / "baseline.json").read_text(encoding="utf-8"))
    results = {}
    for name, roles in expected.items():
        output = args.output / name
        actual = convert(args.baseline / f"{name}.docx", output, pack)
        if args.policy_output:
            results[name] = check_policy_baseline(
                actual, baseline[name], (args.shadow_reference or args.baseline) / name, output
            )
        else:
            assert actual == baseline[name], f"影子识别改变了源/legacy/输出部件：{name}"
            results[name] = {"output_unchanged": True}
        results[name]["paragraphs"] = check_shadow(output, roles)
    edges = check_edges(args.output / "edges", pack) if args.edge_cases else {}
    write_json(
        args.output / "verification.json", {"passed": True, "fixtures": results, "edges": edges}
    )
    print(json.dumps(results, ensure_ascii=False))


if __name__ == "__main__":
    main()
