"""Build a CRLF context and exercise a real, isolated Docker runtime (not mocked APIs)."""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import tempfile
import time
import uuid
from pathlib import Path

import httpx
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def docker(*args: str) -> str:
    return subprocess.check_output(["docker", *args], text=True).strip()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--image", default="doc-converter:issue3-test")
    parser.add_argument("--build-crlf", action="store_true")
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts/docker-runtime")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    name = "doc-converter-smoke-" + uuid.uuid4().hex[:10]
    volume = name + "-data"
    events: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="doc-converter-crlf-") as scratch:
        temp = Path(scratch)
        if args.build_crlf:
            for directory in ("frontend", "src"):
                shutil.copytree(
                    ROOT / directory,
                    temp / directory,
                    ignore=shutil.ignore_patterns("node_modules", "__pycache__", "dist"),
                )
            for filename in (
                "Dockerfile",
                "docker-entrypoint.sh",
                "pyproject.toml",
                "README.md",
                ".dockerignore",
            ):
                shutil.copyfile(ROOT / filename, temp / filename)
            entry = temp / "docker-entrypoint.sh"
            entry.write_bytes(entry.read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"))
            assert b"#!/bin/sh\r\n" in entry.read_bytes()
            subprocess.run(["docker", "build", "--tag", args.image, str(temp)], check=True)
            events.append({"check": "build-with-real-CRLF-entrypoint", "passed": True})
        docker("volume", "create", volume)
        try:
            docker(
                "run",
                "--detach",
                "--init",
                "--name",
                name,
                "--publish",
                "127.0.0.1::8765",
                "--volume",
                volume + ":/data",
                args.image,
            )
            details = json.loads(docker("inspect", name))[0]
            port = details["NetworkSettings"]["Ports"]["8765/tcp"][0]["HostPort"]
            base = f"http://127.0.0.1:{port}"

            def wait_healthy() -> None:
                deadline = time.monotonic() + 120
                while time.monotonic() < deadline:
                    state = json.loads(docker("inspect", name))[0]["State"]
                    if not state["Running"]:
                        raise RuntimeError(docker("logs", name))
                    if state.get("Health", {}).get("Status") == "healthy":
                        return
                    time.sleep(1)
                raise RuntimeError("Container did not reach Docker HEALTHCHECK healthy")

            wait_healthy()
            docker(
                "exec",
                name,
                "python",
                "-c",
                'from pathlib import Path; assert bytes([13]) not in Path("/usr/local/bin/docker-entrypoint.sh").read_bytes()',
            )
            with httpx.Client(base_url=base, timeout=120) as client:
                assert client.get("/health").status_code == 200
                home = client.get("/")
                assert home.status_code == 200
                assert home.headers["cache-control"] == "no-cache"
                assert "/assets/" in home.text
                assert client.get("/api/style-packs").status_code == 200
                pack_response = client.post(
                    "/api/style-packs", json={"name": "容器验收", "format_id": "docx"}
                )
                pack_response.raise_for_status()
                pack = pack_response.json()
                assert pack["roles"]["heading-1"]["font"]
                docx_path = temp / "smoke.docx"
                document = Document()
                document.add_heading("容器验收报告", 0)
                document.add_heading("一、范围", 1)
                document.add_paragraph("正文含金额100元。")
                document.save(docx_path)
                docker("cp", str(docx_path), name + ":/tmp/smoke.docx")
                docker(
                    "exec",
                    name,
                    "libreoffice",
                    "--headless",
                    "-env:UserInstallation=file:///tmp/smoke-profile",
                    "--convert-to",
                    "doc",
                    "--outdir",
                    "/tmp",
                    "/tmp/smoke.docx",
                )
                docker("cp", name + ":/tmp/smoke.doc", str(temp / "smoke.doc"))
                sources = [
                    (
                        "smoke.md",
                        "# 容器验收报告\n\n一、范围\n\n正文含金额100元。\n\n二、结果\n\n完成核查。".encode(),
                    ),
                    ("smoke.doc", (temp / "smoke.doc").read_bytes()),
                    ("smoke.docx", docx_path.read_bytes()),
                ]
                jobs = []
                for filename, payload in sources:
                    imported = client.post(
                        "/api/jobs",
                        data={
                            "style_pack_id": pack["id"],
                            "style_pack_version": str(pack["version"]),
                        },
                        files={"source": (filename, payload)},
                    )
                    imported.raise_for_status()
                    job = imported.json()
                    jobs.append(job["id"])
                    endpoint = (
                        "normalize-docx-direct"
                        if job["conversion_mode"] == "docx-direct"
                        else "normalize-and-render"
                    )
                    rendered = client.post(f"/api/jobs/{job['id']}/{endpoint}", json={})
                    rendered.raise_for_status()
                    completed = rendered.json()
                    assert completed["status"] == "completed", completed
                    assert completed["preview_ready"], completed
                    downloaded = client.get(
                        f"/api/jobs/{job['id']}/artifacts/{completed['output_name']}"
                    )
                    downloaded.raise_for_status()
                    assert downloaded.content.startswith(b"PK")
                    preview = client.get(f"/api/jobs/{job['id']}/preview")
                    preview.raise_for_status()
                    assert preview.content.startswith(b"%PDF")
                    if endpoint == "normalize-and-render":
                        verification = client.get(
                            f"/api/jobs/{job['id']}/artifacts/output-verification.json"
                        )
                        verification.raise_for_status()
                        assert verification.json()["passed"], verification.text
                        repeated = client.post(f"/api/jobs/{job['id']}/{endpoint}", json={})
                        repeated.raise_for_status()
                        assert repeated.json()["status"] == "completed"
                    (args.output / filename.replace(".doc", "-source.doc")).write_bytes(payload)
                    (args.output / (filename + ".output.docx")).write_bytes(downloaded.content)
                    events.append(
                        {"check": filename + "-import-render-download-preview", "passed": True}
                    )
                docker("restart", name)
                wait_healthy()
                restarted = json.loads(docker("inspect", name))[0]
                restarted_port = restarted["NetworkSettings"]["Ports"]["8765/tcp"][0]["HostPort"]
                # Docker may allocate a different random host port on restart.
                # Open a fresh client against the observed binding, not a stale pooled connection.
                with httpx.Client(
                    base_url=f"http://127.0.0.1:{restarted_port}", timeout=120
                ) as after_restart:
                    assert after_restart.get("/health").status_code == 200
                    for job_id in jobs:
                        assert (
                            after_restart.get("/api/jobs/" + job_id).json()["status"] == "completed"
                        )
                events.append(
                    {
                        "check": "restart-persistent-volume",
                        "passed": True,
                        "initial_port": port,
                        "restarted_port": restarted_port,
                    }
                )
        finally:
            subprocess.run(
                ["docker", "logs", name],
                stdout=(args.output / "container.log").open("w"),
                stderr=subprocess.STDOUT,
                check=False,
            )
            subprocess.run(["docker", "rm", "--force", name], check=False)
            docker("volume", "rm", volume)
            (args.output / "verification.json").write_text(json.dumps({"checks": events}, indent=2))
    print(json.dumps(events, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
