"""Evaluate pre-labelled inputs, actual repairs and independently expected DOCX styles."""

from __future__ import annotations

import argparse
import contextlib
import importlib.util
import json
import re
import subprocess
import sys
import tempfile
import time
import tracemalloc
from collections import Counter, defaultdict, deque
from pathlib import Path
from typing import Any

from docx import Document
from markdown_it import MarkdownIt

from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.ooxml_patch import parse_xml_payload, read_zip_member
from doc_converter.docx_styles import ROLE_TO_STYLE_ID
from doc_converter.normalizer import normalize_markdown
from doc_converter.renderer import render_docx
from doc_converter.style_library import create_style_pack

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures/markdown_scoring"
COMPLEX = ROOT / "tests/fixtures/markdown_complex"
METADATA = {
    "document-title",
    "document-subtitle",
    "author",
    "organization",
    "department",
    "date",
    "abstract-label",
    "keywords-label",
}


def ratio(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def summarize(records: list[dict[str, Any]]) -> dict[str, Any]:
    headings = [r for r in records if r["role"].startswith("heading-")]
    predicted = [r for r in records if r["actual_role"].startswith("heading-")]
    tp = sum(r["actual_role"].startswith("heading-") for r in headings)
    precision, recall = ratio(tp, len(predicted)), ratio(tp, len(headings))
    needed = [r for r in records if r["needs_repair"]]

    def correct(r: dict[str, Any]) -> bool:
        return bool(r["role"] == r["actual_role"])

    repair_coverage = ratio(sum(r["repair_success"] for r in needed), len(needed))
    return {
        "annotated_items": len(records),
        "correct_roles": sum(correct(r) for r in records),
        "role_accuracy": ratio(sum(correct(r) for r in records), len(records)),
        "heading_items": len(headings),
        "heading_precision": precision,
        "heading_recall": recall,
        "heading_f1": 2 * tp / (len(predicted) + len(headings)) if predicted or headings else None,
        "heading_layer_accuracy": ratio(sum(correct(r) for r in headings), len(headings)),
        "required_repairs": len(needed),
        "correct_auto_repair_coverage": repair_coverage,
        "correct_auto_coverage": repair_coverage,  # compatibility name; v2 uses the corrected denominator
        "false_repairs": sum(r["changed"] and not correct(r) for r in records),
        "missed_repairs": sum(not r["repair_success"] for r in needed),
        "manual_corrections": sum(
            not correct(r) or r.get("output_correct") is False for r in records
        ),
        "metadata_accuracy": ratio(
            sum(correct(r) for r in records if r["role"] in METADATA),
            sum(r["role"] in METADATA for r in records),
        ),
        "body_numbering_false_heading_rate": ratio(
            sum(
                r["actual_role"].startswith("heading-")
                for r in records
                if r["role"] == "body-paragraph"
                and re.match(r"^(?:\d|[（(]|[一二三四五六七八九十]+、)", r.get("text", ""))
            ),
            sum(
                r["role"] == "body-paragraph"
                and bool(re.match(r"^(?:\d|[（(]|[一二三四五六七八九十]+、)", r.get("text", "")))
                for r in records
            ),
        ),
        "true_list_errors": sum(not correct(r) for r in records if r["role"] == "preserved"),
        "output_checked_items": sum(r.get("output_correct") is not None for r in records),
        "output_errors": sum(r.get("output_correct") is False for r in records),
        "errors_by_role": dict(Counter(r["role"] for r in records if not correct(r))),
    }


def plain(text: str) -> str:
    text = re.sub(r"^#{1,6}\s+", "", text)
    children = MarkdownIt().parseInline(text)[0].children or []
    return re.sub(
        r"\s+",
        " ",
        "".join(
            c.content for c in children if c.type in {"text", "code_inline", "image", "softbreak"}
        ),
    ).strip()


def check_output(path: Path, expected: list[dict[str, Any]]) -> list[dict[str, Any]]:
    document = Document(path)
    resolver = EffectiveFormatResolver(parse_xml_payload(read_zip_member(path, "word/styles.xml")))
    paragraphs: dict[str, deque[Any]] = defaultdict(deque)
    for p in document.paragraphs:
        paragraphs[re.sub(r"\s+", " ", p.text).strip()].append(p)
    output = []
    for e in expected:
        role, text = e["role"], plain(e["text"])
        if role == "preserved":
            # Integrity is asserted by syntax-specific tests, not a fake DOCX style match.
            output.append({"output_correct": None, "output_check": "protected-not-restyled"})
            continue
        if role == "body-table":
            output.append(
                {"output_correct": bool(document.tables), "output_check": "table-present"}
            )
            continue
        if role in {"organization", "department", "author", "date"}:
            text = re.sub(r"^(?:单位|机构|部门|作者|日期|时间)\s*[:：]\s*", "", text)
        if ": " in text and e.get("front_matter"):
            text = text.split(": ", 1)[1]
        queue = paragraphs.get(text)
        p = queue.popleft() if queue else None
        actual = p.style.style_id if p else None
        outline = (
            resolver.resolve_paragraph(p._p, paragraph_style_id=None).get("outline_level")
            if p
            else None
        )
        expected_style = ROLE_TO_STYLE_ID.get(
            "organization-department" if role in {"organization", "department"} else role
        )
        expected_outline = int(role[-1]) - 1 if role.startswith("heading-") else None
        matches = actual == expected_style and (
            expected_outline is None or outline == expected_outline
        )
        # Source numeric prose may remain an actual list with numbering; inspect its body semantics.
        if role == "body-paragraph" and p is None and re.match(r"^\d+[.)]\s+", text):
            queue = paragraphs.get(re.sub(r"^\d+[.)]\s+", "", text))
            p = queue.popleft() if queue else None
            matches = bool(
                p
                and p._p.xpath("./w:pPr/w:numPr")
                and not p.style.style_id.startswith("DC_Heading")
            )
        output.append(
            {
                "output_correct": matches,
                "expected_style": expected_style,
                "actual_style": actual,
                "actual_outline": outline,
                "expected_outline": expected_outline,
            }
        )
    return output


def load_sample(
    sample: dict[str, Any], fixtures: Path, destination: Path
) -> tuple[str, list[dict[str, Any]]]:
    if sample.get("input_kind") != "doc":
        return (fixtures / sample["file"]).read_text(encoding="utf-8"), sample["expected"]
    from doc_converter.importers.doc import import_doc
    from doc_converter.importers.libreoffice import resolve_libreoffice
    from doc_converter.markdown_scoring import parse_blocks

    doc = Document()
    for e in sample["expected"]:
        doc.add_paragraph(e["text"], style="Title" if e["role"] == "document-title" else "Normal")
    source_docx = destination / (sample["file"] + ".docx")
    doc.save(str(source_docx))
    subprocess.run(
        [
            str(resolve_libreoffice()),
            "--headless",
            "-env:UserInstallation=" + (destination / "lo-profile").resolve().as_uri(),
            "--convert-to",
            "doc",
            "--outdir",
            str(destination.resolve()),
            str(source_docx.resolve()),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    source_doc = source_docx.with_suffix(".doc").resolve()
    assert source_doc.read_bytes().startswith(bytes.fromhex("d0cf11e0"))
    source = import_doc(source_doc, job_directory=destination.resolve() / "doc-import").markdown
    blocks = parse_blocks(source.splitlines())
    by_text = {b.text: b for b in blocks}
    expected = []
    for e in sample["expected"]:
        if e["role"] == "document-title":
            from doc_converter.normalizer import _extract_front_matter

            meta, _, _ = _extract_front_matter(source)
            if meta.get("title") == e["text"]:
                line = next(
                    i + 1 for i, t in enumerate(source.splitlines()) if t.startswith("title:")
                )
                expected.append(
                    {
                        **e,
                        "line": line,
                        "text": source.splitlines()[line - 1],
                        "front_matter": True,
                        "needs_repair": False,
                    }
                )
                continue
        block = by_text.get(e["text"])
        if block is None:
            raise ValueError("DOC import content changed or merged: " + e["text"])
        expected.append({**e, "line": block.start + 1, "text": block.inline})
    (destination / "doc-imported.md").write_text(source, encoding="utf-8")
    return source, expected


def evaluate(
    fixtures: Path = FIXTURES, *, render: bool = False, output_dir: Path | None = None
) -> dict[str, Any]:
    truth = json.loads((fixtures / "ground_truth.json").read_text(encoding="utf-8"))
    records: list[dict[str, Any]] = []
    performance = []
    with tempfile.TemporaryDirectory(prefix="markdown-quality-") as temp:
        destination = output_dir or Path(temp)
        destination.mkdir(parents=True, exist_ok=True)
        for sample in truth["samples"]:
            source, expected = load_sample(sample, fixtures, destination)
            tracemalloc.start()
            start = time.perf_counter()
            result = normalize_markdown(source, filename=sample["file"])
            elapsed = time.perf_counter() - start
            _, peak = tracemalloc.get_traced_memory()
            tracemalloc.stop()
            performance.append(
                {
                    "file": sample["file"],
                    "source_bytes": len(source.encode()),
                    "normalization_seconds": elapsed,
                    "peak_python_bytes": peak,
                }
            )
            by_line = {d["line"]: d for d in result.report["decisions"]}
            output_checks: list[dict[str, Any]] = []
            if render:
                path = destination / (sample["file"] + ".docx")
                start = time.perf_counter()
                rendered = render_docx(
                    markdown=result.markdown,
                    style_pack=create_style_pack(name="质量验收", format_id="docx"),
                    output_path=path,
                    work_dir=destination / (sample["file"] + ".work"),
                    role_decisions=result.report["decisions"],
                )
                output_checks = check_output(path, expected)
                performance[-1]["render_seconds"] = time.perf_counter() - start
                (destination / (sample["file"] + ".written.json")).write_text(
                    json.dumps(rendered.verification, ensure_ascii=False, indent=2)
                )
            for i, e in enumerate(expected):
                d = by_line.get(e["line"], {})
                role = d.get("role", "missing")
                original = re.match(r"^(#{1,6})\s+", e["text"])
                markup_repair = (
                    e["role"].startswith("heading-")
                    and (not original or len(original.group(1)) != int(e["role"][-1]))
                    or e["role"] in METADATA
                    and not d.get("features", {}).get("front_matter")
                )
                needed = bool(e.get("needs_repair", markup_repair))
                style_repair = e["role"] in {
                    "table-caption",
                    "figure-caption",
                    "reference-heading",
                    "reference-entry",
                }
                output_check = output_checks[i] if render else {}
                applied = d.get("applied_change", False)
                success = (
                    needed
                    and role == e["role"]
                    and (applied or style_repair and output_check.get("output_correct", False))
                    and output_check.get("output_correct", True)
                )
                records.append(
                    {
                        "file": sample["file"],
                        "document_type": sample.get("document_type", "synthetic"),
                        "input_kind": sample.get("input_kind", "md"),
                        **e,
                        "needs_repair": needed,
                        "actual_role": role,
                        "correct": role == e["role"],
                        "action": d.get("action", "missing"),
                        "changed": applied
                        or style_repair
                        and render
                        and output_check.get("output_correct", False),
                        "repair_success": success,
                        "score": d.get("score"),
                        "margin": d.get("margin"),
                        "reason": d.get("reason"),
                        **output_check,
                    }
                )
    metrics = summarize(records)
    return {
        "schema": "markdown-quality/v2",
        "scope": truth["scope"],
        "metrics": metrics,
        "groups": {
            kind: summarize([r for r in records if r["document_type"] == kind])
            for kind in sorted({r["document_type"] for r in records})
        },
        "groups_by_input": {
            kind: summarize([r for r in records if r["input_kind"] == kind])
            for kind in sorted({r["input_kind"] for r in records})
        },
        "records": records,
        "performance": performance,
        "limitations": [
            "独立人工标注受控样本，不代表真实业务文档总体正确率。",
            "修复覆盖率分母是应修复项；仅分类正确或action=mapped不算正确修复。",
            "未指定--render时未核验样式修复项；未标注项不扩大指标分母。",
            "tracemalloc统计Python内存；单次计时含采集开销，不是吞吐基准。",
        ],
    }


@contextlib.contextmanager
def baseline_module(path: Path):
    global normalize_markdown
    original_normalize = normalize_markdown
    originals = {}
    try:
        for suffix in ("markdown_preprocess", "normalizer", "markdown_scoring"):
            name = "doc_converter." + suffix
            source = path.with_name(path.name.replace("markdown_scoring", suffix))
            if not source.is_file():
                raise ValueError(f"Missing frozen baseline source: {source}")
            originals[name] = sys.modules.get(name)
            spec = importlib.util.spec_from_file_location(name, source)
            assert spec and spec.loader
            module = importlib.util.module_from_spec(spec)
            sys.modules[name] = module
            spec.loader.exec_module(module)
            if suffix == "normalizer":
                normalize_markdown = module.normalize_markdown
        yield
    finally:
        normalize_markdown = original_normalize
        for name, original in originals.items():
            if original is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = original


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--fixtures", type=Path, default=FIXTURES)
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--baseline-source", type=Path)
    parser.add_argument("--baseline-ref")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.baseline_source or args.baseline_ref:
        with tempfile.TemporaryDirectory() as tmp:
            path = args.baseline_source or Path(tmp) / "markdown_scoring.py"
            if args.baseline_ref:
                for module_name in ("markdown_scoring", "markdown_preprocess", "normalizer"):
                    frozen = path.with_name(module_name + ".py")
                    frozen.write_bytes(
                        subprocess.check_output(
                            [
                                "git",
                                "show",
                                args.baseline_ref + ":src/doc_converter/" + module_name + ".py",
                            ],
                            cwd=ROOT,
                        )
                    )
            with baseline_module(path):
                old = evaluate(
                    args.fixtures, render=args.render, output_dir=args.output / "baseline"
                )
        (args.output / "baseline.json").write_text(json.dumps(old, ensure_ascii=False, indent=2))
    report = evaluate(args.fixtures, render=args.render, output_dir=args.output / "current")
    (args.output / "quality.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report["metrics"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
