"""Reproduce and verify Direct picture paragraph styling without rebuilding media."""

from __future__ import annotations

import argparse
import base64
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

from doc_converter.docx_direct.analyzer import analyze_docx
from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.style_library import create_style_pack

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
)


def fixture(path: Path) -> tuple[list[int], list[int]]:
    document = Document()
    wanted, preserved = [], []

    def picture(container, *, accept: bool = True, text: str = "", style: str | None = None):
        p = container.add_paragraph(text, style)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(21)
        p.paragraph_format.space_after = Pt(22)
        p.paragraph_format.line_spacing = Pt(5)
        p.paragraph_format.first_line_indent = Pt(25)
        p.add_run().add_picture(BytesIO(PNG), width=Inches(1))
        (wanted if accept else preserved).append(p)
        return p

    picture(document)
    dirty = document.styles.add_style("DirtyImageSource", 1)
    dirty.font.size = Pt(27)
    dirty.font.bold = True
    picture(document, style=dirty.name)
    p = picture(document)
    p.add_run().add_picture(BytesIO(PNG), width=Inches(0.5))
    table = document.add_table(rows=1, cols=1)
    picture(table.cell(0, 0))
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    picture(nested.cell(0, 0))
    picture(document, accept=False, text="正文中的图片")
    picture(document, accept=False, style="List Bullet")
    p = picture(document, accept=False)
    inline = p._p.find(".//" + qn("wp:inline"))
    inline.tag = qn("wp:anchor")
    p = picture(document, accept=False)
    data = p._p.find(".//" + qn("a:graphicData"))
    data.set("uri", "http://schemas.openxmlformats.org/drawingml/2006/chart")
    p = picture(document, accept=False)
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    p.add_run()._r.append(field)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(end)
    picture(document, accept=False).add_run().add_break()
    inherited = document.styles.add_style("InheritedPictureList", 1)
    inherited.base_style = document.styles["List Bullet"]
    picture(document, accept=False, style=inherited.name)
    begin = document.add_paragraph().add_run()._r
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    begin.append(field)
    picture(document, accept=False)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    document.add_paragraph().add_run()._r.append(end)
    p = picture(document, accept=False)
    p.add_run()._r.append(OxmlElement("w:pict"))
    empty = document.add_paragraph()
    preserved.append(empty)
    document.add_paragraph("图1 图片标题")
    document.save(path)
    paragraphs = list(document._element.body.iter(qn("w:p")))
    return ([paragraphs.index(p._p) for p in wanted], [paragraphs.index(p._p) for p in preserved])


def verify(source: Path, output: Path, wanted: list[int], preserved: list[int]) -> dict:
    pack = create_style_pack(name="图片段落验收", format_id="docx")
    pack.roles["image"]["image"].update(alignment="center", space_before_pt=7, space_after_pt=9)
    result = normalize_docx_direct(source, pack, output, audit_dir=output.parent / "audit")
    after = analyze_docx(output)
    for index in wanted:
        element = result.analysis.elements[index]
        assert element.selected_decision.role == "image", (element.element_id, element.selected_decision)
        assert element.transformation.style_application_status in {"applied", "guarded"}, element.transformation
        assert element.transformation.effective_format_match is True, element.transformation
        actual = after.elements[index]
        assert actual.original_style_id == "DC_Image", actual.original_style_id
        assert actual.effective_format["alignment"] == "center"
        assert actual.effective_format["paragraph_spacing"]["before"] == "140"
        assert actual.effective_format["paragraph_spacing"]["after"] == "180"
        assert actual.effective_format["paragraph_spacing"]["lineRule"] == "auto"
        assert actual.effective_format["paragraph_indent"]["firstLine"] == "0"
    for index in preserved:
        assert result.analysis.elements[index].selected_decision.role != "image", index
    with ZipFile(source) as left, ZipFile(output) as right:
        assert set(left.namelist()) == set(right.namelist())
        for name in left.namelist():
            if name not in {"word/document.xml", "word/styles.xml"}:
                assert left.read(name) == right.read(name), name
        old = etree.fromstring(left.read("word/document.xml"))
        new = etree.fromstring(right.read("word/document.xml"))
        # Canonicalize detached subtrees to ignore declarations added on the document root.
        def drawings(root):
            return [etree.tostring(etree.fromstring(etree.tostring(n)), method="c14n", exclusive=True)
                    for n in root.iter(qn("w:drawing"))]
        assert drawings(old) == drawings(new), "drawing payload, dimensions or relationships changed"
        old_paragraphs, new_paragraphs = list(old.iter(qn("w:p"))), list(new.iter(qn("w:p")))
        for index in wanted:
            a, b = deepcopy(old_paragraphs[index]), deepcopy(new_paragraphs[index])
            for p in (a, b):
                if p.find(qn("w:pPr")) is not None:
                    p.remove(p.find(qn("w:pPr")))
            assert etree.tostring(a, method="c14n", exclusive=True) == etree.tostring(b, method="c14n", exclusive=True)
    return {"ok": True, "mapped_images": len(wanted), "excluded_cases": len(preserved),
            "drawing_xml_and_other_package_parts_unchanged": True, "output": str(output)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--source", type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    source = args.output / "fixture.docx"
    wanted, preserved = fixture(source)
    report = verify(source, args.output / "fixture_V2_图片样式修复.docx", wanted, preserved)
    if args.source:
        sample = analyze_docx(args.source)
        wanted = [i for i, e in enumerate(sample.elements) if e.has_drawing]
        sample_dir = args.output / "sample"
        sample_dir.mkdir()
        report["sample"] = verify(args.source, sample_dir / (args.source.stem + "_V7_图片样式修复.docx"), wanted, [])
        report["sample"]["source_sha256"] = sample.source_sha256
    (args.output / "acceptance.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
