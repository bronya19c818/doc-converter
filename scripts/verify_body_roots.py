"""Issue #4 independent end-to-end corpus, including a frozen Issue #3 baseline."""

from __future__ import annotations

import argparse
import io
import json
import os
import subprocess
import sys
import tarfile
import tempfile
import time
from pathlib import Path

from docx import Document

from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.docx_direct.ooxml_patch import parse_xml_payload, read_zip_member
from doc_converter.normalizer import normalize_markdown
from doc_converter.renderer import render_docx
from doc_converter.style_library import create_style_pack

ROOT = Path(__file__).resolve().parents[1]


def check_output(path, expected):
    doc = Document(path)
    resolver = EffectiveFormatResolver(parse_xml_payload(read_zip_member(path, "word/styles.xml")))
    checks = []
    assert [p.text for p in doc.paragraphs] == [p[0] for p in expected]
    for p, (text, source, target) in zip(doc.paragraphs, expected, strict=True):
        actual = resolver.resolve_paragraph(p._p, paragraph_style_id=None).get("outline_level")
        style = p.style.style_id
        checks.append(
            {
                "text": text,
                "expected_level": target,
                "actual_outline": actual,
                "actual_style": style,
                "required_repair": source != target,
                "correct": (actual == target - 1 and style == f"DC_Heading{target}")
                if target
                else actual in {None, 9},
            }
        )
    return checks


def evaluate(output):
    pack = create_style_pack(name="根层级对照", format_id="docx")
    cases = json.loads((ROOT / "tests/fixtures/root_hierarchy/cases.json").read_text())
    records, timings = [], []
    for case in cases:
        for route in ["markdown", "docx"]:
            dest = output / route / case["name"]
            dest.mkdir(parents=True, exist_ok=True)
            paragraphs = case["paragraphs"]
            before = time.perf_counter()
            if route == "markdown":
                source = "\n\n".join(
                    ("#" * level + " " if level else "") + text for text, level, _ in paragraphs
                )
                normalized = normalize_markdown(source, filename=case["name"] + ".md")
                (dest / "normalized.md").write_text(normalized.markdown)
                render_docx(
                    markdown=normalized.markdown,
                    style_pack=pack,
                    output_path=dest / "output.docx",
                    work_dir=dest / "render",
                    role_decisions=normalized.report["decisions"],
                )
                (dest / "audit.json").write_text(
                    json.dumps(normalized.report, ensure_ascii=False, indent=2)
                )
            else:
                doc = Document()
                for text, level, _ in paragraphs:
                    doc.add_heading(text, level) if level else doc.add_paragraph(text)
                doc.save(dest / "source.docx")
                normalize_docx_direct(
                    dest / "source.docx", pack, dest / "output.docx", audit_dir=dest / "audit"
                )
            timings.append(
                {"case": case["name"], "route": route, "seconds": time.perf_counter() - before}
            )
            records.extend(
                {"case": case["name"], "route": route, **c}
                for c in check_output(dest / "output.docx", paragraphs)
            )
    return {
        "scope": "six-independently-labelled-root-cases-two-routes",
        "records": records,
        "timings": timings,
        "metrics": {
            route: {
                "checks": len(rows),
                "errors": sum(not r["correct"] for r in rows),
                "required_repairs": sum(r["required_repair"] for r in rows),
                "correct_repairs": sum(r["required_repair"] and r["correct"] for r in rows),
            }
            for route in ["markdown", "docx"]
            if (rows := [r for r in records if r["route"] == route])
        },
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--baseline-ref")
    parser.add_argument(
        "--check", action="store_true", help="Fail if current written outputs differ"
    )
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.baseline_ref:
        # Freeze the entire source tree: the direct classifier also imports the
        # shared resolver. Freezing only the scoring file would be a false baseline.
        with tempfile.TemporaryDirectory() as tmp:
            payload = subprocess.check_output(
                ["git", "archive", args.baseline_ref, "src"], cwd=ROOT
            )
            with tarfile.open(fileobj=io.BytesIO(payload)) as archive:
                archive.extractall(tmp, filter="data")
            env = {**os.environ, "PYTHONPATH": str(Path(tmp) / "src")}
            subprocess.run(
                [sys.executable, __file__, "--output", str(args.output / "baseline")],
                env=env,
                check=True,
            )
    report = evaluate(args.output / "current")
    (args.output / "quality.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report["metrics"], ensure_ascii=False))
    if args.check and any(m["errors"] for m in report["metrics"].values()):
        raise SystemExit("Current root hierarchy output differs from independent labels")


if __name__ == "__main__":
    main()
