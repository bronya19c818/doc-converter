import json
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from doc_converter import api as api_module
from doc_converter.api import create_app


def test_render_always_uses_first_import_and_style_switch_is_idempotent(tmp_path, monkeypatch):
    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda *_: False)
    client = TestClient(create_app(data_root=tmp_path))
    packs = client.get('/api/style-packs').json()
    source = '# 实施方案\n\n部门：物资部\n\n一、背景\n\n（一）范围\n\n正文保留123。'
    ready = client.post('/api/jobs', data={'style_pack_id': packs[0]['id'], 'style_pack_version': 1}, files={'source': ('方案.md', source, 'text/markdown')}).json()
    root = tmp_path / 'jobs' / ready['id']
    render_url = f"/api/jobs/{ready['id']}/normalize-and-render"
    first = client.post(render_url, json={})
    assert first.status_code == 200, first.text
    normalized = (root / 'document.normalized.md').read_bytes()
    original = (root / 'source/imported.md').read_bytes()
    (root / 'work/source.md').write_text('# 污染工作文件\n\n污染正文。', encoding='utf-8')
    for index, pack in enumerate([packs[0], packs[1], packs[0]], start=2):
        response = client.post(render_url, json={'style_pack_id': pack['id'], 'style_pack_version': pack['version']})
        assert response.status_code == 200, response.text
        completed = response.json()
        assert completed['render_revision'] == index
        assert completed['style_pack_id'] == pack['id']
        assert (root / 'document.normalized.md').read_bytes() == normalized
        assert (root / 'source/imported.md').read_bytes() == original
        document = Document(BytesIO(client.get(f"/api/jobs/{ready['id']}/artifacts/{completed['output_name']}").content))
        assert sum(p.text == '实施方案' for p in document.paragraphs) == 1
        assert any(p.text == '正文保留123。' for p in document.paragraphs)
        assert not any('污染' in p.text for p in document.paragraphs)
    report = json.loads((root / 'normalization-report.json').read_text())
    assert report['decisions'] and report['counts']['changed'] > 0
    assert 'confirmation.json' not in first.json()['artifacts']


def test_removed_write_and_confirmation_contracts_fail_explicitly(tmp_path, monkeypatch):
    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda *_: False)
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.get('/api/style-packs').json()[0]
    ready = client.post('/api/jobs', data={'style_pack_id': pack['id'], 'style_pack_version': 1}, files={'source': ('方案.md', '# 方案\n\n正文。', 'text/markdown')}).json()
    assert ready['suggestions'] == []
    assert not (tmp_path / 'jobs' / ready['id'] / 'work/source.md').exists()
    assert client.put(f"/api/jobs/{ready['id']}/draft", json={'source_markdown': '# 修改'}).status_code == 410
    for payload in [{'source_markdown': '# 修改'}, {'confirmed_roles': {'1': 'document-title'}}]:
        assert client.post(f"/api/jobs/{ready['id']}/normalize-and-render", json=payload).status_code == 422


def test_old_job_without_imported_md_recovers_from_original_upload(tmp_path, monkeypatch):
    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda *_: False)
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.get('/api/style-packs').json()[0]
    ready = client.post('/api/jobs', data={'style_pack_id': pack['id'], 'style_pack_version': 1}, files={'source': ('恢复.md', '# 恢复方案\n\n正文。', 'text/markdown')}).json()
    root = tmp_path / 'jobs' / ready['id']
    (root / 'source/imported.md').unlink()
    (root / 'source.md').write_text('# 旧编辑草稿\n\n不得用作源输入。', encoding='utf-8')
    response = client.post(f"/api/jobs/{ready['id']}/normalize-and-render", json={})
    assert response.status_code == 200
    assert response.json()['output_name'] == '恢复方案.docx'


def test_optional_diagnostic_write_failure_does_not_block_docx(tmp_path, monkeypatch):
    from pathlib import Path

    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda *_: False)
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.get('/api/style-packs').json()[0]
    ready = client.post('/api/jobs', data={'style_pack_id': pack['id'], 'style_pack_version': 1}, files={'source': ('方案.md', '# 方案\n\n正文。', 'text/markdown')}).json()
    write_text = Path.write_text

    def deny_report(path, *args, **kwargs):
        if path.name == 'normalization-report.json':
            raise OSError('diagnostic disk error')
        return write_text(path, *args, **kwargs)

    monkeypatch.setattr(Path, 'write_text', deny_report)
    response = client.post(f"/api/jobs/{ready['id']}/normalize-and-render", json={})
    assert response.status_code == 200, response.text
    assert 'normalization-report.json' not in response.json()['artifacts']
    assert response.json()['processing_warnings']


def test_failed_regeneration_does_not_overwrite_last_successful_docx(tmp_path, monkeypatch):
    from doc_converter.renderer import RendererError

    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda *_: False)
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.get('/api/style-packs').json()[0]
    ready = client.post('/api/jobs', data={'style_pack_id': pack['id'], 'style_pack_version': 1}, files={'source': ('方案.md', '# 方案\n\n正文。', 'text/markdown')}).json()
    url = f"/api/jobs/{ready['id']}/normalize-and-render"
    first = client.post(url, json={}).json()
    output = tmp_path / 'jobs' / ready['id'] / first['output_name']
    original = output.read_bytes()

    def fail_render(*, output_path, **kwargs):
        output_path.write_bytes(b'incomplete DOCX')
        raise RendererError('render failed')

    monkeypatch.setattr(api_module, 'render_docx', fail_render)
    assert client.post(url, json={}).status_code == 422
    assert output.read_bytes() == original
