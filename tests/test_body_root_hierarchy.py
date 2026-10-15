"""Final output contracts: an offset is repaired without flattening the tree."""

from pathlib import Path

import pytest
from docx import Document

from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.normalizer import normalize_markdown as _normalize_markdown
from doc_converter.renderer import render_docx
from doc_converter.style_library import create_style_pack

FIXTURE = Path(__file__).parent / "fixtures/root_hierarchy/complete.md"


def normalize_markdown(source):
    return _normalize_markdown(source, filename="根层级.md")


def heading_roles(result):
    return [d["role"] for d in result.report["decisions"] if d["role"].startswith("heading-")]


@pytest.mark.parametrize("root,child", [(2, 3), (3, 4), (2, 4)])
def test_markdown_complete_tree_shift_preserves_gaps(root, child):
    source = (
        FIXTURE.read_text()
        .replace("####", "@CHILD@")
        .replace("##", "#" * root)
        .replace("@CHILD@", "#" * child)
    )
    r = normalize_markdown(source)
    assert heading_roles(r) == ["heading-1", f"heading-{child - root + 1}", "heading-1"]
    assert normalize_markdown(r.markdown).markdown == r.markdown
    audit = r.report["context"]["body_root_normalization"]
    assert audit["offset"] == root - 1
    assert len(audit["affected_ids"]) == 3


@pytest.mark.parametrize("mode", ["legacy", "shadow", "scored", "hybrid"])
@pytest.mark.parametrize("root,child", [(2, 3), (3, 4), (2, 4)])
def test_written_docx_tree_and_second_conversion(tmp_path, root, child, mode):
    pack = create_style_pack(name="根层级", format_id="docx")
    d = Document()
    for text, level in [
        ("总体要求", root),
        ("本文件说明完整的实施范围。", None),
        ("适用范围", child),
        ("保留原有层级差。", None),
        ("实施安排", root),
        ("按阶段落实交付工作。", None),
    ]:
        d.add_heading(text, level) if level else d.add_paragraph(text)
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    d.save(source)
    r = normalize_docx_direct(
        source, pack, output, audit_dir=tmp_path / "audit", decision_mode=mode
    )
    expected = [1, child - root + 1, 1]
    elements = [e for e in r.analysis.elements if e.features["native_heading_level"]]
    assert [e.selected_decision.role for e in elements] == [f"heading-{n}" for n in expected]
    assert [e.transformation.effective_format_after["outline_level"] for e in elements] == [
        n - 1 for n in expected
    ]
    written = Document(output)
    assert [
        p.style.style_id
        for p in written.paragraphs
        if p.text in {"总体要求", "适用范围", "实施安排"}
    ] == [f"DC_Heading{n}" for n in expected]
    assert [p.text for p in written.paragraphs] == [p.text for p in d.paragraphs]
    twice = normalize_docx_direct(
        output, pack, tmp_path / "twice.docx", audit_dir=tmp_path / "audit2", decision_mode=mode
    )
    assert [e.selected_decision.role for e in twice.analysis.elements] == [
        e.selected_decision.role for e in r.analysis.elements
    ]


def test_markdown_actual_written_styles_and_outline(tmp_path):
    pack = create_style_pack(name="根层级", format_id="docx")
    r = normalize_markdown(FIXTURE.read_text())
    rendered = render_docx(
        markdown=r.markdown,
        style_pack=pack,
        output_path=tmp_path / "out.docx",
        work_dir=tmp_path / "work",
        role_decisions=r.report["decisions"],
    )
    assert rendered.verification["passed"]
    checks = [c for c in rendered.verification["checks"] if c["role"].startswith("heading-")]
    assert [(c["actual_style"], c["actual_outline"]) for c in checks] == [
        ("DC_Heading1", 0),
        ("DC_Heading3", 2),
        ("DC_Heading1", 0),
    ]


@pytest.mark.parametrize(
    "source,expected",
    [
        ("# 总体要求\n\n### 适用范围\n\n正文。", ["heading-1", "heading-3"]),
        ("## 第五章 实施要求（节选）\n\n### 适用范围\n\n正文。", ["heading-2", "heading-3"]),
        ("## 5.1 业务范围\n\n### 5.1.1 内容\n\n正文。", ["heading-2", "heading-3"]),
        ("## 单独业务主题\n\n正文。\n\n### 深层内容\n\n正文。", ["heading-2", "heading-3"]),
        ("普通正文，没有标题。", []),
    ],
)
def test_incomplete_or_existing_root_is_not_mechanically_shifted(source, expected):
    assert heading_roles(normalize_markdown(source)) == expected


def test_appendix_h1_does_not_mask_offset_body():
    r = normalize_markdown(FIXTURE.read_text() + "\n# 附录A 资料\n\n## 对照表\n\n说明。")
    assert heading_roles(r) == ["heading-1", "heading-3", "heading-1", "heading-1", "heading-2"]


def test_independent_title_is_kept_outside_the_body_tree():
    r = normalize_markdown("# 年度报告\n\n" + FIXTURE.read_text())
    assert r.metadata["title"] == "年度报告"
    assert heading_roles(r) == ["heading-1", "heading-3", "heading-1"]


def test_intro_is_body_heading_even_before_numbered_child():
    r = normalize_markdown("# 总体要求\n\n## 1 范围\n\n范围正文。\n\n## 2 职责\n\n职责正文。")
    assert "title" not in r.metadata
    assert heading_roles(r) == ["heading-1", "heading-2", "heading-2"]


def test_direct_excerpt_keeps_native_anchors_after_candidate_compression(tmp_path):
    d = Document()
    d.add_heading("总体要求（节选）", 2)
    d.add_heading("适用范围", 3)
    d.add_paragraph("范围正文。")
    d.add_heading("实施安排", 2)
    d.add_paragraph("安排正文。")
    source = tmp_path / "excerpt.docx"
    d.save(source)
    r = normalize_docx_direct(
        source, create_style_pack(name="节选", format_id="docx"), tmp_path / "out.docx"
    )
    assert [
        e.selected_decision.role for e in r.analysis.elements if e.features["native_heading_level"]
    ] == ["heading-2", "heading-3", "heading-2"]
    assert r.analysis.document_baselines["body_root_normalization"]["reason"] == "explicit-excerpt"


def test_independent_regimes_are_not_merged_into_one_shift():
    from doc_converter.structure.body_roots import FinalParagraph, normalize_body_roots

    levels, audit = normalize_body_roots(
        [
            FinalParagraph("a", "总体要求", "heading-2"),
            FinalParagraph("b", "这里是完整正文。", "body-paragraph"),
            FinalParagraph("c", "独立章节", "heading-3", facts={"hierarchy_reset_marker": True}),
            FinalParagraph("d", "这里是完整正文。", "body-paragraph"),
        ]
    )
    assert levels == {"a": 2, "c": 3}
    assert audit["reason"] == "independent-roots-at-different-levels"


def test_long_tree_preserves_all_edges_content_and_style_switch(tmp_path):
    source = (FIXTURE.parent / "long_complete.md").read_text(encoding="utf-8")
    first = normalize_markdown(source)
    assert heading_roles(first) == ["heading-1"] + [
        r for _ in range(60) for r in ["heading-1", "heading-3"]
    ]
    assert normalize_markdown(first.markdown).markdown == first.markdown
    tree = first.report["context"]["final_heading_tree"]
    assert all(tree[i + 1]["parent_id"] == tree[i]["element_id"] for i in range(1, len(tree), 2))
    for name in ["标准", "替换样式"]:
        pack = create_style_pack(name=name, format_id="docx")
        out = render_docx(
            markdown=first.markdown,
            style_pack=pack,
            output_path=tmp_path / (name + ".docx"),
            work_dir=tmp_path / name,
            role_decisions=first.report["decisions"],
        )
        assert out.verification["passed"]
        assert (
            len([c for c in out.verification["checks"] if c["role"].startswith("heading-")]) == 121
        )


@pytest.mark.parametrize(
    "titles",
    [
        ("1. 总体要求", "1.1 适用范围", "2. 实施安排"),
        ("一、总体要求", "（一）适用范围", "二、实施安排"),
    ],
)
def test_numbered_native_gap_matches_markdown_and_docx(tmp_path, titles):
    paragraphs = [
        (titles[0], 2),
        ("这是总体要求的业务正文。", 0),
        (titles[1], 4),
        ("这是范围的业务正文。", 0),
        (titles[2], 2),
        ("这是安排的业务正文。", 0),
    ]
    source_md = "\n\n".join(
        ("#" * level + " " if level else "") + text for text, level in paragraphs
    )
    assert heading_roles(normalize_markdown(source_md)) == ["heading-1", "heading-3", "heading-1"]
    d = Document()
    for text, level in paragraphs:
        d.add_heading(text, level) if level else d.add_paragraph(text)
    source = tmp_path / "numbered.docx"
    d.save(source)
    r = normalize_docx_direct(
        source, create_style_pack(name="编号", format_id="docx"), tmp_path / "out.docx"
    )
    assert [
        e.selected_decision.role for e in r.analysis.elements if e.features["native_heading_level"]
    ] == ["heading-1", "heading-3", "heading-1"]
    twice = normalize_docx_direct(
        r.output_path,
        create_style_pack(name="再次转换", format_id="docx"),
        tmp_path / "twice.docx",
        audit_dir=tmp_path / "twice-audit",
    )
    assert [e.selected_decision.role for e in twice.analysis.elements] == [
        e.selected_decision.role for e in r.analysis.elements
    ]


def test_appendix_word_in_body_or_business_heading_is_not_boundary():
    source = "## 总体要求\n\n附件应当包括完整的技术说明。\n\n#### 适用范围\n\n范围正文。\n\n## 附件管理要求\n\n附件管理正文。\n\n## 实施安排\n\n安排正文。"
    assert heading_roles(normalize_markdown(source)) == [
        "heading-1",
        "heading-3",
        "heading-1",
        "heading-1",
    ]


def test_field_only_heading_vetoes_global_offset_before_ooxml_guard(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    p = d.add_heading("", 2)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DOCVARIABLE body_heading")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "总体要求"
    run.append(text)
    field.append(run)
    p._p.append(field)
    d.add_paragraph("这是总体要求的业务正文。")
    d.add_heading("适用范围", 4)
    d.add_paragraph("这是适用范围的业务正文。")
    d.add_heading("实施安排", 2)
    d.add_paragraph("这是实施安排的业务正文。")
    source = tmp_path / "fields.docx"
    d.save(source)
    r = normalize_docx_direct(
        source,
        create_style_pack(name="保护", format_id="docx"),
        tmp_path / "out.docx",
        decision_mode="legacy",
    )
    audit = r.analysis.document_baselines["body_root_normalization"]
    assert audit["offset"] == 0
    assert audit["reason"] == "unmapped-heading"
    written = Document(r.output_path)
    assert written.paragraphs[0].style.style_id == "Heading2"
    assert written.paragraphs[2].style.style_id == "DC_Heading4"
    assert written.paragraphs[4].style.style_id == "DC_Heading2"
    assert written.paragraphs[0]._p.find(qn("w:fldSimple")) is not None


@pytest.mark.parametrize("mode", ["legacy", "shadow", "scored", "hybrid"])
@pytest.mark.parametrize("mixed", [False, True])
def test_direct_outline_anchors_keep_peer_tree_and_gap(tmp_path, mode, mixed):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    for text, level in [
        ("总体要求", 2),
        ("这是总体要求的业务正文。", 0),
        ("适用范围", 4),
        ("这是范围的业务正文。", 0),
        ("实施安排", 2),
        ("这是安排的业务正文。", 0),
    ]:
        native_child = mixed and text == "适用范围"
        p = d.add_heading(text, level) if native_child else d.add_paragraph(text)
        if level and not native_child:
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), str(level - 1))
            p._p.get_or_add_pPr().append(outline)
    source = tmp_path / "outline.docx"
    d.save(source)
    r = normalize_docx_direct(
        source,
        create_style_pack(name="大纲", format_id="docx"),
        tmp_path / "out.docx",
        decision_mode=mode,
    )
    headings = [e for e in r.analysis.elements if e.features["root_source_level"]]
    assert [e.selected_decision.role for e in headings] == ["heading-1", "heading-3", "heading-1"]
    assert [e.transformation.effective_format_after["outline_level"] for e in headings] == [0, 2, 0]


@pytest.mark.parametrize("route", ["markdown", "docx"])
def test_later_named_chapter_is_excerpt_without_literal_excerpt_label(tmp_path, route):
    paragraphs = [
        ("第五章 实施要求", 2),
        ("这是实施要求的业务正文。", 0),
        ("适用范围", 3),
        ("这是范围的业务正文。", 0),
        ("第六章 交付要求", 2),
        ("这是交付要求的业务正文。", 0),
    ]
    if route == "markdown":
        source = "\n\n".join(
            ("#" * level + " " if level else "") + text for text, level in paragraphs
        )
        r = normalize_markdown(source)
        assert heading_roles(r) == ["heading-2", "heading-3", "heading-2"]
        assert r.report["context"]["body_root_normalization"]["reason"] == "later-chapter-excerpt"
    else:
        d = Document()
        for text, level in paragraphs:
            d.add_heading(text, level) if level else d.add_paragraph(text)
        source = tmp_path / "chapters.docx"
        d.save(source)
        r = normalize_docx_direct(
            source, create_style_pack(name="章节", format_id="docx"), tmp_path / "out.docx"
        )
        assert [
            e.selected_decision.role
            for e in r.analysis.elements
            if e.features["native_heading_level"]
        ] == ["heading-2", "heading-3", "heading-2"]
        assert [
            e.transformation.effective_format_after["outline_level"]
            for e in r.analysis.elements
            if e.features["native_heading_level"]
        ] == [1, 2, 1]
