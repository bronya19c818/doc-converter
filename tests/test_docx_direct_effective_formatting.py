from __future__ import annotations

from pathlib import Path
from typing import Any, Callable
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from doc_converter.docx_direct.effective_format import EffectiveFormatResolver

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def _w(tag: str) -> str:
    return f'{{{W_NS}}}{tag}'


def _read_member(path: Path, member: str) -> bytes:
    with ZipFile(path) as archive:
        return archive.read(member)


def _first_paragraph(path: Path) -> Any:
    root = parse_xml(_read_member(path, 'word/document.xml'))
    return root.find('.//' + qn('w:body') + '/' + qn('w:p'))


def _add_style(
    root: Any,
    style_id: str,
    *,
    style_type: str = 'paragraph',
    based_on: str | None = None,
    size: str | None = None,
    alignment: str | None = None,
    outline: str | None = None,
    fonts: dict[str, str] | None = None,
) -> None:
    style = OxmlElement('w:style')
    style.set(qn('w:type'), style_type)
    style.set(qn('w:styleId'), style_id)
    name = OxmlElement('w:name')
    name.set(qn('w:val'), style_id)
    style.append(name)
    if based_on is not None:
        based = OxmlElement('w:basedOn')
        based.set(qn('w:val'), based_on)
        style.append(based)
    if alignment is not None or outline is not None:
        ppr = OxmlElement('w:pPr')
        if alignment is not None:
            jc = OxmlElement('w:jc')
            jc.set(qn('w:val'), alignment)
            ppr.append(jc)
        if outline is not None:
            outline_node = OxmlElement('w:outlineLvl')
            outline_node.set(qn('w:val'), outline)
            ppr.append(outline_node)
        style.append(ppr)
    if size is not None or fonts:
        rpr = OxmlElement('w:rPr')
        if size is not None:
            size_node = OxmlElement('w:sz')
            size_node.set(qn('w:val'), size)
            rpr.append(size_node)
        if fonts:
            fonts_node = OxmlElement('w:rFonts')
            for key, value in fonts.items():
                fonts_node.set(qn('w:' + key), value)
            rpr.append(fonts_node)
        style.append(rpr)
    root.append(style)


def _set_doc_defaults(root: Any, *, size: str | None = None, alignment: str | None = None) -> None:
    defaults = root.find(qn('w:docDefaults'))
    if defaults is None:
        defaults = OxmlElement('w:docDefaults')
        root.insert(0, defaults)
    if size is not None:
        rpr_default = defaults.find(qn('w:rPrDefault'))
        if rpr_default is None:
            rpr_default = OxmlElement('w:rPrDefault')
            defaults.append(rpr_default)
        rpr = rpr_default.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr')
            rpr_default.append(rpr)
        size_node = rpr.find(qn('w:sz'))
        if size_node is None:
            size_node = OxmlElement('w:sz')
            rpr.append(size_node)
        size_node.set(qn('w:val'), size)
    if alignment is not None:
        ppr_default = defaults.find(qn('w:pPrDefault'))
        if ppr_default is None:
            ppr_default = OxmlElement('w:pPrDefault')
            defaults.append(ppr_default)
        ppr = ppr_default.find(qn('w:pPr'))
        if ppr is None:
            ppr = OxmlElement('w:pPr')
            ppr_default.append(ppr)
        jc = ppr.find(qn('w:jc'))
        if jc is None:
            jc = OxmlElement('w:jc')
            ppr.append(jc)
        jc.set(qn('w:val'), alignment)


def _remove_doc_default_size(root: Any) -> None:
    defaults = root.find(qn('w:docDefaults'))
    rpr_default = defaults.find(qn('w:rPrDefault')) if defaults is not None else None
    rpr = rpr_default.find(qn('w:rPr')) if rpr_default is not None else None
    if rpr is not None:
        for tag in ('w:sz', 'w:szCs'):
            size = rpr.find(qn(tag))
            if size is not None:
                rpr.remove(size)


def _make_docx(
    tmp_path: Path,
    *,
    style_mutator: Callable[[Any], None] | None = None,
    paragraph_style: str | None = None,
    text: str | None = None,
    runs: list[dict[str, Any]] | None = None,
) -> Path:
    document = Document()
    paragraph = document.add_paragraph()
    if paragraph_style is not None:
        ppr = paragraph._p.get_or_add_pPr()
        pstyle = OxmlElement('w:pStyle')
        pstyle.set(qn('w:val'), paragraph_style)
        ppr.append(pstyle)
    if runs is None:
        runs = [{'text': text or ''}]
    for spec in runs:
        run = paragraph.add_run(spec['text'])
        if spec.get('size') is not None:
            run.font.size = Pt(spec['size'])
        if spec.get('superscript'):
            run.font.superscript = True
        if spec.get('subscript'):
            run.font.subscript = True
    source = tmp_path / 'source.docx'
    document.save(source)
    if style_mutator is None:
        return source

    with ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    styles_root = parse_xml(members['word/styles.xml'])
    style_mutator(styles_root)
    members['word/styles.xml'] = etree.tostring(
        styles_root,
        encoding='UTF-8',
        xml_declaration=True,
        standalone=True,
    )
    patched = tmp_path / 'patched.docx'
    with ZipFile(patched, 'w', ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return patched


def _style_with_values(style_id: str, **kwargs: str) -> Callable[[Any], None]:
    def mutate(root: Any) -> None:
        _add_style(root, style_id, **kwargs)

    return mutate


def test_resolver_inherits_size_and_alignment_from_paragraph_style(tmp_path: Path) -> None:
    source = _make_docx(
        tmp_path,
        style_mutator=_style_with_values(
            'TitleInherited',
            size='44',
            alignment='center',
        ),
        paragraph_style='TitleInherited',
        text='继承格式标题',
    )
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(_read_member(source, 'word/styles.xml'))

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id='TitleInherited')

    assert result['font_size_pt'] == 22
    assert result['alignment'] == 'center'
    assert result['style_chain'] == ['TitleInherited']


def test_resolver_inherits_doc_defaults_when_style_has_no_direct_values(tmp_path: Path) -> None:
    def mutate(root: Any) -> None:
        _set_doc_defaults(root, size='28', alignment='center')

    source = _make_docx(tmp_path, style_mutator=mutate, text='默认格式段落')
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(_read_member(source, 'word/styles.xml'))

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id=None)

    assert result['font_size_pt'] == 14
    assert result['alignment'] == 'center'


def test_resolver_follows_based_on_chain_and_keeps_heading_as_evidence(tmp_path: Path) -> None:
    def mutate(root: Any) -> None:
        _add_style(root, 'BaseHeading', size='32', alignment='center', outline='1')
        _add_style(root, 'CustomTitle', based_on='BaseHeading')

    source = _make_docx(
        tmp_path,
        style_mutator=mutate,
        paragraph_style='CustomTitle',
        text='自定义继承标题',
    )
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(_read_member(source, 'word/styles.xml'))

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id='CustomTitle')

    assert result['font_size_pt'] == 16
    assert result['alignment'] == 'center'
    assert result['outline_level'] == 1
    assert result['style_chain'] == ['BaseHeading', 'CustomTitle']
    assert result['based_on_heading'] is True


def test_resolver_stops_a_based_on_cycle_without_recursing_forever(tmp_path: Path) -> None:
    def mutate(root: Any) -> None:
        _remove_doc_default_size(root)
        _add_style(root, 'CycleA', based_on='CycleB')
        _add_style(root, 'CycleB', based_on='CycleA')

    source = _make_docx(
        tmp_path,
        style_mutator=mutate,
        paragraph_style='CycleA',
        text='循环样式',
    )
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(
        _read_member(source, 'word/styles.xml'),
        max_depth=4,
    )

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id='CycleA')

    assert result['font_size_pt'] is None
    assert len(result['style_chain']) <= 4


def test_resolver_prefers_east_asia_font_and_falls_back_in_order(tmp_path: Path) -> None:
    def mutate(root: Any) -> None:
        _add_style(
            root,
            'EastAsiaStyle',
            fonts={'ascii': 'AsciiFont', 'hAnsi': 'HAnsiFont', 'eastAsia': 'EastAsiaFont'},
        )

    source = _make_docx(
        tmp_path,
        style_mutator=mutate,
        paragraph_style='EastAsiaStyle',
        text='中文字体',
    )
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(_read_member(source, 'word/styles.xml'))

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id='EastAsiaStyle')

    assert result['font_name'] == 'EastAsiaFont'


def test_resolver_uses_character_weighted_dominant_size_not_maximum(tmp_path: Path) -> None:
    source = _make_docx(
        tmp_path,
        runs=[
            {'text': '正文正文正文正文', 'size': 12},
            {'text': '大', 'size': 30},
        ],
    )
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(_read_member(source, 'word/styles.xml'))

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id=None)

    assert result['dominant_font_size_pt'] == 12
    assert result['dominant_font_size_ratio'] == 8 / 9
    assert result['max_font_size_pt'] == 30


def test_resolver_excludes_superscript_from_dominant_size_but_keeps_evidence(
    tmp_path: Path,
) -> None:
    source = _make_docx(
        tmp_path,
        runs=[
            {'text': '正文正文正文', 'size': 12},
            {'text': '2', 'size': 30, 'superscript': True},
        ],
    )
    paragraph = _first_paragraph(source)
    resolver = EffectiveFormatResolver.from_styles_xml(_read_member(source, 'word/styles.xml'))

    result = resolver.resolve_paragraph(paragraph, paragraph_style_id=None)

    assert result['dominant_font_size_pt'] == 12
    assert result['max_font_size_pt'] == 30
    assert any(
        run['vert_align'] == 'superscript'
        for run in result['run_statistics']['runs']
    )
