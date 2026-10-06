"""Real DOCX/API/CLI regressions, written before semantic implementation.

Catch dirty Heading author rows, overly broad name matching, metadata without
context, OLE/VML/OMML loss, formula numbering pollution and risk-review omissions.
"""
import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.style_library import create_style_pack
from scripts.verify_docx_selected import verify_sample

FIXTURES = Path(__file__).parent / "fixtures/docx_direct_regression"
TRUTH = json.loads((FIXTURES / "ground_truth.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("name", list(TRUTH))
def test_semantic_fixture_ground_truth_through_api(name, tmp_path, monkeypatch):
    monkeypatch.setenv("DOC_CONVERTER_LIBREOFFICE", str(tmp_path / "no-preview"))
    verify_sample(FIXTURES / name, tmp_path / "api", TRUTH[name])
    analysis = json.loads((tmp_path / "api/docx-analysis.json").read_text(encoding="utf-8"))
    if name == "complex_numbering.docx":
        assert analysis["document_baselines"]["numbering_family_levels"]["ARABIC_BRACKET"] == 3
        for e in analysis["elements"]:
            if e["text"] in {"（1）", "（2）", "（3）", "（4）"}:
                assert e["features"]["has_complex_object"]
                assert e["has_embedded_object"] or e["has_math"]
                assert not any(v["feature"] == "numbering_family_match"
                               for c in e["candidates"] for v in c["evidence"])
    else:
        e = next(e for e in analysis["elements"] if e["text"] == "林硕，张圣洁")
        assert e["features"]["author_name_list_pattern"]
        assert e["features"]["author_name_count"] == 2


@pytest.mark.parametrize("text,count", [
    ("林硕，张圣洁", 2), ("林硕、张圣洁", 2), ("林硕, 张圣洁", 2),
    ("林硕 张圣洁", 2), ("张三 李四 王五", 3),
    ("林硕，张圣洁。", 0), ("林硕，Smith", 0),
    ("张三 李四 王五 张三 李四 王五 张三 李四 王五", 0),
])
def test_author_list_formats_through_docx_conversion(text, count, tmp_path):
    doc = Document(FIXTURES / "author_list_dirty_heading.docx")
    doc.paragraphs[1].text = text
    source = tmp_path / "source.docx"
    doc.save(source)
    result = normalize_docx_direct(source, create_style_pack(name="姓名验收", format_id="docx"),
                                   tmp_path / "normalized.docx")
    e = next(e for e in result.analysis.elements if e.text == text)
    assert e.features["author_name_count"] == count
    assert e.features["author_name_list_pattern"] == bool(count)
    if count:
        assert e.selected_decision.role == "author"
        assert e.selected_decision.action in {"mapped", "uncertain"}
    else:
        assert e.selected_decision.role != "author"


def test_name_list_without_metadata_context_stays_preserved(tmp_path):
    doc = Document(FIXTURES / "author_list_dirty_heading.docx")
    for p in list(doc.paragraphs)[2:4]:
        p._p.getparent().remove(p._p)
    source = tmp_path / "unconfirmed.docx"
    doc.save(source)
    result = normalize_docx_direct(source, create_style_pack(name="上下文验收", format_id="docx"),
                                   tmp_path / "normalized.docx")
    e = next(e for e in result.analysis.elements if e.text == "林硕，张圣洁")
    assert (e.selected_decision.role, e.selected_decision.action) == ("unknown", "preserved")


@pytest.mark.parametrize("followup", ["date-after-body-heading", "native-subtitle"])
def test_author_context_does_not_cross_body_and_accepts_native_subtitle(followup, tmp_path):
    doc = Document(FIXTURES / "author_list_dirty_heading.docx")
    for p in list(doc.paragraphs)[2:4]:
        p._p.getparent().remove(p._p)
    if followup == "date-after-body-heading":
        date = doc.add_paragraph("2026年10月")
        doc.paragraphs[2]._p.addnext(date._p)
    else:
        subtitle = doc.add_paragraph("业务研究副标题", style="Subtitle")
        doc.paragraphs[2]._p.addprevious(subtitle._p)
    source = tmp_path / "context.docx"
    doc.save(source)
    result = normalize_docx_direct(source, create_style_pack(name="上下文边界", format_id="docx"),
                                   tmp_path / "normalized.docx")
    e = next(e for e in result.analysis.elements if e.text == "林硕，张圣洁")
    if followup == "date-after-body-heading":
        assert (e.selected_decision.role, e.selected_decision.action) == ("unknown", "preserved")
    else:
        assert e.selected_decision.role == "author"
        assert e.selected_decision.action in {"mapped", "uncertain"}


def test_legacy_complex_mapping_is_guarded_and_risk_cli_marks_high(tmp_path):
    result = normalize_docx_direct(
        FIXTURES / "complex_numbering.docx", create_style_pack(name="风险验收", format_id="docx"),
        tmp_path / "normalized.docx", decision_mode="legacy",
    )
    e = next(e for e in result.analysis.elements if e.text == "（1）")
    assert e.transformation.transform_safety in {"guarded", "preserve"}
    review = tmp_path / "decision-risk-review.json"
    subprocess.run([
        sys.executable, str(Path(__file__).parents[1] / "scripts/review_docx_decision_risks.py"),
        "--analysis", str(result.analysis_path), "--report", str(result.normalization_report_path),
        "--output", str(review),
    ], check=True)
    payload = json.loads(review.read_text(encoding="utf-8"))
    risk = next(r for r in payload["elements"] if r["element_id"] == e.element_id)
    assert risk["risk_level"] == "HIGH"
    assert "complex_object_mapped" in risk["risk_flags"]
    assert risk["context_before"] and risk["context_after"]


@pytest.mark.parametrize("problem", ["missing-hash", "mode-mismatch"])
def test_risk_cli_rejects_unverifiable_audit_pairs(problem, tmp_path):
    result = normalize_docx_direct(
        FIXTURES / "author_list_dirty_heading.docx",
        create_style_pack(name="审计配对验收", format_id="docx"), tmp_path / "normalized.docx",
    )
    analysis = json.loads(result.analysis_path.read_text(encoding="utf-8"))
    report = json.loads(result.normalization_report_path.read_text(encoding="utf-8"))
    if problem == "missing-hash":
        analysis.pop("source_sha256")
        report.pop("source_sha256")
    else:
        report["classification_mode"] = "legacy"
    result.analysis_path.write_text(json.dumps(analysis), encoding="utf-8")
    result.normalization_report_path.write_text(json.dumps(report), encoding="utf-8")
    output = tmp_path / "rejected-review.json"
    process = subprocess.run([
        sys.executable, str(Path(__file__).parents[1] / "scripts/review_docx_decision_risks.py"),
        "--analysis", str(result.analysis_path), "--report", str(result.normalization_report_path),
        "--output", str(output),
    ], capture_output=True, text=True)
    assert process.returncode != 0
    assert not output.exists()
