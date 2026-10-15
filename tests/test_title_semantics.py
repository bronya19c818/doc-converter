"""Independent title/subtitle/H1 contracts, checked at final and written output."""

import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from doc_converter.docx_direct.analyzer import analyze_docx
from doc_converter.docx_direct.classifier import classify_document
from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.docx_direct.ooxml_patch import parse_xml_payload, read_zip_member
from doc_converter.docx_styles import ROLE_TO_STYLE_ID
from doc_converter.normalizer import normalize_markdown
from doc_converter.style_library import create_style_pack

CASES = json.loads((Path(__file__).parent / "fixtures/title_semantics/cases.json").read_text())


def markdown_source(case):
    return "\n\n".join(
        ("# " if style in {"Heading1", "Title"} else "") + text
        for text, style, _ in case["paragraphs"]
    )


def make_docx(case, path):
    doc = Document()
    for text, style, _ in case["paragraphs"]:
        doc.add_paragraph(text, style=style.replace("Heading1", "Heading 1"))
    doc.save(path)
    return path


def canonical(text):
    return "".join(text.split()).removeprefix("作者：").removeprefix("单位：")


def markdown_roles(result, case):
    roles = {}
    for d in result.report["decisions"]:
        role = (
            "organization-department" if d["role"] in {"organization", "department"} else d["role"]
        )
        for paragraph in str(d["before"]).removeprefix("# ").split("\n\n"):
            roles[canonical(paragraph)] = role
    return [roles.get(canonical(text), "missing") for text, _, _ in case["paragraphs"]]


@pytest.mark.parametrize("case", CASES, ids=lambda c: c["name"])
@pytest.mark.parametrize("route", ["markdown", "docx"])
def test_final_cover_roles_and_repeated_conversion(case, route, tmp_path):
    expected = [r for _, _, r in case["paragraphs"]]
    if route == "markdown":
        one = normalize_markdown(markdown_source(case), filename="标题样本.md")
        assert markdown_roles(one, case) == expected
        two = normalize_markdown(one.markdown, filename="标题样本.md")
        assert two.markdown == one.markdown
    else:
        path = make_docx(case, tmp_path / "source.docx")
        result = classify_document(analyze_docx(path))
        assert [e.selected_decision.role for e in result.elements] == expected


def test_render_preserves_cover_paragraph_order_and_multiline_title(tmp_path):
    from doc_converter.renderer import render_docx
    from doc_converter.style_library import create_style_pack

    for name in ["cover-with-metadata", "multiline-title", "manual-break-title"]:
        case = next(c for c in CASES if c["name"] == name)
        normalized = normalize_markdown(markdown_source(case), filename="封面.md")
        out = render_docx(
            markdown=normalized.markdown,
            style_pack=create_style_pack(name="封面", format_id="docx"),
            output_path=tmp_path / (name + ".docx"),
            work_dir=tmp_path / name,
            role_decisions=normalized.report["decisions"],
        )
        assert [canonical(p.text) for p in Document(out.output_path).paragraphs] == [
            canonical(t) for t, _, _ in case["paragraphs"]
        ]
        assert out.verification["passed"]


@pytest.mark.parametrize("route", ["markdown", "docx"])
@pytest.mark.parametrize("viewpoint_child", [False, True])
def test_thematic_h1_owning_h2_children_is_not_a_cover(tmp_path, route, viewpoint_child):
    paragraphs = [
        ("面向应急响应的区域资源共享", 1),
        ("研究方法", 2),
        ("这里是方法的完整正文。", 0),
        ("研究结果", 2),
        ("这里是结果的完整正文。", 0),
    ]
    expected = ["heading-1", "heading-2", "body-paragraph", "heading-2", "body-paragraph"]
    if viewpoint_child:
        paragraphs = [(text, level + 1 if level else 0) for text, level in paragraphs]
        paragraphs[0] = ("面向应急响应的区域资源共享", 1)
        paragraphs.insert(1, ("基于区域网络的实践观察", 2))
        expected = [
            "heading-1",
            "heading-2",
            "heading-3",
            "body-paragraph",
            "heading-3",
            "body-paragraph",
        ]
    if route == "markdown":
        result = normalize_markdown(
            "\n\n".join(("#" * level + " " if level else "") + text for text, level in paragraphs),
            filename="正文.md",
        )
        assert [d["role"] for d in result.report["decisions"]] == expected
    else:
        doc = Document()
        for text, level in paragraphs:
            doc.add_heading(text, level) if level else doc.add_paragraph(text)
        path = tmp_path / "section.docx"
        doc.save(path)
        result = classify_document(analyze_docx(path))
        assert [e.selected_decision.role for e in result.elements] == expected


@pytest.mark.parametrize("case", CASES, ids=lambda c: c["name"])
def test_written_docx_roles_outline_and_repeat_conversion(case, tmp_path):
    source = make_docx(case, tmp_path / "source.docx")
    pack = create_style_pack(name="封面回归", format_id="docx")
    for name in ("first", "second"):
        output = tmp_path / (name + ".docx")
        normalize_docx_direct(source, pack, output, audit_dir=tmp_path / name)
        doc = Document(output)
        resolver = EffectiveFormatResolver(
            parse_xml_payload(read_zip_member(output, "word/styles.xml"))
        )
        assert [p.text for p in doc.paragraphs] == [t for t, _, _ in case["paragraphs"]]
        for paragraph, (_, _, role) in zip(doc.paragraphs, case["paragraphs"], strict=True):
            assert paragraph.style.style_id == ROLE_TO_STYLE_ID[role]
            outline = resolver.resolve_paragraph(paragraph._p, paragraph_style_id=None).get(
                "outline_level"
            )
            assert outline == 0 if role == "heading-1" else outline in {None, 9}
        source = output


def test_cover_reclassification_preserves_complex_body_ooxml(tmp_path):
    case = next(c for c in CASES if c["name"] == "cover-with-metadata")
    source = make_docx(case, tmp_path / "source.docx")
    doc = Document(source)
    p = doc.add_paragraph("复杂正文与字段保持完整。")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "42")
    bookmark.set(qn("w:name"), "TitleRegression")
    p._p.append(bookmark)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " DATE ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "2026"
    run.append(text)
    field.append(run)
    p._p.append(field)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "42")
    p._p.append(end)
    doc.save(source)
    output = tmp_path / "output.docx"
    normalize_docx_direct(
        source,
        create_style_pack(name="保护", format_id="docx"),
        output,
        audit_dir=tmp_path / "audit",
    )
    before = parse_xml_payload(read_zip_member(source, "word/document.xml"))
    after = parse_xml_payload(read_zip_member(output, "word/document.xml"))
    for tag in ("bookmarkStart", "bookmarkEnd", "fldSimple"):
        assert [etree.tostring(n) for n in before.iter(qn("w:" + tag))] == [
            etree.tostring(n) for n in after.iter(qn("w:" + tag))
        ]
    assert read_zip_member(source, "word/_rels/document.xml.rels") == read_zip_member(
        output, "word/_rels/document.xml.rels"
    )


def test_explicit_metadata_and_no_blank_front_matter():
    source = "---\ntitle: 已确认的主题\nsubtitle: 已确认的观察视角\n---\n# 总体要求\n正文内容完整。\n# 工作安排\n安排工作。"
    one = normalize_markdown(source, filename="文前.md")
    assert one.metadata["title"] == "已确认的主题"
    assert one.metadata["subtitle"] == "已确认的观察视角"
    assert [d["role"] for d in one.report["decisions"] if d["role"].startswith("heading-")] == [
        "heading-1",
        "heading-1",
    ]
    assert normalize_markdown(one.markdown, filename="文前.md").markdown == one.markdown
    inferred = normalize_markdown(
        "# 从单点响应走向综合服务\n以城乡服务网络为研究对象\n# 研究背景\n记录服务需求。",
        filename="文前.md",
    )
    assert inferred.metadata["title"] == "从单点响应走向综合服务"
    assert inferred.metadata["subtitle"] == "以城乡服务网络为研究对象"


@pytest.mark.parametrize("route", ["markdown", "docx"])
def test_abstract_boundary_supports_cover_without_consuming_real_body(route, tmp_path):
    case = {
        "paragraphs": [
            ("从分散储备走向跨仓协同", "Heading1", "document-title"),
            ("以区域保障网络为观察窗口", "Heading1", "document-subtitle"),
            ("摘要：本文记录区域保障网络的协作过程。", "Normal", "abstract-label"),
            ("研究背景", "Heading1", "heading-1"),
            ("记录网络运行条件。", "Normal", "body-paragraph"),
            ("研究方法", "Heading1", "heading-1"),
            ("开展现场数据采集。", "Normal", "body-paragraph"),
        ]
    }
    if route == "markdown":
        result = normalize_markdown(markdown_source(case), filename="论文.md")
        actual = markdown_roles(result, case)
    else:
        result = classify_document(analyze_docx(make_docx(case, tmp_path / "paper.docx")))
        actual = [e.selected_decision.role for e in result.elements]
    assert actual == [r for _, _, r in case["paragraphs"]]
