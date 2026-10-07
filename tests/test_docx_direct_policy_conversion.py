"""Full-source DOCX acceptance: policies may change format, never structure."""

from __future__ import annotations

import json
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.style_library import create_style_pack

PRESSURE = Path(__file__).parent / "fixtures/DOCX转Markdown压力测试_V1_结构兼容性.docx"


def node(tag: str, **attrs: str):
    result = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        result.set(qn("w:" + key), value)
    return result


def paragraph_text(p) -> str:
    return "".join(n.text or "" for n in p.iter() if n.tag in {qn("w:t"), qn("w:delText")})


def paragraph(root, marker: str):
    return next(p for p in root.iter(qn("w:p")) if marker in paragraph_text(p))


def package(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as z:
        assert z.testzip() is None
        return {n: z.read(n) for n in z.namelist()}


def make_policy_fixture(path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    for text, size in (
        ("转换策略验收报告", 32),
        ("段落与字符策略", 24),
        ("作者：张三", 12),
        ("单位：示例公司", 12),
        ("2026年10月3日", 12),
    ):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(size)
    base = doc.styles.add_style("BadCharBase", WD_STYLE_TYPE.CHARACTER)
    base.font.name = "Courier New"
    base.font.size = Pt(26)
    emphasis = doc.styles.add_style("LocalSemantic", WD_STYLE_TYPE.CHARACTER)
    emphasis.base_style = base
    emphasis.font.bold = True
    for level in range(1, 7):
        p = doc.add_heading(f"策略层级{level}", level=level)
        p.runs[0].font.name = "Arial"
        p.runs[0].font.size = Pt(29)
        p.runs[0].font.color.rgb = RGBColor.from_string("FF0000")
        p.runs[0].bold = False
        p.add_run("局部源强调").bold = True
        p.add_run("字符样式").style = emphasis
        p.paragraph_format.left_indent = Pt(40)
        p.paragraph_format.line_spacing = Pt(45)
        p.paragraph_format.keep_with_next = False
    p = doc.add_paragraph("混合正文包含")
    p.add_run("重点").bold = True
    p.add_run("斜体").italic = True
    p.add_run("字符强调").style = emphasis
    p.add_run("上标").font.superscript = True
    p.add_run("下标").font.subscript = True
    p.add_run("彩色").font.color.rgb = RGBColor.from_string("009900")
    p.add_run("下划线").underline = True
    p.add_run("背景")._r.get_or_add_rPr().append(node("highlight", val="yellow"))
    p.add_run("，仍然是一段正常正文。")
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(21)
    p.paragraph_format.left_indent = Pt(34)
    p.paragraph_format.line_spacing = Pt(40)
    p = doc.add_paragraph("全段视觉覆盖应清除并使用正文样式。")
    for tag, val in (
        ("b", "1"),
        ("i", "1"),
        ("color", "FF0000"),
        ("u", "single"),
        ("highlight", "yellow"),
    ):
        p.runs[0]._r.get_or_add_rPr().append(node(tag, val=val))
    p = doc.add_paragraph("复杂正文仍然有正常文字")
    relationship_id = doc.part.relate_to(
        "https://example.invalid/reference",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = node("hyperlink")
    link.set(qn("r:id"), relationship_id)
    run = node("r")
    props = node("rPr")
    props.append(node("rStyle", val="LocalSemantic"))
    run.append(props)
    text = node("t")
    text.text = "有效链接"
    run.append(text)
    link.append(run)
    p._p.append(link)
    p._p.append(node("bookmarkStart", id="33", name="PolicyBookmark"))
    p._p.append(node("bookmarkEnd", id="33"))
    p.add_run().add_tab()
    p.add_run().add_break()
    field = node("fldSimple", instr=" DATE ")
    field_run = node("r")
    field_run.append(node("rPr"))
    field_text = node("t")
    field_text.text = "2026"
    field_run.append(field_text)
    field.append(field_run)
    p._p.append(field)
    for kind, value in (
        ("begin", ""),
        (None, " DATE "),
        ("separate", ""),
        (None, "2026"),
        ("end", ""),
    ):
        r = p.add_run()._r
        if kind:
            r.append(node("fldChar", fldCharType=kind))
        else:
            n = node("instrText" if value.strip() == "DATE" else "t")
            n.text = value
            r.append(n)
    pressure = package(PRESSURE)
    pressure_xml = parse_xml(pressure["word/document.xml"])
    for tag in ("ins", "del"):
        original = next(iter(pressure_xml.iter(qn("w:" + tag))), None)
        if original is not None:
            p._p.append(deepcopy(original))
    p.add_run("，这些结构均应保留。")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "水平合并单元格"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵向合并单元格"
    table.cell(2, 2).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套表格正文"
    p = table.cell(1, 1).paragraphs[0]
    p.text = "表内编号段落"
    p._p.get_or_add_pPr().append(node("numPr"))  # structure present without active numId
    p._p.get_or_add_pPr().append(node("ind", left="400", hanging="200"))
    for p in (cell.paragraphs[0] for row in table.rows for cell in row.cells):
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(25)
    p = doc.add_paragraph("列表内容保持原样", style="List Number")
    p.runs[0].font.size = Pt(27)
    custom = doc.styles.add_style("UnknownPolicy", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("未知样式保留", style=custom).runs[0].font.size = Pt(9)
    doc.sections[0].header.paragraphs[0].text = "页眉必须保留"
    doc.sections[0].footer.paragraphs[0].text = "页脚必须保留"
    image = next(
        v for k, v in pressure.items() if k.startswith("word/media/") and k.endswith(".png")
    )
    doc.add_paragraph("图形正文包含文字及图片。").add_run().add_picture(BytesIO(image))
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 参考条目的基础字体也应统一。").runs[0].font.size = Pt(22)
    doc.save(path)
    # Add valid note parts/relationships from the real stress package; no rebuild.
    parts = package(path)
    root = parse_xml(parts["word/document.xml"])
    complex_p = paragraph(root, "复杂正文")
    relations = etree.fromstring(parts["word/_rels/document.xml.rels"])
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    for kind in ("footnotes", "endnotes"):
        member = f"word/{kind}.xml"
        if member not in pressure:
            continue
        parts[member] = pressure[member]
        note_root = parse_xml(parts[member])
        singular = kind[:-1]
        note_id = next(n.get(qn("w:id")) for n in note_root if int(n.get(qn("w:id"), "-1")) > 0)
        ref_run = node("r")
        ref_run.append(node(singular + "Reference", id=note_id))
        complex_p.append(ref_run)
        etree.SubElement(
            relations,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="rIdPolicy" + kind,
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + kind,
            Target=kind + ".xml",
        )
        etree.SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/" + member,
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml."
            + kind
            + "+xml",
        )
    for member, xml in (
        ("word/document.xml", root),
        ("word/_rels/document.xml.rels", relations),
        ("[Content_Types].xml", content_types),
    ):
        parts[member] = etree.tostring(xml, xml_declaration=True, encoding="UTF-8")
    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for name, value in parts.items():
            z.writestr(name, value)


def assert_package_structures(source: Path, output: Path) -> None:
    before, after = package(source), package(output)
    assert before.keys() == after.keys()
    for name in before:
        if name not in {"word/document.xml", "word/styles.xml"}:
            assert before[name] == after[name], name
    original_styles = parse_xml(before["word/styles.xml"])
    converted_styles = parse_xml(after["word/styles.xml"])
    style_lookup = {s.get(qn("w:styleId")): s for s in converted_styles.findall(qn("w:style"))}
    for style in original_styles.findall(qn("w:style")):
        assert etree.tostring(style, method="c14n", exclusive=True) == etree.tostring(
            style_lookup[style.get(qn("w:styleId"))], method="c14n", exclusive=True
        )
    roots = [parse_xml(p["word/document.xml"]) for p in (before, after)]
    for tag in (
        "numPr",
        "bookmarkStart",
        "bookmarkEnd",
        "fldSimple",
        "fldChar",
        "instrText",
        "footnoteReference",
        "endnoteReference",
        "drawing",
        "br",
        "tab",
        "ins",
        "del",
        "sectPr",
    ):
        assert [etree.tostring(n, method="c14n") for n in roots[0].iter(qn("w:" + tag))] == [
            etree.tostring(n, method="c14n") for n in roots[1].iter(qn("w:" + tag))
        ], tag
    for tag in ("gridSpan", "vMerge", "tblPr", "tblGrid", "tcPr", "trPr"):
        assert [etree.tostring(n, method="c14n") for n in roots[0].iter(qn("w:" + tag))] == [
            etree.tostring(n, method="c14n") for n in roots[1].iter(qn("w:" + tag))
        ], tag
    assert [paragraph_text(p) for p in roots[0].iter(qn("w:p"))] == [
        paragraph_text(p) for p in roots[1].iter(qn("w:p"))
    ]
    for marker in ("列表内容保持原样", "未知样式保留"):
        assert etree.tostring(paragraph(roots[0], marker), method="c14n") == etree.tostring(
            paragraph(roots[1], marker), method="c14n"
        )


def verify_policy_conversion(folder: Path, decision_mode="hybrid") -> dict[str, int | bool]:
    folder.mkdir(parents=True, exist_ok=True)
    source = folder / "source.docx"
    make_policy_fixture(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = folder / "normalized.docx"
    pack = create_style_pack(name="策略验收", format_id="docx")
    pack.roles["body-paragraph"]["font"]["size_pt"] = 14
    pack.roles["body-table"]["font"]["size_pt"] = 13
    for level in range(1, 7):
        pack.roles[f"heading-{level}"]["font"]["bold"] = True
    result = normalize_docx_direct(source, pack, output, decision_mode=decision_mode)
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert_package_structures(source, output)
    normalized = Document(output)
    heading = next(p for p in normalized.paragraphs if p.text.startswith("策略层级1"))
    for r in heading.runs:
        assert r._r.find("w:rPr/w:rFonts", heading._p.nsmap) is None
        assert r._r.find("w:rPr/w:sz", heading._p.nsmap) is None
        assert r._r.find("w:rPr/w:b", heading._p.nsmap) is None
    from doc_converter.docx_direct.analyzer import analyze_docx

    actual = analyze_docx(output)
    h = next(e for e in actual.elements if e.text.startswith("策略层级1"))
    assert h.effective_format["dominant_bold"] is True
    for element in actual.elements:
        if element.text.startswith("策略层级"):
            level = int(element.text[4])
            assert element.effective_format["dominant_bold"] is True
            assert all(
                r["font_size_pt"] == pack.roles[f"heading-{level}"]["font"]["size_pt"]
                for r in element.effective_format["run_statistics"]["runs"]
            )
        if element.table_context and element.text:
            assert all(
                r["font_size_pt"] == 13 for r in element.effective_format["run_statistics"]["runs"]
            )
    assert all(
        r["font_size_pt"] == 14
        for r in next(e for e in actual.elements if e.text.startswith("混合正文")).effective_format[
            "run_statistics"
        ]["runs"]
    )
    body = next(p for p in normalized.paragraphs if p.text.startswith("混合正文"))
    for token, prop in (
        ("重点", "b"),
        ("斜体", "i"),
        ("上标", "vertAlign"),
        ("下标", "vertAlign"),
        ("彩色", "color"),
        ("下划线", "u"),
        ("背景", "highlight"),
    ):
        assert (
            next(r for r in body.runs if r.text == token)._r.find("w:rPr/w:" + prop, body._p.nsmap)
            is not None
        )
    uniform = next(p for p in normalized.paragraphs if p.text.startswith("全段视觉"))
    assert all(
        uniform._p.find(".//w:" + tag, uniform._p.nsmap) is None
        for tag in ("b", "i", "color", "u", "highlight")
    )
    records = json.loads(result.normalization_report_path.read_text(encoding="utf-8"))["elements"]
    assert all(r["policy_name"] and r["policy_family"] for r in records)
    assert (
        next(r for r in records if r["text"].startswith("策略层级1"))["policy_family"]
        == "heading-strong"
    )
    assert (
        next(r for r in records if r["text"].startswith("复杂正文"))["transform_safety"]
        == "preserve"
    )
    if decision_mode in {"hybrid", "shadow"}:
        roots = [parse_xml(package(p)["word/document.xml"]) for p in (source, output)]
        assert etree.tostring(paragraph(roots[0], "复杂正文"), method="c14n") == etree.tostring(
            paragraph(roots[1], "复杂正文"), method="c14n",
        )
    assert (
        next(r for r in records if r["text"] == "列表内容保持原样")["policy_family"] == "preserve"
    )
    assert all(r["scored_decision"] and r["legacy_decision"] for r in records)
    assert any(r["removed_properties"] for r in records)
    heading_record = next(r for r in records if r["text"].startswith("策略层级1"))
    assert {"rFonts", "sz", "b"} <= {
        p["property"]
        for p in heading_record["removed_properties"]
        if p["level"] == "character-style"
    }
    summary = {
        "passed": True,
        "paragraphs": len(records),
        "source_unchanged": True,
        "structures_preserved": True,
    }
    (folder / "verification.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


@pytest.mark.parametrize("mode", ["shadow", "hybrid"])
def test_full_docx_policy_conversion(tmp_path: Path, mode) -> None:
    verify_policy_conversion(tmp_path, mode)


def test_inherited_numbering_and_page_break_survive_style_mapping(tmp_path: Path) -> None:
    doc = Document()
    source_style = doc.styles["Heading 1"]._element.get_or_add_pPr()
    numbering = node("numPr")
    numbering.append(node("ilvl", val="0"))
    numbering.append(node("numId", val="5"))
    source_style.append(numbering)
    source_style.append(node("ind", left="900", hanging="300"))
    source_style.append(node("pageBreakBefore", val="1"))
    doc.add_heading("继承编号标题", level=1)
    source = tmp_path / "inherited.docx"
    output = tmp_path / "normalized.docx"
    doc.save(source)
    normalize_docx_direct(source, create_style_pack(name="继承验收", format_id="docx"), output)
    converted = Document(output).paragraphs[0]
    assert converted.style.style_id == "DC_Heading1"
    properties = converted._p.find(qn("w:pPr"))
    for tag in ("numPr", "ind"):
        preserved = properties.find(qn("w:" + tag))
        assert preserved is not None, tag
        assert etree.tostring(preserved, method="c14n", exclusive=True) == etree.tostring(
            source_style.find(qn("w:" + tag)), method="c14n", exclusive=True
        ), tag
    assert package(source)["word/numbering.xml"] == package(output)["word/numbering.xml"]


def test_complex_field_results_across_paragraphs_are_untouched(tmp_path: Path) -> None:
    doc = Document()
    start = doc.add_paragraph("域开始之前的普通正文。")
    start.add_run()._r.append(node("fldChar", fldCharType="begin"))
    instruction = node("instrText")
    instruction.text = ' TOC \\o "1-3" '
    start.add_run()._r.append(instruction)
    start.add_run()._r.append(node("fldChar", fldCharType="separate"))
    result = doc.add_paragraph("跨段落的目录域结果应该保留原样。")
    result.runs[0].font.size = Pt(26)
    result.runs[0].font.name = "Arial"
    expected = etree.tostring(result.runs[0]._r, method="c14n", exclusive=True)
    doc.add_paragraph().add_run()._r.append(node("fldChar", fldCharType="end"))
    doc.add_paragraph("域结束之后是正常业务正文。")
    source, output = tmp_path / "fields.docx", tmp_path / "normalized.docx"
    doc.save(source)
    normalize_docx_direct(source, create_style_pack(name="域保护验收", format_id="docx"), output)
    converted = next(p for p in Document(output).paragraphs if p.text.startswith("跨段落"))
    assert etree.tostring(converted.runs[0]._r, method="c14n", exclusive=True) == expected
