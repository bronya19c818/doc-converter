"""Rebuild small synthetic semantic fixtures without real document/binary content."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from lxml import etree

HERE = Path(__file__).resolve().parent


def main():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    p = doc.add_paragraph("基于库容指数的物流仓储风险研究")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(24)
    p = doc.add_paragraph("林硕，张圣洁", style="Heading 3")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(14)
    p.runs[0].bold = False
    for text in ("某研究单位", "2026年10月"):
        doc.add_paragraph(text).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("一、研究背景", level=1)
    doc.add_paragraph("本段说明研究目的、基本措施和验证方式，不包含真实项目数据。")
    doc.save(HERE / "author_list_dirty_heading.docx")

    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading("一、正文", level=1)
    doc.add_paragraph("1. 基本范围")
    doc.add_paragraph("本段是完整正文，作为正文格式的计算基线。")
    for text, kind in (("（1）", "ole"), ("（2）", "ole"),
                       ("（3）", "math"), ("（4）", "pict")):
        p = doc.add_paragraph(text, style="Heading 6" if text == "（1）" else "Normal")
        if kind == "ole":
            obj = OxmlElement("w:object")
            obj.append(etree.Element("{urn:schemas-microsoft-com:office:office}OLEObject",
                                     ProgID="Equation.DSMT4", Type="Embed"))
            p.add_run()._r.append(obj)
        elif kind == "math":
            math_para = OxmlElement("m:oMathPara")
            math = OxmlElement("m:oMath")
            run = OxmlElement("m:r")
            token = OxmlElement("m:t")
            token.text = "x=1"
            run.append(token)
            math.append(run)
            math_para.append(math)
            p._p.append(math_para)
        else:
            pict = OxmlElement("w:pict")
            pict.append(etree.Element("{urn:schemas-microsoft-com:vml}shape", id="fixture-shape"))
            p.add_run()._r.append(pict)
    doc.add_paragraph("（5）正常三级标题")
    doc.add_paragraph("本段继续说明业务事实，公式编号不能影响本段前的标题层级。")
    doc.save(HERE / "complex_numbering.docx")


if __name__ == "__main__":
    main()
