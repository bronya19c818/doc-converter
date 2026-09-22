from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

import doc_converter.importers.libreoffice as libreoffice
from doc_converter.importers.base import ImporterError
from doc_converter.importers.docx import import_docx


def test_docx_import_promotes_high_confidence_centered_title_and_subtitle(tmp_path) -> None:
    document = Document()
    title = document.add_paragraph("DOCX → Markdown 压力测试文档")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(20)
    subtitle = document.add_paragraph("V1 · 结构兼容性 / 边界场景 / 可定位验收")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    document.add_paragraph()
    usage = document.add_paragraph("用途：验证结构保真度。")
    usage.alignment = WD_ALIGN_PARAGRAPH.CENTER
    usage.runs[0].font.size = Pt(11)
    document.add_heading("CASE-TOC-001 目录域", level=1)
    source_path = tmp_path / "压力测试.docx"
    document.save(source_path)

    result = import_docx(source_path, job_directory=tmp_path)

    assert result.markdown.startswith(
        "---\ntitle: DOCX → Markdown 压力测试文档\nsubtitle: V1 · 结构兼容性 / 边界场景 / 可定位验收\n---"
    )
    assert result.markdown.count("DOCX → Markdown 压力测试文档") == 1
    assert "用途：验证结构保真度。" in result.markdown
    assert "\n\n用途：验证结构保真度。\n\n# CASE-TOC-001" in result.markdown


def test_docx_import_marks_tabs_and_page_breaks_for_roundtrip(tmp_path) -> None:
    document = Document()
    paragraph = document.add_paragraph("A")
    paragraph.add_run().add_tab()
    paragraph.add_run("B")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("After")
    source_path = tmp_path / "扩展.docx"
    document.save(source_path)

    result = import_docx(source_path, job_directory=tmp_path)

    assert "A&#9;B" in result.markdown
    assert "<!-- doc-converter:page-break -->" in result.markdown


def test_docx_import_promotes_title_without_subtitle_and_splits_soft_line_preface(tmp_path) -> None:
    document = Document()
    title = document.add_paragraph("只有标题的文档")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(20)
    preface = document.add_paragraph()
    preface.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = preface.add_run("V2 · 结构说明")
    subtitle.font.size = Pt(14)
    subtitle.add_break()
    usage = preface.add_run("用途说明")
    usage.font.size = Pt(11)
    document.add_heading("一、正文", level=1)
    source_path = tmp_path / "软换行.docx"
    document.save(source_path)

    result = import_docx(source_path, job_directory=tmp_path)

    assert "title: 只有标题的文档" in result.markdown
    assert "subtitle: V2 · 结构说明" in result.markdown
    assert "用途说明" in result.markdown
    assert "subtitle: 用途说明" not in result.markdown


def test_docx_import_keeps_title_when_no_subtitle_exists(tmp_path) -> None:
    document = Document()
    title = document.add_paragraph("只有一个标题")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(20)
    document.add_heading("一、正文", level=1)
    source_path = tmp_path / "无副标题.docx"
    document.save(source_path)

    result = import_docx(source_path, job_directory=tmp_path)

    assert "title: 只有一个标题" in result.markdown
    assert "subtitle:" not in result.markdown


def test_docx_import_skips_unlabeled_organization_preface_when_promoting_title(tmp_path) -> None:
    document = Document()
    organization = document.add_paragraph("采购部")
    organization.alignment = WD_ALIGN_PARAGRAPH.CENTER
    organization.runs[0].font.size = Pt(28)
    title = document.add_paragraph("正式报告标题")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(20)
    document.add_heading("一、正文", level=1)
    source_path = tmp_path / "文前单位.docx"
    document.save(source_path)

    result = import_docx(source_path, job_directory=tmp_path)

    assert "title: 正式报告标题" in result.markdown
    assert "title: 采购部" not in result.markdown
    assert "采购部" in result.markdown


def test_docx_import_rejects_invalid_zip_as_importer_error(tmp_path) -> None:
    source_path = tmp_path / "坏文件.docx"
    source_path.write_bytes(b"not a zip")

    try:
        import_docx(source_path, job_directory=tmp_path)
    except ImporterError as error:
        assert "有效的 Office Open XML" in str(error)
    else:
        raise AssertionError("invalid DOCX should raise ImporterError")


def test_resolve_libreoffice_prefers_console_launcher_before_gui_executable(monkeypatch) -> None:
    monkeypatch.delenv("DOC_CONVERTER_LIBREOFFICE", raising=False)
    calls: list[str] = []

    def fake_which(name: str) -> str | None:
        calls.append(name)
        if name == "soffice.com":
            return r"C:\Program Files\LibreOffice\program\soffice.com"
        if name == "soffice.exe":
            return r"C:\Program Files\LibreOffice\program\soffice.exe"
        return None

    monkeypatch.setattr(libreoffice.shutil, "which", fake_which)

    resolved = libreoffice.resolve_libreoffice()

    assert resolved == Path(r"C:\Program Files\LibreOffice\program\soffice.com")
    assert calls == ["soffice.com"]
