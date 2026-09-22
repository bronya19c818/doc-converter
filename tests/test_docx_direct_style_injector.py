from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from doc_converter.docx_direct.ooxml_patch import read_zip_member
from doc_converter.docx_direct.style_injector import inject_styles_into_docx
from doc_converter.style_library import create_style_pack

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def make_source_docx(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Existing normal paragraph")
    document.add_heading("Existing Heading", level=1)
    document.styles.add_style("LegacyCustom", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Existing custom paragraph", style="LegacyCustom")
    source = tmp_path / "source.docx"
    document.save(source)
    return source


def read_style_ids(path: Path) -> set[str]:
    root = ElementTree.fromstring(read_zip_member(path, "word/styles.xml"))
    return {
        style.attrib[f"{W_NS}styleId"]
        for style in root.findall(f"{W_NS}style")
        if f"{W_NS}styleId" in style.attrib
    }


def package_payloads(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_style_injection_adds_dc_styles_without_replacing_existing_styles(tmp_path: Path) -> None:
    source = make_source_docx(tmp_path)
    output = tmp_path / "injected.docx"
    pack = create_style_pack(name="test", format_id="docx")

    original_ids = read_style_ids(source)
    report = inject_styles_into_docx(source, pack, output)
    output_ids = read_style_ids(output)

    assert original_ids <= output_ids
    assert {"DC_DocumentTitle", "DC_Heading1", "DC_BodyParagraph"} <= output_ids
    assert report["inserted_count"] > 0
    assert report["replaced_count"] == 0


def test_style_injection_is_idempotent_and_preserves_non_style_payloads(tmp_path: Path) -> None:
    source = make_source_docx(tmp_path)
    first_output = tmp_path / "first.docx"
    second_output = tmp_path / "second.docx"
    pack = create_style_pack(name="test", format_id="docx")

    first_report = inject_styles_into_docx(source, pack, first_output)
    second_report = inject_styles_into_docx(first_output, pack, second_output)

    assert first_report["inserted_count"] > 0
    assert second_report["inserted_count"] == 0
    assert read_style_ids(first_output) == read_style_ids(second_output)

    source_payloads = package_payloads(source)
    second_payloads = package_payloads(second_output)
    for name, payload in source_payloads.items():
        if name != "word/styles.xml":
            assert second_payloads[name] == payload


def test_read_zip_member_reports_missing_member(tmp_path: Path) -> None:
    source = make_source_docx(tmp_path)

    try:
        read_zip_member(source, "word/not-found.xml")
    except KeyError as error:
        assert "word/not-found.xml" in str(error)
    else:
        raise AssertionError("missing ZIP members must raise KeyError")
