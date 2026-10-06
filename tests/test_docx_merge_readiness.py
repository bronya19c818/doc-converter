"""DOCX upload/export and risk CLI checks, specified before implementation.

Failure modes: missing surnames, partial-list overconfidence, business-word false
authors, context-free mapping, false HIGH alerts, repeated table floods, changing
totals after filtering, non-deterministic sampling and unreadable review output.
"""
import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from scripts.verify_docx_selected import verify_sample

ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "tests/fixtures/docx_direct_regression/author_list_dirty_heading.docx"


@pytest.mark.parametrize("text,matches,role,action", [
    ("刘伟，黄强", 2, "author", "mapped"),
    ("郭峰、梁晨", 2, "author", "mapped"),
    ("高伟 罗峰", 2, "author", "mapped"),
    ("欧阳明，司马晨", 2, "author", "mapped"),
    ("某某，张圣洁", 1, "author", "preserved"),
    ("采购 仓储", 0, "unknown", "preserved"),
    ("计划 合同", 0, "unknown", "preserved"),
    ("库存 调配", 0, "unknown", "preserved"),
    ("设备 材料", 0, "unknown", "preserved"),
])
def test_generalized_author_evidence_through_api(text, matches, role, action, tmp_path, monkeypatch):
    monkeypatch.setenv("DOC_CONVERTER_LIBREOFFICE", str(tmp_path / "no-preview"))
    doc = Document(FIXTURE)
    doc.paragraphs[1].text = text
    source = tmp_path / "source.docx"
    doc.save(source)
    verify_sample(source, tmp_path / "api", [])
    analysis = json.loads((tmp_path / "api/docx-analysis.json").read_text(encoding="utf-8"))
    e = next(e for e in analysis["elements"] if e["text"] == text)
    f = e["features"]
    assert f["author_name_list_shape"]
    assert f["author_name_count"] == 2
    assert f["author_surname_match_count"] == matches
    assert f["author_surname_match_ratio"] == matches / 2
    assert f["author_name_list_pattern"] == (matches == 2)
    assert (e["selected_decision"]["role"], e["selected_decision"]["action"]) == (role, action)
    if matches == 1:
        assert e["scored_decision"]["confidence_band"] == "medium"
        assert e["scored_decision"]["action"] == "uncertain"


@pytest.mark.parametrize("text", ["欧阳明，司马晨", "某某，张圣洁", "采购 仓储"])
def test_name_shape_without_context_preserved_through_api(text, tmp_path, monkeypatch):
    monkeypatch.setenv("DOC_CONVERTER_LIBREOFFICE", str(tmp_path / "no-preview"))
    doc = Document(FIXTURE)
    doc.paragraphs[1].text = text
    for p in list(doc.paragraphs)[2:4]:
        p._p.getparent().remove(p._p)
    source = tmp_path / "source.docx"
    doc.save(source)
    verify_sample(source, tmp_path / "api", [])
    analysis = json.loads((tmp_path / "api/docx-analysis.json").read_text(encoding="utf-8"))
    e = next(e for e in analysis["elements"] if e["text"] == text)
    assert (e["selected_decision"]["role"], e["selected_decision"]["action"]) == ("unknown", "preserved")


def _review(folder, name, *args):
    output = folder / (name + ".json")
    markdown = folder / (name + ".md")
    subprocess.run([
        sys.executable, str(ROOT / "scripts/review_docx_decision_risks.py"),
        "--analysis", str(folder / "docx-analysis.json"),
        "--report", str(folder / "docx-normalization-report.json"),
        "--output", str(output), "--markdown", str(markdown), *args,
    ], check=True)
    return json.loads(output.read_text(encoding="utf-8")), markdown.read_text(encoding="utf-8")


def test_risk_tiers_sampling_filters_and_markdown_through_api_cli(tmp_path, monkeypatch):
    monkeypatch.setenv("DOC_CONVERTER_LIBREOFFICE", str(tmp_path / "no-preview"))
    doc = Document(FIXTURE)
    table = doc.add_table(rows=120, cols=2)
    for i, row in enumerate(table.rows):
        row.cells[0].text = f"物料{i}"
        row.cells[1].text = "数量"
    source = tmp_path / "source.docx"
    doc.save(source)
    folder = tmp_path / "api"
    verify_sample(source, folder, [])
    default, md = _review(folder, "default")
    repeated, _ = _review(folder, "repeat")
    full, _ = _review(folder, "full", "--full")
    only_p1, _ = _review(folder, "p1", "--level", "P1")
    small, _ = _review(folder, "small", "--max-per-flag", "2")
    assert default == repeated
    assert default["summary"]["high"] == 0  # repaired dirty Heading author is not HIGH
    assert default["summary"]["p1"] > 0
    assert default["summary"]["sample_total"] >= 200
    assert len(default["elements"]) < 40
    assert len(full["elements"]) == full["summary"]["flagged"]
    for key in ("total_analyzed", "high", "p1", "p2_total", "sample_total", "flagged"):
        assert default["summary"][key] == full["summary"][key] == only_p1["summary"][key]
    assert {e["risk_level"] for e in only_p1["elements"]} == {"P1"}
    assert len(only_p1["elements"]) == only_p1["summary"]["p1"]
    assert small["summary"]["sample_emitted"] <= 4
    assert len({e["element_id"] for e in full["elements"]}) == len(full["elements"])
    assert "legacy:" in md and "scored:" in md and "selected:" in md
    assert "numbering_family" in md and "selection_source" in md
    assert '"features"' not in md
    author = next(e for e in full["elements"] if e["text"] == "林硕，张圣洁")
    assert author["risk_level"] == "P1"
    assert "scored_takeover" in author["risk_flags"]
