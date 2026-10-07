"""Actual-package contracts, written before stage-two executor changes.

Failure modes: stale managed styles, source outline/page overrides, lost numbering
or inline structures, invented table headers, unsafe label splitting, growing
character projections, and audit success based only on a style identifier.
"""

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from doc_converter.docx_direct.analyzer import analyze_docx
from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.normalizer import normalize_docx_direct
from doc_converter.docx_styles import ROLE_TO_STYLE_ID
from doc_converter.style_library import create_style_pack


def xml(path, member="word/document.xml"):
    with ZipFile(path) as archive:
        return parse_xml(archive.read(member))


def pack(size=14):
    result = create_style_pack(name="转换契约", format_id="test")
    for role in result.roles.values():
        if "font" in role:
            role["font"].update(chinese_font="微软雅黑", western_font="Cambria", size_pt=size)
    result.roles["abstract-label"]["font"]["size_pt"] = size + 3
    result.roles["keywords-label"]["font"]["size_pt"] = size + 3
    return result


def convert(source, output, style=None, mode="hybrid"):
    return normalize_docx_direct(source, style or pack(), output, audit_dir=output.with_suffix(""), decision_mode=mode)


def prop(tag, value):
    node = OxmlElement("w:" + tag)
    node.set(qn("w:val"), str(value))
    return node


def element(result, text):
    return next(e for e in result.analysis.elements if text in e.text)


def test_visual_outline_and_managed_style_upsert_end_to_end(tmp_path: Path):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    title = doc.add_paragraph("角色转换测试报告")
    title.runs[0].font.size = Pt(32)
    author = doc.add_paragraph("作者：林硕，张圣洁", "Heading 3")
    author.runs[0].font.size = Pt(12)
    author.runs[0].bold = False
    for tag, value in (("outlineLvl", 2), ("pageBreakBefore", 1), ("widowControl", 0)):
        author._p.get_or_add_pPr().append(prop(tag, value))
    doc.add_heading("一、管理要求", 1)
    body = doc.add_paragraph("正文需要规范化，同时保留这里的局部强调。")
    body.add_run("重点").bold = True
    body.runs[0].font.size = Pt(28)
    for _ in range(4):
        doc.add_paragraph("这是用于确定正文基础字号的普通业务段落，正文内容连续完整，并且结束于句号。")
    source = tmp_path / "source.docx"
    doc.save(source)
    a = convert(source, tmp_path / "a.docx")
    a2 = convert(a.output_path, tmp_path / "a2.docx")
    b = convert(a2.output_path, tmp_path / "b.docx", pack(18))
    b2 = convert(b.output_path, tmp_path / "b2.docx", pack(18))
    for result, size in ((a, 14), (b, 18)):
        author = element(result, "作者")
        assert author.selected_decision.role == "author"
        actual = next(e for e in analyze_docx(result.output_path).elements if "作者" in e.text)
        assert actual.effective_format["outline_level"] == 9
        assert actual.effective_format["page_break_before"] is False
        assert actual.effective_format["widow_control"] is True
        assert actual.effective_format["dominant_font_size_pt"] == size
        record = next(r for r in result.records if "作者" in r.text)
        assert record.style_application_status == "applied"
        assert record.effective_format_match is True
    for first, second in ((a, a2), (b, b2)):
        assert etree.tostring(xml(first.output_path, "word/styles.xml"), method="c14n") == etree.tostring(xml(second.output_path, "word/styles.xml"), method="c14n")
    original = {s.get(qn("w:styleId")): etree.tostring(s, method="c14n") for s in xml(source, "word/styles.xml").findall(qn("w:style"))}
    for style in xml(b2.output_path, "word/styles.xml").findall(qn("w:style")):
        sid = style.get(qn("w:styleId"))
        if sid in original:
            assert etree.tostring(style, method="c14n") == original[sid]
    report = json.loads(b.normalization_report_path.read_text(encoding="utf-8"))
    definitions = {s.get(qn("w:styleId")): s for s in xml(b.output_path, "word/styles.xml").findall(qn("w:style"))}
    for sid in ("DC_Point", "DC_TableRowHeader"):
        assert definitions[sid].find("w:rPr/w:sz", definitions[sid].nsmap).get(qn("w:val")) == "36"
    assert report["role_capabilities"]["page"]["status"] == "out-of-scope"
    assert report["role_capabilities"]["table-row-header"]["status"] == "preserve-only"


@pytest.mark.parametrize("label", ["摘要：", "关键词:"])
@pytest.mark.parametrize("split", [False, True])
def test_composite_metadata_and_character_projection(tmp_path, label, split):
    doc = Document()
    base = doc.styles.add_style("SourceInline", WD_STYLE_TYPE.CHARACTER)
    base.font.size = Pt(30)
    base.font.bold = True
    p = doc.add_paragraph()
    if split:
        for char in label:
            p.add_run(char)
        p.add_run(" 普通内容 ")
    else:
        p.add_run(label + " 普通内容 ")
    p.add_run("重点").style = base
    source = tmp_path / "source.docx"
    doc.save(source)
    a = convert(source, tmp_path / "a.docx")
    a2 = convert(a.output_path, tmp_path / "a2.docx")
    b = convert(a2.output_path, tmp_path / "b.docx", pack(18))
    prefix = "abstract" if label.startswith("摘要") else "keywords"
    assert a.analysis.elements[0].selected_decision.role == prefix + "-label"
    assert [s.role for s in a.analysis.elements[0].inline_segments] == [prefix + "-label", prefix + "-content"]
    for result, size in ((a, 14), (b, 18)):
        actual = analyze_docx(result.output_path).elements[0]
        assert actual.text == label + " 普通内容 重点"
        runs = actual.effective_format["run_statistics"]["runs"]
        assert runs[0]["font_size_pt"] == size + 3
        assert all(r["font_size_pt"] == size for r in runs if r["text"].strip() in {"普通内容", "重点"})
        assert runs[-1]["bold"] is True
        assert result.records[0].effective_format_match is True
        for t in xml(result.output_path).iter(qn("w:t")):
            if t.text and (t.text.startswith(" ") or t.text.endswith(" ")):
                assert t.get(qn("xml:space")) == "preserve"
    def ids(path):
        return {s.get(qn("w:styleId")) for s in xml(path, "word/styles.xml").findall(qn("w:style"))}
    assert ids(a.output_path) == ids(a2.output_path) == ids(b.output_path)


def test_only_explicit_nearest_table_header_is_selected(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"单元格{i}{j}"
    table.rows[1]._tr.get_or_add_trPr().append(prop("tblHeader", 1))
    table.rows[2]._tr.get_or_add_trPr().append(prop("tblHeader", 0))
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套普通内容"
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx")
    for e in result.analysis.elements:
        if not e.text:
            continue
        expected = "table-column-header" if e.text.startswith("单元格1") else "body-table"
        assert e.selected_decision.role == expected
    for tag in ("tblPr", "tblGrid", "trPr", "tcPr"):
        before = [etree.tostring(n, method="c14n") for n in xml(source).iter(qn("w:" + tag))]
        after = [etree.tostring(n, method="c14n") for n in xml(result.output_path).iter(qn("w:" + tag))]
        assert before == after


def test_unsafe_inline_label_boundary_preserves_paragraph(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "摘要：内容"
    run.append(text)
    link.append(run)
    p._p.append(link)
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx")
    assert etree.tostring(next(xml(source).iter(qn("w:p"))), method="c14n") == etree.tostring(next(xml(result.output_path).iter(qn("w:p"))), method="c14n")
    assert result.records[0].style_application_status == "preserved"
    assert result.records[0].effective_format_match is None


def test_dirty_heading_peer_recovery_applies_actual_outline(tmp_path):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading("一、管理制度", 1)
    doc.add_heading("（一）管理职责", 2)
    for number, level in ((7, 3), (8, 3), (9, 6)):
        p = doc.add_heading(f"{number}. 主题职责", level)
        if number == 9:
            p._p.get_or_add_pPr().append(prop("outlineLvl", 5))
            p._p.get_or_add_pPr().append(prop("pageBreakBefore", 1))
            p.runs[0].font.size = Pt(30)
        doc.add_paragraph("这是职责的详细业务正文，业务正文需要由责任人按要求执行并完成记录。")
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx")
    selected = element(result, "9. 主题职责")
    assert selected.selected_decision.role == "heading-3"
    actual = next(e for e in analyze_docx(result.output_path).elements if e.text == selected.text)
    assert actual.original_style_id == "DC_Heading3"
    assert actual.effective_format["outline_level"] == 2
    assert actual.effective_format["page_break_before"] is False
    assert next(r for r in result.records if r.text == selected.text).effective_format_match is True


def test_composite_content_link_and_field_are_not_reconstructed(tmp_path):
    doc = Document()
    p = doc.add_paragraph("摘要：普通内容 ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "bookmark")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rpr.append(prop("color", "0000FF"))
    rpr.append(prop("u", "single"))
    rpr.append(prop("sz", 60))
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "链接内容"
    run.append(text)
    link.append(run)
    p._p.append(link)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    frun = OxmlElement("w:r")
    # A field with an explicit, self-contained font can remain unchanged while
    # surrounding ordinary content acquires the target role.
    fpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for slot, value in (("eastAsia", "微软雅黑"), ("ascii", "Cambria"), ("hAnsi", "Cambria"), ("cs", "Cambria")):
        fonts.set(qn("w:" + slot), value)
    fpr.append(fonts)
    fpr.append(prop("sz", 24))
    fpr.append(prop("szCs", 24))
    for tag in ("b", "bCs", "i", "iCs"):
        fpr.append(prop(tag, 0))
    frun.append(fpr)
    ftext = OxmlElement("w:t")
    ftext.text = "受保护域结果"
    frun.append(ftext)
    field.append(frun)
    p._p.append(field)
    expected = etree.tostring(field, method="c14n", exclusive=True)
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx")
    root = xml(result.output_path)
    assert etree.tostring(next(root.iter(qn("w:fldSimple"))), method="c14n", exclusive=True) == expected
    actual = analyze_docx(result.output_path).elements[0]
    link_format = next(r for r in actual.effective_format["run_statistics"]["runs"] if r["text"] == "链接内容")
    assert link_format["font_size_pt"] == 14
    assert link_format["color"] == "0000FF"
    assert next(root.iter(qn("w:hyperlink"))).attrib == link.attrib
    assert result.records[0].style_application_status == "guarded"
    assert result.records[0].effective_format_match is True
    assert any(c["excluded_reason"] for c in result.records[0].format_checks)


def test_protected_inherited_field_format_preserves_whole_paragraph(tmp_path):
    doc = Document()
    p = doc.add_paragraph("摘要：普通文字")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "继承源格式的域结果"
    run.append(text)
    field.append(run)
    p._p.append(field)
    source = tmp_path / "source.docx"
    doc.save(source)
    a = convert(source, tmp_path / "a.docx")
    assert a.records[0].style_application_status == "preserved"
    assert a.records[0].effective_format_match is None
    assert etree.tostring(next(xml(source).iter(qn("w:p"))), method="c14n") == etree.tostring(next(xml(a.output_path).iter(qn("w:p"))), method="c14n")


def test_default_style_true_inherited_numbering_is_kept(tmp_path):
    doc = Document()
    normal = doc.styles["Normal"]._element
    normal.set(qn("w:default"), "true")
    numbering = OxmlElement("w:numPr")
    numbering.append(prop("numId", 5))
    numbering.append(prop("ilvl", 0))
    normal.get_or_add_pPr().append(numbering)
    p = doc.add_paragraph("普通业务正文应当保留默认样式所继承的编号结构，不因应用样式而丢失编号。")
    p.runs[0].font.size = Pt(12)
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx", mode="shadow")
    assert result.analysis.elements[0].selected_decision.action == "mapped"
    output = next(xml(result.output_path).iter(qn("w:p")))
    actual = output.find("w:pPr/w:numPr", output.nsmap)
    assert actual is not None
    assert etree.tostring(actual, method="c14n", exclusive=True) == etree.tostring(numbering, method="c14n", exclusive=True)


def test_complex_script_emphasis_is_audited_per_run(tmp_path):
    doc = Document()
    p = doc.add_paragraph("这是完整普通业务正文，需要统一基础字体并保留局部的复杂文字强调。")
    p.add_run("强调")._r.get_or_add_rPr().append(prop("bCs", 1))
    for run in p.runs:
        run.font.size = Pt(12)
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx", mode="shadow")
    assert result.analysis.elements[0].selected_decision.action == "mapped"
    last = analyze_docx(result.output_path).elements[0].effective_format["run_statistics"]["runs"][-1]
    assert last["bold_cs"] is True
    checks = result.records[0].format_checks
    assert any(c["property"].endswith("bold_cs") and c["excluded_reason"] for c in checks)


def test_managed_style_used_by_protected_field_is_not_indirectly_rewritten(tmp_path):
    doc = Document()
    sid = ROLE_TO_STYLE_ID["abstract-label"]
    style = doc.styles.add_style(sid, WD_STYLE_TYPE.PARAGRAPH)
    style.font.size = Pt(10)
    p = doc.add_paragraph("摘要：普通文字", style)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "域结果"
    run.append(t)
    field.append(run)
    p._p.append(field)
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx")
    def definitions(path):
        return next(s for s in xml(path, "word/styles.xml").findall(qn("w:style")) if s.get(qn("w:styleId")) == sid)
    assert etree.tostring(definitions(source), method="c14n") == etree.tostring(definitions(result.output_path), method="c14n")
    report = json.loads(result.normalization_report_path.read_text(encoding="utf-8"))
    assert sid in report["style_injection"]["guarded_style_ids"]
    assert result.records[0].style_application_status == "preserved"


def test_content_character_toggle_resets_bold_label_baseline(tmp_path):
    doc = Document()
    doc.add_paragraph("摘要：普通内容应使用内容角色的非粗体字体。")
    source = tmp_path / "source.docx"
    doc.save(source)
    result = convert(source, tmp_path / "out.docx")
    p = next(xml(result.output_path).iter(qn("w:p")))
    actual = EffectiveFormatResolver(xml(result.output_path, "word/styles.xml"), structural_guard=True).resolve_paragraph(p, paragraph_style_id=None)
    assert actual["run_statistics"]["runs"][0]["bold"] is True
    assert actual["run_statistics"]["runs"][-1]["bold"] is False
    assert result.records[0].effective_format_after["run_statistics"]["runs"][-1]["bold"] is False


@pytest.mark.parametrize("bold", [True, False])
def test_heading_bold_overrides_document_default_toggle(tmp_path, bold):
    doc = Document()
    defaults = doc.styles._element.find("w:docDefaults/w:rPrDefault/w:rPr", doc.styles._element.nsmap)
    defaults.append(prop("b", 1))
    defaults.append(prop("bCs", 1))
    doc.add_heading("一、规范化标题", 1)
    source = tmp_path / "source.docx"
    doc.save(source)
    style = pack()
    style.roles["heading-1"]["font"]["bold"] = bold
    result = convert(source, tmp_path / "out.docx", style)
    p = next(xml(result.output_path).iter(qn("w:p")))
    actual = EffectiveFormatResolver(xml(result.output_path, "word/styles.xml"), structural_guard=True).resolve_paragraph(p, paragraph_style_id=None)
    assert actual["run_statistics"]["runs"][0]["bold"] is bold
    assert actual["run_statistics"]["runs"][0]["bold_cs"] is bold
    assert result.records[0].effective_format_match is True


def test_math_inherited_complex_font_blocks_managed_upsert(tmp_path):
    doc = Document()
    sid = ROLE_TO_STYLE_ID["abstract-label"]
    style = doc.styles.add_style(sid, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Cambria"
    style.font.size = Pt(14)
    style.font.bold = True
    style.font.italic = False
    props = style._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    fonts.set(qn("w:cs"), "Courier New")
    props.append(prop("szCs", 28))
    props.append(prop("bCs", 1))
    props.append(prop("iCs", 0))
    p = doc.add_paragraph("摘要：", style)
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x+1"
    run.append(text)
    math.append(run)
    p._p.append(math)
    source = tmp_path / "source.docx"
    doc.save(source)
    desired = pack()
    desired.roles["abstract-label"]["font"]["size_pt"] = 14
    result = convert(source, tmp_path / "out.docx", desired)
    report = json.loads(result.normalization_report_path.read_text(encoding="utf-8"))
    assert sid in report["style_injection"]["guarded_style_ids"]
    assert result.records[0].style_application_status == "preserved"
    source_style = next(s for s in xml(source, "word/styles.xml").findall(qn("w:style")) if s.get(qn("w:styleId")) == sid)
    output_style = next(s for s in xml(result.output_path, "word/styles.xml").findall(qn("w:style")) if s.get(qn("w:styleId")) == sid)
    assert etree.tostring(source_style, method="c14n") == etree.tostring(output_style, method="c14n")


def test_inline_character_inheritance_keeps_semantic_layers(tmp_path):
    doc = Document()
    base = doc.styles.add_style("SemanticBase", WD_STYLE_TYPE.CHARACTER)
    base.font.bold = True
    base.font.size = Pt(30)
    child = doc.styles.add_style("SemanticChild", WD_STYLE_TYPE.CHARACTER)
    child.base_style = base
    child.font.bold = True
    p = doc.add_paragraph("摘要：普通内容")
    p.add_run("双层字符语义").style = child
    source = tmp_path / "source.docx"
    doc.save(source)
    a = convert(source, tmp_path / "a.docx")
    a2 = convert(a.output_path, tmp_path / "a2.docx")
    run = list(xml(a.output_path).iter(qn("w:r")))[-1]
    sid = run.find("w:rPr/w:rStyle", run.nsmap).get(qn("w:val"))
    styles = {s.get(qn("w:styleId")): s for s in xml(a.output_path, "word/styles.xml").findall(qn("w:style"))}
    layers = []
    while sid != "DC_AbstractContent":
        style = styles[sid]
        layers.append(style.find("w:rPr/w:b", style.nsmap))
        sid = style.find(qn("w:basedOn")).get(qn("w:val"))
    assert len(layers) == 2 and all(n is not None for n in layers)
    assert etree.tostring(xml(a.output_path, "word/styles.xml"), method="c14n") == etree.tostring(xml(a2.output_path, "word/styles.xml"), method="c14n")
