"""User-specified DOCX/API end-to-end cases, written before local structure code.

Failure modes: adjacent numbered items become headings; keyword-based rules
damage true headings; a sequence swallows its containing/closing heading;
same-family levels leak across sections; decimal depth loses a native offset;
table or complex objects become parents; native anchors get overwritten.
"""
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

from scripts.verify_docx_selected import verify_sample


def _convert(tmp_path, monkeypatch, paragraphs):
    monkeypatch.setenv("DOC_CONVERTER_LIBREOFFICE", str(tmp_path / "no-preview"))
    document = Document()
    document.styles["Normal"].font.size = Pt(12)
    for text, native_level in paragraphs:
        if native_level:
            document.add_heading(text, level=native_level)
        else:
            p = document.add_paragraph(text)
            if native_level == 0:
                p._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    source = tmp_path / "source.docx"
    document.save(source)
    verify_sample(source, tmp_path / "api", [])
    analysis = json.loads((tmp_path / "api/docx-analysis.json").read_text(encoding="utf-8"))
    return {**{e["text"]: e for e in analysis["elements"]}, "_analysis": analysis}


@pytest.mark.parametrize("items", [
    ["1. 《A办法》", "2. 《B规定》", "3. 《C标准》", "4. 《D规范》"],
    ["1. 北区仓位", "2. 南区仓位", "3. 东区仓位", "4. 西区仓位"],
    ["（一）《A标准》", "（二）《B办法》", "（三）《C规定》", "（四）《D规范》"],
])
def test_parallel_manual_sequence_is_body_through_api(items, tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("管理依据", 1), *((text, None) for text in items),
        ("执行要求", 1), ("这里是执行要求的完整业务正文。", None),
    ])
    for text in items:
        e = actual[text]
        assert e["selected_decision"]["role"] == "body-paragraph"
        assert e["selected_decision"]["action"] == "mapped"
        assert e["features"]["manual_sequence_candidate"]
        assert e["features"]["manual_sequence_length"] >= 3
        assert not e["features"]["manual_sequence_has_child_content"]


def test_numbered_headings_with_child_body_remain_headings(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("1. 项目背景", None), ("这里是项目背景正文。", None),
        ("2. 建设目标", None), ("这里是建设目标正文。", None),
        ("3. 实施安排", None), ("这里是实施安排正文。", None),
    ])
    for text in ("1. 项目背景", "2. 建设目标", "3. 实施安排"):
        assert actual[text]["selected_decision"]["role"] == "heading-1"
        assert actual[text]["selected_decision"]["action"] == "mapped"
        assert not actual[text]["features"]["manual_sequence_candidate"]


@pytest.mark.parametrize("anchor", [None, 1])
def test_nested_manual_sequence_does_not_swallow_outer_headings(anchor, tmp_path, monkeypatch):
    items = ["1. 《A办法》", "2. 《B规定》", "3. 《C标准》"]
    actual = _convert(tmp_path, monkeypatch, [
        ("3. 管理要求", anchor), *((text, None) for text in items),
        ("4. 实施要求", anchor), ("这里是实施要求正文。", None),
    ])
    for text in items:
        assert actual[text]["selected_decision"]["role"] == "body-paragraph"
    for text in ("3. 管理要求", "4. 实施要求"):
        assert actual[text]["selected_decision"]["role"] == "heading-1"
        assert actual[text]["selected_decision"]["action"] == "mapped"


@pytest.mark.parametrize("paragraphs,expected", [
    ([
        ("一、总体要求", None), ("（一）总体原则", None), ("1. 工作要求", None), ("（1）具体要求", None),
        ("这里是工作要求的完整正文。", None),
    ], {"一、总体要求": 1, "（一）总体原则": 2, "1. 工作要求": 3, "（1）具体要求": 4}),
    ([
        ("第一章 总体要求", None), ("一、工作要求", None), ("（一）工作原则", None),
        ("这里是工作原则的完整正文。", None),
    ], {"第一章 总体要求": 1, "一、工作要求": 2, "（一）工作原则": 3}),
    ([
        ("第一章 总体情况", None), ("第一节 工作原则", None), ("一、具体要求", None),
        ("（一）服务内容", None), ("这里是服务内容的完整正文。", None),
    ], {"第一章 总体情况": 1, "第一节 工作原则": 2, "一、具体要求": 3, "（一）服务内容": 4}),
    ([
        ("一、总体要求", None), ("（1）具体要求", None), ("这里是具体要求的完整正文。", None),
    ], {"一、总体要求": 1, "（1）具体要求": 2}),
])
def test_user_specified_local_family_hierarchy_through_api(paragraphs, expected, tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, paragraphs)
    for text, level in expected.items():
        e = actual[text]
        assert e["selected_decision"]["role"] == f"heading-{level}"
        assert e["selected_decision"]["action"] == "mapped"
    for text, level in expected.items():
        assert actual[text]["features"]["local_heading_level"] == level


def test_numeric_schema_does_not_reserve_absent_layers_for_style(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("5. 技术评分支撑材料", 2), ("5.1 响应情况", 3), ("这里是响应情况正文。", None),
        ("5.2 合同业绩", None), ("5.2.1 业绩说明", None), ("这里是业绩说明正文。", None),
        ("6. 其他材料", 2), ("这里是其他材料正文。", None),
    ])
    for text, level in {"5. 技术评分支撑材料": 1, "5.1 响应情况": 2,
                        "5.2 合同业绩": 2, "5.2.1 业绩说明": 3, "6. 其他材料": 1}.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"


def test_parent_heading_above_six_is_not_clamped_or_mapped(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("第一章 父章节", None), ("第一节 工作原则", None), ("一、基本要求", None),
        ("（一）具体要求", None), ("1. 服务内容", None), ("（1）技术要求", None),
        ("① 超范围子主题", None), ("这里是子主题正文。", None),
    ])
    for text in ("① 超范围子主题",):
        e = actual[text]
        assert e["selected_decision"]["action"] != "mapped"
        assert e["features"]["local_hierarchy_warning"]


@pytest.mark.parametrize("paragraphs,expected", [
    ([("前言", None), ("这里是前言正文。", None), ("一、工作背景", None),
      ("（一）基本情况", None), ("这里是基本情况正文。", None)],
     {"前言": 1, "一、工作背景": 1, "（一）基本情况": 2}),
    ([("第一章 总体情况", None), ("总则", None), ("这里是总则正文。", None)],
     {"第一章 总体情况": 1, "总则": 2}),
    ([("第一章 总体情况", None), ("第一节 工作原则", None), ("说明", None),
      ("这里是说明正文。", None)],
     {"第一章 总体情况": 1, "第一节 工作原则": 2, "说明": 3}),
])
def test_relative_headings_do_not_pad_numbered_schema(paragraphs, expected, tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, paragraphs)
    for text, level in expected.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"
    for regime in actual["_analysis"]["document_baselines"]["hierarchy_regimes"]:
        assert "RELATIVE_HEADING" not in regime["active_families"]


def test_manual_lists_cannot_activate_a_heading_family(tmp_path, monkeypatch):
    items = ["1. 《A办法》", "2. 《B规定》", "3. 《C标准》"]
    actual = _convert(tmp_path, monkeypatch, [
        ("一、管理依据", None), *((text, None) for text in items),
        ("二、实施要求", None), ("这里是实施要求正文。", None),
    ])
    audit = actual["_analysis"]["document_baselines"]
    assert len(audit["hierarchy_schemas"]) == 3
    assert all("ARABIC_ENUM" not in r["active_families"] for r in audit["hierarchy_regimes"])
    node_ids = {n["element_id"] for n in audit["heading_tree"]}
    assert all(actual[t]["element_id"] not in node_ids for t in items)
    assert actual["一、管理依据"]["selected_decision"]["role"] == "heading-1"
    assert actual["二、实施要求"]["selected_decision"]["role"] == "heading-1"


def test_heading_body_contrast_compresses_two_active_families(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、总体要求", None), ("1. 服务内容", None), ("这里是服务内容正文。", None),
        ("2. 质量要求", None), ("这里是质量要求正文。", None),
    ])
    assert actual["一、总体要求"]["selected_decision"]["role"] == "heading-1"
    for text in ("1. 服务内容", "2. 质量要求"):
        assert actual[text]["selected_decision"]["role"] == "heading-2"
        assert actual[text]["features"]["parent_heading_id"] == actual["一、总体要求"]["element_id"]


def test_mixed_chinese_and_decimal_attachment_regimes(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、工作要求", None), ("（一）基本原则", None), ("这里是基本原则正文。", None),
        ("附件A", None), ("1 技术参数", None), ("1.1 性能要求", None),
        ("1.1.1 测试方法", None), ("这里是测试方法正文。", None),
    ])
    for text, level in {"一、工作要求": 1, "（一）基本原则": 2, "1 技术参数": 1,
                        "1.1 性能要求": 2, "1.1.1 测试方法": 3}.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"
    a = actual["一、工作要求"]["features"]["hierarchy_regime_id"]
    b = actual["1 技术参数"]["features"]["hierarchy_regime_id"]
    assert a != b
    assert actual["1 技术参数"]["features"]["hierarchy_schema"] == "decimal_technical"


def test_nested_regime_return_restores_parent_tree(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、总方案", 1), ("设备章节", 2), ("一、设备分类", None),
        ("（一）技术指标", None), ("这里是技术指标正文。", None),
        ("交付章节", 2), ("这里是交付章节正文。", None),
    ])
    assert actual["一、设备分类"]["selected_decision"]["role"] == "heading-3"
    assert actual["（一）技术指标"]["selected_decision"]["role"] == "heading-4"
    inner_id = actual["一、设备分类"]["features"]["hierarchy_regime_id"]
    inner = next(r for r in actual["_analysis"]["document_baselines"]["hierarchy_regimes"] if r["regime_id"] == inner_id)
    assert inner["effective_levels"]["CN_ENUM"] == 3
    assert inner["effective_levels"]["CN_BRACKET"] == 4
    assert actual["交付章节"]["features"]["parent_heading_id"] == actual["一、总方案"]["element_id"]


@pytest.mark.parametrize("section_break", [False, True])
def test_stable_decimal_reset_is_a_separate_regime(section_break, tmp_path, monkeypatch):
    paragraphs = [("一、正文要求", None), ("（一）工作原则", None), ("这里是工作原则正文。", None)]
    if section_break:
        paragraphs.append(("", 0))
    paragraphs.extend([("1 技术范围", None), ("1.1 基本参数", None), ("1.1.1 参数测试", None),
                       ("这里是参数测试正文。", None)])
    actual = _convert(tmp_path, monkeypatch, paragraphs)
    assert actual["1 技术范围"]["selected_decision"]["role"] == "heading-1"
    assert actual["1.1 基本参数"]["selected_decision"]["role"] == "heading-2"
    assert actual["1 技术范围"]["features"]["hierarchy_regime_id"] != actual["一、正文要求"]["features"]["hierarchy_regime_id"]


def test_parallel_decimal_clauses_are_not_activated_as_headings(tmp_path, monkeypatch):
    items = ["1.1.1 核对库存，记录差异", "1.1.2 检查批次，核实数量", "1.1.3 复查储位，归档记录"]
    actual = _convert(tmp_path, monkeypatch, [
        ("1.1 日常管理", 2), *((text, None) for text in items),
        ("2 附则", 1), ("这里是附则正文。", None),
    ])
    for text in items:
        assert actual[text]["selected_decision"]["role"] == "body-paragraph"
        assert actual[text]["features"]["manual_sequence_candidate"]


def test_dirty_native_heading_list_cannot_fall_back_to_legacy(tmp_path, monkeypatch):
    items = ["1. 《A办法》", "2. 《B规定》", "3. 《C标准》"]
    actual = _convert(tmp_path, monkeypatch, [(text, 3) for text in items])
    for text in items:
        assert not actual[text]["selected_decision"]["role"].startswith("heading-")
        assert actual[text]["selected_decision"]["action"] in {"mapped", "preserved"}
        assert actual[text]["features"]["manual_sequence_candidate"]


def test_native_unnumbered_headings_use_effective_parent_levels(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("独立业务章节", 3), ("这里是独立业务章节正文。", None),
        ("5. 技术体系", 2), ("业务实现说明", 3), ("这里是业务实现说明正文。", None),
    ])
    assert actual["独立业务章节"]["selected_decision"]["role"] == "heading-1"
    assert actual["业务实现说明"]["selected_decision"]["role"] == "heading-2"
    assert actual["业务实现说明"]["features"]["parent_heading_id"] == actual["5. 技术体系"]["element_id"]


def test_numeric_parent_without_native_style_can_own_a_nested_schema(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、总论", 1), ("5. 技术支撑", 2), ("5.1 配套要求", None),
        ("（一）实施要点", 4), ("这里是实施要点正文。", None),
        ("（二）验收要点", 4), ("这里是验收要点正文。", None),
        ("5.2 其他要求", None), ("这里是其他要求正文。", None),
    ])
    for text, level in {"一、总论": 1, "5. 技术支撑": 2, "5.1 配套要求": 3,
                        "（一）实施要点": 4, "（二）验收要点": 4, "5.2 其他要求": 3}.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"
    assert actual["（一）实施要点"]["features"]["parent_heading_id"] == actual["5.1 配套要求"]["element_id"]
    assert actual["5.2 其他要求"]["features"]["parent_heading_id"] == actual["5. 技术支撑"]["element_id"]


def test_two_adjacent_items_with_only_trailing_body_are_uncertain(tmp_path, monkeypatch):
    items = ["1、检查仓位", "2、核对数量"]
    actual = _convert(tmp_path, monkeypatch, [
        *((text, None) for text in items), ("这段正文描述上述清单。", None),
    ])
    for text in items:
        assert actual[text]["selected_decision"]["action"] == "preserved"
        assert not actual[text]["features"]["heading_candidate"]


def test_nested_schema_cannot_restart_below_an_unsupported_parent(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("第一章 父章节", 1), ("第一节 工作原则", None), ("一、基本要求", None),
        ("（一）具体要求", None), ("1. 服务内容", None), ("（1）技术要求", None),
        ("说明", 6), ("一、嵌套范围", None), ("这里是嵌套范围正文。", None),
    ])
    e = actual["一、嵌套范围"]
    assert e["selected_decision"]["action"] == "preserved"
    assert e["features"]["effective_heading_level"] is None
    assert "unsupported_parent" in e["features"]["local_hierarchy_warning"]


def test_new_numeric_root_closes_the_previous_nested_chinese_scope(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("5. 技术支撑", 2), ("5.1 配套要求", None),
        ("（一）实施要点", 4), ("这里是实施要点正文。", None),
        ("6 其他要求", None), ("6.1 交付材料", None), ("这里是交付材料正文。", None),
    ])
    for text, level in {"5. 技术支撑": 1, "5.1 配套要求": 2,
                        "（一）实施要点": 3, "6 其他要求": 1, "6.1 交付材料": 2}.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"
    assert actual["6 其他要求"]["features"]["parent_heading_id"] is None


def test_isolated_number_after_section_break_cannot_borrow_global_peers(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、总体要求", None), ("（一）工作原则", None), ("这里是工作原则正文。", None),
        ("", 0), ("1. 检查范围", None), ("这里是范围检查的完整正文。", None),
    ])
    assert actual["1. 检查范围"]["selected_decision"]["action"] == "preserved"
    assert not actual["1. 检查范围"]["features"]["heading_candidate"]


def test_native_numeric_children_continue_until_the_true_numeric_parent_peer(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("7. 服务方案", 1), ("7.4 碳服务", 2), ("7.4.3 计算方法", 3),
        ("1. 参数说明", 4), ("这里是参数说明正文。", None),
        ("2. 计算安排", 4), ("（1）电力计算", None), ("这里是电力计算正文。", None),
        ("（2）燃料计算", None), ("这里是燃料计算正文。", None),
        ("7.4.4 其他方法", 3), ("这里是其他方法正文。", None),
    ])
    for text, level in {"1. 参数说明": 4, "2. 计算安排": 4, "（1）电力计算": 5,
                        "（2）燃料计算": 5, "7.4.4 其他方法": 3}.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"


def test_same_native_level_numbered_peer_does_not_nest_under_unnumbered_heading(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、总论", 1), ("5. 技术支撑", 2), ("编制说明", 3), ("这里是编制说明正文。", None),
        ("5.1 响应情况", 3), ("这里是响应情况正文。", None),
        ("5.2 业绩情况", None), ("这里是业绩情况正文。", None),
    ])
    for text in ("5.1 响应情况", "5.2 业绩情况"):
        assert actual[text]["selected_decision"]["role"] == "heading-3"
        assert actual[text]["features"]["parent_heading_id"] == actual["5. 技术支撑"]["element_id"]


def test_decimal_in_chinese_schema_requires_its_actual_prefix_parent(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("一、总体要求", None), ("（一）基本原则", None), ("1. 服务要点", None),
        ("这里是服务要点正文。", None), ("缺失编号的主题", None),
        ("4.1 工作安排", None), ("这里是工作安排正文。", None),
        ("4.2 验收安排", None), ("这里是验收安排正文。", None),
    ])
    for text in ("4.1 工作安排", "4.2 验收安排"):
        assert actual[text]["selected_decision"]["action"] == "preserved"
        assert "missing_numbered_parent" in actual[text]["features"]["local_hierarchy_warning"]


def test_restarted_numbering_under_native_parent_is_a_local_child_series(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("7. 服务方案", 1), ("7.4 服务细则", 2), ("7.4.9 保密方案", 3),
        ("1. 总则", 4), ("1. 目的", None), ("这里是目的正文。", None),
        ("2. 范围", None), ("这里是范围正文。", None),
        ("2. 组织架构", 4), ("1. 安全小组", None), ("这里是安全小组正文。", None),
        ("7.4.10 其他方案", 3), ("这里是其他方案正文。", None),
    ])
    for text, level in {"1. 总则": 4, "1. 目的": 5, "2. 范围": 5,
                        "2. 组织架构": 4, "1. 安全小组": 5, "7.4.10 其他方案": 3}.items():
        assert actual[text]["selected_decision"]["role"] == f"heading-{level}"
        assert actual[text]["selected_decision"]["action"] == "mapped"


def test_plain_numeric_child_series_under_an_anchored_decimal_parent(tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("7. 服务方案", 1), ("7.4 服务细则", 2), ("7.4.7 沟通方案", 3),
        ("1. 沟通内容", None), ("这里是沟通内容正文。", None),
        ("2. 沟通安排", None), ("这里是沟通安排正文。", None),
        ("7.4.8 验收方案", 3), ("这里是验收方案正文。", None),
    ])
    for text in ("1. 沟通内容", "2. 沟通安排"):
        assert actual[text]["selected_decision"]["role"] == "heading-4"
        assert actual[text]["features"]["parent_heading_id"] == actual["7.4.7 沟通方案"]["element_id"]
    assert actual["7.4.8 验收方案"]["selected_decision"]["role"] == "heading-3"


@pytest.mark.parametrize("missing", ["1.2.1 缺失父主题的内容", "2.1 缺失根主题的内容"])
def test_decimal_depth_cannot_attach_to_a_different_prefix_parent(missing, tmp_path, monkeypatch):
    actual = _convert(tmp_path, monkeypatch, [
        ("1 技术范围", 1), ("1.1 正常主题", None), ("这里是正常主题正文。", None),
        (missing, None), ("这里是缺失父主题内容的正文。", None),
        ("1.3 后续主题", None), ("这里是后续主题正文。", None),
    ])
    assert actual[missing]["selected_decision"]["action"] == "preserved"
    assert "missing_numbered_parent" in actual[missing]["features"]["local_hierarchy_warning"]
    assert actual["1.3 后续主题"]["selected_decision"]["role"] == "heading-2"
