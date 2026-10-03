import struct
import zlib
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient

import doc_converter.api as api_module
import doc_converter.importers.doc as doc_importer_module
from doc_converter.api import _suggest_semantic_roles, create_app


def test_local_api_creates_editable_job_saves_draft_and_renders_named_job(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))

    created_pack = client.post(
        "/api/style-packs",
        json={"name": "测试正式报告", "format_id": "docx"},
    )
    assert created_pack.status_code == 201
    pack = created_pack.json()

    created_job = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("实施方案.md", "# 实施方案\n\n# 一、建设背景\n\n正文。", "text/markdown")},
    )
    assert created_job.status_code == 201
    job = created_job.json()
    assert job["status"] == "ready"
    assert job["suggestions"] == [
        {"line": 1, "role": "document-title", "score": 0.9, "reason": "首个无编号 H1"}
    ]
    job_directory = tmp_path / "jobs" / job["id"]
    assert (job_directory / "source" / "original.md").read_text(encoding="utf-8") == (
        "# 实施方案\n\n# 一、建设背景\n\n正文。"
    )
    assert (job_directory / "source" / "imported.md").read_text(encoding="utf-8") == (
        "# 实施方案\n\n# 一、建设背景\n\n正文。"
    )
    assert (job_directory / "work" / "source.md").read_text(encoding="utf-8") == (
        "# 实施方案\n\n# 一、建设背景\n\n正文。"
    )

    draft = client.put(
        f"/api/jobs/{job['id']}/draft",
        json={"source_markdown": "---\ntitle: 实施方案\n---\n\n# 一、建设背景\n\n正文。"},
    )
    assert draft.status_code == 200
    assert draft.json()["draft_updated_at"] is not None

    completed_job = client.post(
        f"/api/jobs/{job['id']}/normalize-and-render",
        json={"source_markdown": "---\ntitle: 实施方案\n---\n\n# 一、建设背景\n\n正文。", "confirmed_roles": {}},
    )
    assert completed_job.status_code == 200
    assert completed_job.json()["status"] == "completed"
    assert completed_job.json()["output_name"] == "实施方案.docx"
    assert "实施方案.docx" in completed_job.json()["artifacts"]

    docx = client.get(f"/api/jobs/{job['id']}/artifacts/实施方案.docx")
    assert docx.status_code == 200
    assert docx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_api_imports_docx_media_and_returns_backend_working_copy(tmp_path, monkeypatch) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "导入测试", "format_id": "docx"}).json()
    source_bytes = _make_docx_bytes("中文标题", image_color=(240, 10, 10))
    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda _source, _output: False)

    created = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"]), "conversion_mode": "round-trip"},
        files={"source": ("中文材料.docx", source_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert created.status_code == 201, created.text
    job = created.json()
    assert job["source_format"] == "docx"
    assert job["source_original_name"] == "中文材料.docx"
    assert job["import_warnings"] == []
    job_directory = tmp_path / "jobs" / job["id"]
    assert (job_directory / "source" / "original.docx").read_bytes() == source_bytes

    draft = client.get(f"/api/jobs/{job['id']}/draft")

    assert draft.status_code == 200
    payload = draft.json()
    assert payload["source_format"] == "docx"
    assert payload["warnings"] == []
    assert "# 中文标题" in payload["source_markdown"]
    assert "正文含中文。" in payload["source_markdown"]
    assert "甲项目" in payload["source_markdown"]
    assert "source/assets/" in payload["source_markdown"]
    imported_markdown = (job_directory / "source" / "imported.md").read_text(encoding="utf-8")
    assert imported_markdown == payload["source_markdown"]
    assert len(list((job_directory / "source" / "assets").iterdir())) == 1

    rendered = client.post(
        f"/api/jobs/{job['id']}/normalize-and-render",
        json={"source_markdown": payload["source_markdown"], "confirmed_roles": {}},
    )

    assert rendered.status_code == 200, rendered.text
    output = client.get(f"/api/jobs/{job['id']}/artifacts/{rendered.json()['output_name']}")
    document = Document(BytesIO(output.content))
    assert len(document.inline_shapes) == 1
    assert "甲项目" in "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert (job_directory / "source" / "original.docx").read_bytes() == source_bytes


def test_api_rejects_invalid_docx_upload_with_422(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "坏 DOCX", "format_id": "docx"}).json()

    response = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("坏文件.docx", b"not a zip", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 422
    assert "DOCX ZIP" in response.json()["detail"]


def test_api_direct_docx_job_preserves_source_and_exposes_audit(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda _source, _output: False)
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "直接规范化", "format_id": "docx"}).json()
    document = Document(BytesIO(_make_docx_bytes("原始标题", image_color=(20, 120, 80))))
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "横向合并"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵向合并"
    source_buffer = BytesIO()
    document.save(source_buffer)
    source_bytes = source_buffer.getvalue()

    created = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("原始材料.docx", source_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert created.status_code == 201, created.text
    job = created.json()
    assert job["conversion_mode"] == "docx-direct"
    assert job["status"] == "ready"
    job_directory = tmp_path / "jobs" / job["id"]
    assert not (job_directory / "source" / "imported.md").exists()
    assert client.get(f"/api/jobs/{job['id']}/draft").status_code == 404

    rendered = client.post(
        f"/api/jobs/{job['id']}/normalize-docx-direct",
        json={"style_pack_id": pack["id"], "style_pack_version": pack["version"]},
    )
    assert rendered.status_code == 200, rendered.text
    completed = rendered.json()
    assert completed["status"] == "completed"
    assert completed["output_name"] == "原始材料_规范化.docx"
    assert completed["render_revision"] == 1
    assert set(completed["artifacts"]) == {
        "原始材料_规范化.docx", "docx-analysis.json", "docx-normalization-report.json", "style-pack.snapshot.json"
    }
    assert (job_directory / "source" / "original.docx").read_bytes() == source_bytes
    output = client.get(f"/api/jobs/{job['id']}/artifacts/{completed['output_name']}")
    assert output.status_code == 200
    assert len(Document(BytesIO(output.content)).inline_shapes) == 1
    output_table = Document(BytesIO(output.content)).tables[1]
    source_table = Document(BytesIO(source_bytes)).tables[1]
    assert len(output_table._tbl.xpath(".//w:gridSpan")) == len(source_table._tbl.xpath(".//w:gridSpan")) == 1
    assert len(output_table._tbl.xpath(".//w:vMerge")) == len(source_table._tbl.xpath(".//w:vMerge")) == 2
    assert [len(row._tr.xpath("./w:tc")) for row in output_table.rows] == [
        len(row._tr.xpath("./w:tc")) for row in source_table.rows
    ]
    assert client.get(f"/api/jobs/{job['id']}/artifacts/docx-normalization-report.json").status_code == 200
    assert client.post(f"/api/jobs/{job['id']}/normalize-and-render", json={"confirmed_roles": {}}).status_code == 409


def test_api_rejects_invalid_direct_mode_and_non_docx_source(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "模式限制", "format_id": "docx"}).json()
    fields = {"style_pack_id": pack["id"], "style_pack_version": str(pack["version"]), "conversion_mode": "docx-direct"}
    wrong_source = client.post("/api/jobs", data=fields, files={"source": ("plain.md", "# 标题", "text/markdown")})
    invalid_docx = client.post("/api/jobs", data=fields, files={"source": ("bad.docx", b"bad zip", "application/octet-stream")})
    assert wrong_source.status_code == 422
    assert invalid_docx.status_code == 422
    assert list((tmp_path / "jobs").iterdir()) == []


def test_docx_assets_are_isolated_between_jobs(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "隔离测试", "format_id": "docx"}).json()

    jobs = []
    for title, color in (("红色图片", (240, 10, 10)), ("蓝色图片", (10, 10, 240))):
        response = client.post(
            "/api/jobs",
            data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"]), "conversion_mode": "round-trip"},
            files={"source": (f"{title}.docx", _make_docx_bytes(title, image_color=color), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 201, response.text
        jobs.append(response.json())

    asset_dirs = [tmp_path / "jobs" / job["id"] / "source" / "assets" for job in jobs]
    first_asset = next(path for path in asset_dirs[0].rglob("*") if path.is_file())
    second_asset = next(path for path in asset_dirs[1].rglob("*") if path.is_file())
    assert asset_dirs[0] != asset_dirs[1]
    assert first_asset.read_bytes() != second_asset.read_bytes()


def test_api_converts_doc_to_intermediate_docx_without_changing_upload(tmp_path, monkeypatch) -> None:
    original_bytes = b"legacy binary DOC upload"
    converted_docx = _make_docx_bytes("旧版 Word 中文标题", image_color=(20, 120, 80))

    def fake_convert(source_path: Path, *, output_path: Path) -> Path:
        assert source_path.read_bytes() == original_bytes
        output_path.write_bytes(converted_docx)
        return output_path

    monkeypatch.setattr(doc_importer_module, "convert_doc_to_docx", fake_convert)
    client = TestClient(create_app(data_root=tmp_path / "app-data"))
    pack = client.post("/api/style-packs", json={"name": "旧版格式测试", "format_id": "docx"}).json()

    created = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("旧版材料.doc", original_bytes, "application/msword")},
    )

    assert created.status_code == 201, created.text
    job = created.json()
    assert job["source_format"] == "doc"
    job_directory = tmp_path / "app-data" / "jobs" / job["id"]
    assert (job_directory / "source" / "original.doc").read_bytes() == original_bytes
    assert (job_directory / "source" / "intermediate.docx").is_file()
    draft = client.get(f"/api/jobs/{job['id']}/draft").json()
    assert "旧版 Word 中文标题" in draft["source_markdown"]


def test_doc_import_reports_when_configured_libreoffice_path_is_missing(tmp_path, monkeypatch) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "DOC 错误提示", "format_id": "docx"}).json()
    missing_path = tmp_path / "missing" / "soffice.exe"
    monkeypatch.setenv("DOC_CONVERTER_LIBREOFFICE", str(missing_path))

    response = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("材料.doc", b"legacy Word bytes", "application/msword")},
    )

    assert response.status_code == 422
    assert "DOC_CONVERTER_LIBREOFFICE" in response.json()["detail"]
    assert "LibreOffice" in response.json()["detail"]


def _make_docx_bytes(title: str, *, image_color: tuple[int, int, int]) -> bytes:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph("正文含中文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "甲项目"
    table.cell(1, 1).text = "三项"
    image_buffer = BytesIO(_solid_png(image_color))
    image_buffer.seek(0)
    document.add_picture(image_buffer, width=Inches(0.25))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _solid_png(color: tuple[int, int, int]) -> bytes:
    def chunk(name: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + name
            + payload
            + struct.pack(">I", zlib.crc32(name + payload) & 0xFFFFFFFF)
        )

    row = b"\x00" + bytes(color) * 8
    header = struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(row * 8))
        + chunk(b"IEND", b"")
    )


def test_structure_suggestions_cover_all_title_area_roles() -> None:
    source = """---
title: 项目方案
---

项目实施说明

摘要：这里是摘要内容。

关键词：项目；实施；质量

# 华软钢

重庆市公共资源交易中心

采购管理部

2026年9月18日

# 一、建设背景

正文内容。
"""

    suggestions = _suggest_semantic_roles(source)

    assert {item["role"] for item in suggestions} == {
        "document-subtitle",
        "abstract",
        "keywords",
        "author",
        "organization",
        "department",
        "date",
    }
    assert {item["line"] for item in suggestions} == {5, 7, 9, 11, 13, 15, 17}


def test_api_renders_generated_comprehensive_fixture_without_subtitle_confirmation(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "综合测试", "format_id": "docx"}).json()
    source = (Path(__file__).parent / "fixtures" / "comprehensive.md").read_text(encoding="utf-8")

    created = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("综合测试.md", source, "text/markdown")},
    )
    assert created.status_code == 201
    job = created.json()

    completed = client.post(
        f"/api/jobs/{job['id']}/normalize-and-render",
        json={"source_markdown": source, "confirmed_roles": {}},
    )

    assert completed.status_code == 200
    payload = completed.json()
    assert payload["output_name"] == "人工智能+公共资源交易创新实践.docx"
    assert payload["status"] == "completed"

    response = client.get(f"/api/jobs/{job['id']}/artifacts/{payload['output_name']}")
    document = Document(BytesIO(response.content))
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "基于14个省级创新样本的实践总结" in paragraph_text
    assert "摘要：本文从治理、安全、技术和生态四个方面总结公共资源交易数智化实践，并给出可执行的推进路径。" in paragraph_text
    assert "关键词：人工智能；大模型；公共资源交易；数智化" in paragraph_text


def test_render_uses_current_style_pack_and_advances_render_revision(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    original = client.post("/api/style-packs", json={"name": "项目报告", "format_id": "docx"}).json()
    job = client.post(
        "/api/jobs",
        data={"style_pack_id": original["id"], "style_pack_version": str(original["version"])},
        files={"source": ("方案.md", "# 方案\n\n# 一、范围\n\n正文。", "text/markdown")},
    ).json()
    revised = client.put(
        f"/api/style-packs/{original['id']}",
        json={
            "base_version": original["version"],
            "name": original["name"],
            "format_id": original["format_id"],
            "roles": original["roles"],
        },
    ).json()
    render_request = {
        "source_markdown": "# 方案\n\n# 一、范围\n\n正文。",
        "confirmed_roles": {},
        "style_pack_id": revised["id"],
        "style_pack_version": revised["version"],
    }

    first = client.post(f"/api/jobs/{job['id']}/normalize-and-render", json=render_request)
    second = client.post(f"/api/jobs/{job['id']}/normalize-and-render", json=render_request)

    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["style_pack_version"] == revised["version"]
    assert first.json()["render_revision"] == 1
    assert second.json()["render_revision"] == 2
    snapshot = (
        tmp_path / "jobs" / job["id"] / "style-pack.snapshot.json"
    ).read_text(encoding="utf-8")
    assert f'"version": {revised["version"]}' in snapshot


def test_render_precomputes_inline_pdf_preview_before_marking_job_completed(tmp_path, monkeypatch) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "项目报告", "format_id": "docx"}).json()
    source = "# 方案\n\n# 一、范围\n\n正文。"
    job = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("方案.md", source, "text/markdown")},
    ).json()

    calls: list[tuple[Path, Path]] = []

    def fake_convert(source_path: Path, output_path: Path) -> bool:
        calls.append((source_path, output_path))
        assert source_path.exists()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"%PDF-1.4 preloaded preview")
        return True

    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", fake_convert)
    rendered = client.post(
        f"/api/jobs/{job['id']}/normalize-and-render",
        json={"source_markdown": source, "confirmed_roles": {}},
    )
    assert rendered.status_code == 200
    payload = rendered.json()
    assert len(calls) == 1
    assert payload["preview_ready"] is True
    assert payload["status"] == "completed"

    preview = client.get(f"/api/jobs/{job['id']}/preview?revision=1")

    assert preview.status_code == 200
    assert preview.headers["content-type"].startswith("application/pdf")
    assert preview.headers["content-disposition"].startswith("inline")
    assert preview.content == b"%PDF-1.4 preloaded preview"
    assert len(calls) == 1


def test_preview_endpoint_retains_lazy_conversion_fallback(tmp_path, monkeypatch) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "项目报告", "format_id": "docx"}).json()
    source = "# 方案\n\n# 一、范围\n\n正文。"
    job = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("方案.md", source, "text/markdown")},
    ).json()

    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", lambda _source_path, _output_path: False)
    rendered = client.post(
        f"/api/jobs/{job['id']}/normalize-and-render",
        json={"source_markdown": source, "confirmed_roles": {}},
    )
    assert rendered.status_code == 200
    assert rendered.json()["preview_ready"] is False

    calls: list[tuple[Path, Path]] = []

    def fake_convert(source_path: Path, output_path: Path) -> bool:
        calls.append((source_path, output_path))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"%PDF-1.4 local preview")
        return True

    monkeypatch.setattr(api_module, "_convert_docx_to_pdf", fake_convert)
    preview = client.get(f"/api/jobs/{job['id']}/preview?revision=1")

    assert preview.status_code == 200
    assert preview.headers["content-type"].startswith("application/pdf")
    assert preview.content == b"%PDF-1.4 local preview"
    assert calls and calls[0][0].suffix == ".docx"


def test_application_serves_bundled_frontend_without_a_second_service(tmp_path) -> None:
    frontend_dist = tmp_path / "frontend-dist"
    frontend_dist.mkdir()
    (frontend_dist / "index.html").write_text("<title>Doc Converter</title>", encoding="utf-8")

    client = TestClient(create_app(data_root=tmp_path / "data", frontend_dist=frontend_dist))

    response = client.get("/")

    assert response.status_code == 200
    assert "Doc Converter" in response.text


def test_style_pack_edit_creates_a_new_immutable_api_revision(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    original = client.post(
        "/api/style-packs",
        json={"name": "项目报告", "format_id": "docx", "roles": {"body": {"font_size_pt": 12}}},
    ).json()

    revised = client.put(
        f"/api/style-packs/{original['id']}",
        json={
            "base_version": original["version"],
            "name": "项目报告",
            "format_id": "docx",
            "roles": {"body": {"font_size_pt": 14}},
        },
    )

    assert revised.status_code == 201
    assert revised.json()["version"] == 2
    listed = client.get("/api/style-packs").json()
    assert next(pack for pack in listed if pack["id"] == original["id"])["version"] == 2


def test_draft_endpoint_rejects_missing_source_markdown(tmp_path) -> None:
    client = TestClient(create_app(data_root=tmp_path))
    pack = client.post("/api/style-packs", json={"name": "项目报告", "format_id": "docx"}).json()
    job = client.post(
        "/api/jobs",
        data={"style_pack_id": pack["id"], "style_pack_version": str(pack["version"])},
        files={"source": ("方案.md", "# 方案\n", "text/markdown")},
    ).json()

    response = client.put(f"/api/jobs/{job['id']}/draft", json={})

    assert response.status_code == 422


def test_data_root_environment_override_is_used_for_container_storage(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("DOC_CONVERTER_DATA_ROOT", str(tmp_path))
    client = TestClient(create_app())

    response = client.get("/api/style-packs")

    assert response.status_code == 200
    assert (tmp_path / "style-packs").exists()
