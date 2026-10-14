from docx import Document
from docx.oxml.ns import qn

from doc_converter.normalizer import normalize_markdown
from doc_converter.renderer import render_docx
from doc_converter.style_library import create_style_pack


def test_scored_roles_and_inline_metadata_survive_actual_docx_render(tmp_path):
    source = "# 仓储方案\n\n摘要：保留 **重点**、[链接](https://example.com/abstract)、`代码`。\n\n单位：测试公司\n\n部门：物资部\n\n一、背景\n\n（一）现状\n\n1. 配送业务\n\n（1）问题分析\n\n正文保留[链接](https://example.com/body)。"
    result = normalize_markdown(source, filename="方案.md")
    output = tmp_path / "result.docx"
    render_docx(markdown=result.markdown, style_pack=create_style_pack(name="测试", format_id="docx"), output_path=output, work_dir=tmp_path / "work")
    document = Document(output)
    styles = {p.text: p.style.style_id for p in document.paragraphs}
    assert styles["仓储方案"] == "DC_DocumentTitle"
    assert styles["一、背景"] == "DC_Heading1"
    assert styles["（一）现状"] == "DC_Heading2"
    assert styles["1. 配送业务"] == "DC_Heading3"
    assert styles["（1）问题分析"] == "DC_Heading4"
    assert styles["测试公司"] == styles["物资部"] == "DC_OrganizationDepartment"
    abstract = next(p for p in document.paragraphs if p.text.startswith("摘要："))
    assert abstract.text == "摘要：保留 重点、链接、代码。"
    assert abstract.style.style_id == "DC_AbstractLabel"
    assert len(abstract._p.findall(qn("w:hyperlink"))) == 1
    targets = {r.target_ref for r in document.part.rels.values() if r.reltype.endswith('/hyperlink')}
    assert targets == {"https://example.com/abstract", "https://example.com/body"}


def test_metadata_reference_link_and_footnote_share_body_definitions(tmp_path):
    import zipfile
    source = '# [实施方案][s]\n\n摘要：结论见注释[^note]。\n\n一、正文\n\n正文内容。\n\n[s]: https://example.com/title\n\n[^note]: 注释具体内容。'
    normalized = normalize_markdown(source, filename='方案.md')
    output = tmp_path / 'references.docx'
    render_docx(markdown=normalized.markdown, style_pack=create_style_pack(name='测试', format_id='docx'), output_path=output, work_dir=tmp_path / 'work')
    document = Document(output)
    assert document.paragraphs[0].text == '实施方案'
    assert any(r.target_ref == 'https://example.com/title' for r in document.part.rels.values() if r.reltype.endswith('/hyperlink'))
    assert any(p._p.xpath('.//w:footnoteReference') for p in document.paragraphs if p.text.startswith('摘要：'))
    with zipfile.ZipFile(output) as archive:
        assert '注释具体内容。' in archive.read('word/footnotes.xml').decode()


def test_resolved_caption_and_reference_boundary_roles_are_actual_output_styles(tmp_path):
    source = '# 方案\n\n一、背景\n\nTable 1 Results\n\n| 项目 | 值 |\n| --- | --- |\n| A | 1 |\n\n参考文献\n\n[1] 研究报告。\n\n## 附录\n\n附件正文。'
    normalized = normalize_markdown(source, filename='方案.md')
    output = tmp_path / 'roles.docx'
    render_docx(markdown=normalized.markdown, style_pack=create_style_pack(name='测试', format_id='docx'), output_path=output, work_dir=tmp_path / 'work')
    styles = {p.text: p.style.style_id for p in Document(output).paragraphs}
    assert styles['Table 1 Results'] == 'DC_TableCaption'
    assert styles['附件正文。'] == 'DC_BodyParagraph'


def test_actual_ordered_list_is_two_numbered_paragraphs(tmp_path):
    source = '# 方案\n\n一、背景\n\n1. 完成检查。\n2. 开始配送。'
    normalized = normalize_markdown(source, filename='方案.md')
    output = tmp_path / 'lists.docx'
    render_docx(markdown=normalized.markdown, style_pack=create_style_pack(name='测试', format_id='docx'), output_path=output, work_dir=tmp_path / 'work')
    items = [p for p in Document(output).paragraphs if p.text in {'完成检查。', '开始配送。'}]
    assert len(items) == 2
    assert all(p._p.xpath('./w:pPr/w:numPr') for p in items)


def test_renderer_accepts_relative_output_and_work_directory(tmp_path, monkeypatch):
    from pathlib import Path
    monkeypatch.chdir(tmp_path)
    render_docx(markdown='# 方案\n\n正文。', style_pack=create_style_pack(name='测试', format_id='docx'), output_path=Path('artifacts/result.docx'), work_dir=Path('artifacts/work'))
    assert Document('artifacts/result.docx').paragraphs[0].text == '方案'
