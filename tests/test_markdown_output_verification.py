from docx import Document

from doc_converter.normalizer import normalize_markdown
from doc_converter.renderer import render_docx
from doc_converter.style_library import create_style_pack


def test_render_reports_written_effective_styles_and_heading_levels(tmp_path):
    pack = create_style_pack(name="验收", format_id="docx")
    n = normalize_markdown("一、范围\n\n正文。\n\n二、职责\n\n正文。", filename="验收.md")
    r = render_docx(
        markdown=n.markdown,
        style_pack=pack,
        output_path=tmp_path / "out.docx",
        work_dir=tmp_path / "work",
        role_decisions=n.report["decisions"],
    )
    assert r.verification["passed"]
    headings = [c for c in r.verification["checks"] if c["role"].startswith("heading-")]
    assert len(headings) == 2
    assert all(c["actual_style"] == "DC_Heading1" and c["actual_outline"] == 0 for c in headings)


def test_written_validation_detects_style_and_direct_outline_corruption(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from doc_converter.markdown_verification import verify_rendered_roles

    pack = create_style_pack(name="验收", format_id="docx")
    n = normalize_markdown("一、范围\n\n正文。\n\n二、职责\n\n正文。", filename="验收.md")
    path = tmp_path / "out.docx"
    render_docx(
        markdown=n.markdown,
        style_pack=pack,
        output_path=path,
        work_dir=tmp_path / "work",
        role_decisions=n.report["decisions"],
    )
    d = Document(path)
    d.paragraphs[0].style = "DC_BodyParagraph"
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "4")
    d.paragraphs[2]._p.get_or_add_pPr().append(outline)
    d.save(path)
    v = verify_rendered_roles(path, n.report["decisions"], pack)
    assert not v["passed"]
    assert v["failed_checks"] == 2


def test_written_validation_detects_direct_indent_corruption(tmp_path):
    from docx.shared import Pt

    from doc_converter.markdown_verification import verify_rendered_roles

    pack = create_style_pack(name="验收", format_id="docx")
    n = normalize_markdown("一、范围\n\n正文。\n\n二、职责\n\n正文。", filename="验收.md")
    path = tmp_path / "out.docx"
    render_docx(
        markdown=n.markdown,
        style_pack=pack,
        output_path=path,
        work_dir=tmp_path / "work",
        role_decisions=n.report["decisions"],
    )
    d = Document(path)
    d.paragraphs[0].paragraph_format.left_indent = Pt(50)
    d.save(path)
    v = verify_rendered_roles(path, n.report["decisions"], pack)
    assert not v["passed"]
    assert v["failed_checks"] == 1


def test_written_validation_detects_direct_heading_font_corruption(tmp_path):
    from docx.shared import Pt

    from doc_converter.markdown_verification import verify_rendered_roles

    pack = create_style_pack(name="验收", format_id="docx")
    n = normalize_markdown("一、范围\n\n正文。\n\n二、职责\n\n正文。", filename="验收.md")
    path = tmp_path / "font.docx"
    render_docx(
        markdown=n.markdown,
        style_pack=pack,
        output_path=path,
        work_dir=tmp_path / "work",
        role_decisions=n.report["decisions"],
    )
    d = Document(path)
    d.paragraphs[0].runs[0].font.size = Pt(80)
    d.save(path)
    v = verify_rendered_roles(path, n.report["decisions"], pack)
    assert not v["passed"]
