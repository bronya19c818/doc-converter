import os
import struct
import zipfile
import zlib

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from doc_converter.docx_postprocess import _REFERENCE_HEADING_PATTERN
from doc_converter.renderer import render_docx, resolve_pandoc
from doc_converter.style_library import create_style_pack


def test_renderer_creates_editable_docx_from_normalized_markdown(tmp_path) -> None:
    markdown = """---
schema: normalized-md/v1
title: 项目建设方案
organization: XX公司
department: XX部门
---

# 一、建设背景

正文包含 **重点内容** 和 [项目链接](https://example.com)。

| 项目 | 数量 |
| --- | ---: |
| A | 10 |
"""
    output_path = tmp_path / "项目建设方案.docx"

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="测试样式", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert output_path.exists()
    assert "项目建设方案" in text
    assert "一、建设背景" in text
    assert "重点内容" in text
    assert document.tables[0].cell(1, 1).text == "10"


def test_renderer_adds_half_point_grid_borders_to_tables(tmp_path) -> None:
    output_path = tmp_path / "table-grid.docx"
    render_docx(
        markdown="| 项目 | 数量 |\n| --- | ---: |\n| A | 10 |\n",
        style_pack=create_style_pack(name="表格边框测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    assert len(document.tables) == 1
    borders = document.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    for edge_name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        assert edge is not None
        assert edge.get(qn("w:val")) == "single"
        assert edge.get(qn("w:sz")) == "4"
        assert edge.get(qn("w:color")) == "000000"


def test_renderer_preserves_native_bullet_list_numbering(tmp_path) -> None:
    output_path = tmp_path / "bullet-list.docx"

    render_docx(
        markdown="- CASE-LIST-B1 一级项目符号 A\n- CASE-LIST-B2 一级项目符号 B\n",
        style_pack=create_style_pack(name="项目符号测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text]
    assert [paragraph.text for paragraph in paragraphs] == [
        "CASE-LIST-B1 一级项目符号 A",
        "CASE-LIST-B2 一级项目符号 B",
    ]
    assert all(paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None for paragraph in paragraphs)


def test_renderer_preserves_pandoc_native_subscript_and_superscript(tmp_path) -> None:
    output_path = tmp_path / "subscript-superscript.docx"

    render_docx(
        markdown="CASE-RUN-ANCHOR H~2~O；x^2^\n",
        style_pack=create_style_pack(name="上下角标测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    with zipfile.ZipFile(output_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "CASE-RUN-ANCHOR" in document_xml
    assert 'w:vertAlign w:val="subscript"' in document_xml
    assert 'w:vertAlign w:val="superscript"' in document_xml


def test_renderer_preserves_tabs_and_internal_page_break_extension(tmp_path) -> None:
    output_path = tmp_path / "tabs-page-break.docx"

    render_docx(
        markdown=(
            "CASE-SPACE-001 A&#9;B\n\n"
            "Before page break\n\n"
            "<!-- doc-converter:page-break -->\n\n"
            "After page break\n"
        ),
        style_pack=create_style_pack(name="分页与制表符测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    assert any("A\tB" in paragraph.text for paragraph in document.paragraphs)
    with zipfile.ZipFile(output_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:br w:type="page"' in document_xml


def test_renderer_keeps_multiline_simple_table_cell_with_native_entity(tmp_path) -> None:
    output_path = tmp_path / "multiline-table.docx"

    render_docx(
        markdown="| A | B |\n| --- | --- |\n| T3 | 第一行&#10;第二行&#10;第三行 |\n",
        style_pack=create_style_pack(name="多行表格测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    cell_text = Document(output_path).tables[0].cell(1, 1).text
    assert "第一行" in cell_text
    assert "第二行" in cell_text
    assert "第三行" in cell_text
    assert "\n" in cell_text


def test_renderer_preserves_complex_table_internal_extension(tmp_path) -> None:
    output_path = tmp_path / "complex-table.docx"
    markdown = """表2 复杂表格

<table>
<tr><th>A</th><th>B</th><th>C</th><th>D</th></tr>
<tr><td colspan="2">CASE-TBL-MERGE-H 水平合并 A+B</td><td>C1</td><td>D1</td></tr>
<tr><td rowspan="2">CASE-TBL-MERGE-V 纵向合并两行</td><td>B2</td><td>C2</td><td>D2</td></tr>
<tr><td>B3</td><td>C3</td><td>D3</td></tr>
<tr><td>A4</td><td>CASE-TBL-NESTED 嵌套表：<table><tr><td>N1</td><td>N2</td></tr><tr><td>N3</td><td>N4</td></tr></table></td><td>C4</td><td>CASE-TBL-BLANK 空单元格</td></tr>
</table>
"""

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="复杂表格测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(1, 0).text.startswith("CASE-TBL-MERGE-H")
    assert table.cell(1, 0)._tc.xpath("./w:tcPr/w:gridSpan")
    assert table.cell(2, 0)._tc.xpath("./w:tcPr/w:vMerge")
    assert any("CASE-TBL-BLANK" in cell.text for row in table.rows for cell in row.cells)
    nested_tables = [nested for row in table.rows for cell in row.cells for nested in cell.tables]
    assert len(nested_tables) == 1
    assert nested_tables[0].cell(0, 0).text == "N1"
    assert nested_tables[0].cell(1, 1).text == "N4"


def test_renderer_preserves_complex_table_inline_content_and_nested_page_break(tmp_path) -> None:
    image_path = tmp_path / "cell.png"
    _write_png(image_path, 40, 20)
    output_path = tmp_path / "complex-table-inline.docx"
    markdown = f"""<table>
<tr><td rowspan="2"><strong>粗体</strong> <a href="https://example.com">OpenAI</a> <img src="{image_path.name}" alt="单元格图">A<!-- doc-converter:page-break -->B</td><td>右侧</td></tr>
<tr><td>第二行</td></tr>
</table>
"""

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="复杂表格内联内容", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    cell = document.tables[0].cell(0, 0)
    assert "粗体" in cell.text
    assert "OpenAI" in cell.text
    assert "A" in cell.text and "B" in cell.text
    assert cell._tc.xpath(".//w:drawing")
    assert any(run.bold for paragraph in cell.paragraphs for run in paragraph.runs if "粗体" in run.text)
    with zipfile.ZipFile(output_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        relationships_xml = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert 'w:br w:type="page"' in document_xml
    assert "https://example.com" in relationships_xml


def test_renderer_restores_internal_endnote_as_endnote_part(tmp_path) -> None:
    output_path = tmp_path / "endnote.docx"
    markdown = (
        "脚注文本[^1]，尾注文本[^doc-converter-endnote-1]。\n\n"
        "[^1]: CASE-FN-TEXT-001\n"
        "[^doc-converter-endnote-1]: CASE-EN-TEXT-001\n"
    )

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="尾注测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    with zipfile.ZipFile(output_path) as archive:
        assert "word/endnotes.xml" in archive.namelist()
        endnotes_xml = archive.read("word/endnotes.xml").decode("utf-8")
        footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "CASE-EN-TEXT-001" in endnotes_xml
    assert "CASE-EN-TEXT-001" not in footnotes_xml
    assert "w:endnoteReference" in document_xml
    assert "w:footnoteRef" not in endnotes_xml
    assert "FootnoteReference" not in endnotes_xml
    assert "FootnoteText" not in endnotes_xml


def test_renderer_extracts_title_area_blocks_from_raw_markdown(tmp_path) -> None:
    markdown = """---
title: 原始 Markdown 标题
---

基于本地样本的实施说明

摘要：这是摘要内容。

关键词：确定性转换；DOCX；A4

# 一、正文

正文内容。
"""
    output_path = tmp_path / "raw.docx"

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="原始输入测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "基于本地样本的实施说明" in text
    assert "摘要：这是摘要内容。" in text
    assert "关键词：确定性转换；DOCX；A4" in text
    assert "DC_DocumentSubtitle" in {paragraph.style.style_id for paragraph in document.paragraphs}


def test_renderer_converts_raw_html_table_and_styles_caption_and_metadata_labels(tmp_path) -> None:
    markdown = """---
title: AI模型助力提升物资采购质量
---

何宏山

重庆市公共资源交易中心

采购管理部

摘要：本文说明采购质量提升路径。

关键词：采购；质量；模型

表1 选择AI编写工具时应考虑的几个因素

<table><tr><td>序号</td><td>因素</td><td>描述</td></tr><tr><td>1</td><td>功能性</td><td>工具是否具备自动生成模板、智能检查等功能</td></tr></table>

# 一、正文

正文内容。
"""
    output_path = tmp_path / "metadata-table.docx"

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="元数据表格测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "序号"
    assert table.cell(1, 2).text == "工具是否具备自动生成模板、智能检查等功能"
    caption = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("表1"))
    assert caption.style.style_id == "DC_TableCaption"
    assert {paragraph.style.style_id for paragraph in document.paragraphs} >= {
        "DC_DocumentTitle",
        "DC_Author",
        "DC_OrganizationDepartment",
        "DC_AbstractLabel",
        "DC_KeywordsLabel",
    }
    assert any(run.style and run.style.style_id == "DC_AbstractContent" for paragraph in document.paragraphs for run in paragraph.runs)
    assert any(run.style and run.style.style_id == "DC_KeywordsContent" for paragraph in document.paragraphs for run in paragraph.runs)
    assert any(paragraph.style.style_id == "DC_AbstractLabel" for paragraph in document.paragraphs)
    assert any(paragraph.style.style_id == "DC_KeywordsLabel" for paragraph in document.paragraphs)
    assert document.styles["DC_AbstractLabel"].type == 1  # paragraph style
    assert document.styles["DC_AbstractContent"].type == 2  # character style
    assert document.styles["DC_AbstractLabel"].font.name == "Times New Roman"
    assert document.styles["DC_AbstractContent"].font.name == "Times New Roman"
    assert table.cell(0, 0).paragraphs[0].style.style_id == "DC_TableColumnHeader"
    assert table.cell(1, 0).paragraphs[0].style.style_id == "DC_TableRowHeader"
    assert table.cell(1, 1).paragraphs[0].style.style_id == "DC_TableBody"


def test_renderer_applies_separate_first_row_and_first_column_table_roles(tmp_path) -> None:
    output_path = tmp_path / "table-heading-roles.docx"
    render_docx(
        markdown="| 项目 | 数量 |\n| --- | --- |\n| 甲项 | 3 |\n| 乙项 | 5 |\n",
        style_pack=create_style_pack(name="表格标题角色", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    table = Document(output_path).tables[0]
    assert table.cell(0, 0).paragraphs[0].style.style_id == "DC_TableColumnHeader"
    assert table.cell(0, 1).paragraphs[0].style.style_id == "DC_TableColumnHeader"
    assert table.cell(1, 0).paragraphs[0].style.style_id == "DC_TableRowHeader"
    assert table.cell(2, 0).paragraphs[0].style.style_id == "DC_TableRowHeader"
    assert table.cell(1, 1).paragraphs[0].style.style_id == "DC_TableBody"
    assert table.cell(2, 1).paragraphs[0].style.style_id == "DC_TableBody"


def test_renderer_recognizes_image_syntax_variants_and_applies_image_style(tmp_path) -> None:
    work_dir = tmp_path / "work"
    work_dir.mkdir()
    _write_png(tmp_path / "wide.png", 1200, 600)
    _write_png(tmp_path / "small.png", 80, 40)
    markdown = """![inline](wide.png)

![reference][wide-image]

[wide-image]: wide.png

<IMG src="wide.png" alt="HTML image" />

![split line]
(wide.png)

![small image](small.png)

图1 现场会议图片
"""
    base_pack = create_style_pack(name="图片样式", format_id="docx")
    roles = dict(base_pack.roles)
    roles["image"] = {
        "role_type": "image",
        "image": {
            "alignment": "right",
            "max_width_cm": 6,
            "space_before_pt": 2.5,
            "space_after_pt": 7.5,
        },
    }
    output_path = tmp_path / "images.docx"
    render_docx(
        markdown=markdown,
        style_pack=base_pack.model_copy(update={"roles": roles}),
        output_path=output_path,
        work_dir=work_dir,
    )

    document = Document(output_path)
    image_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph._p.xpath(".//w:drawing")]
    shapes = [document.inline_shapes[index] for index in range(len(document.inline_shapes))]
    widths_cm = [shape.width.cm for shape in shapes]
    assert len(image_paragraphs) == 5
    assert all(paragraph.style.style_id == "DC_Image" for paragraph in image_paragraphs)
    assert all(width <= 6.01 for width in widths_cm)
    assert widths_cm[-1] < 6
    assert all(abs(shape.width.cm / shape.height.cm - 2) < 0.02 for shape in shapes[:-1])
    assert document.styles["DC_Image"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert document.styles["DC_Image"].paragraph_format.space_before.pt == 2.5
    assert document.styles["DC_Image"].paragraph_format.space_after.pt == 7.5
    caption = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("图1"))
    assert caption.style.style_id == "DC_FigureCaption"
    for paragraph in image_paragraphs:
        outer_extent = paragraph._p.xpath(".//wp:extent")[0]
        picture_extent = paragraph._p.xpath(".//a:xfrm/a:ext")[0]
        assert outer_extent.get("cx") == picture_extent.get("cx")
        assert outer_extent.get("cy") == picture_extent.get("cy")


def test_renderer_defaults_image_width_to_body_width_without_upscaling(tmp_path) -> None:
    work_dir = tmp_path / "work"
    work_dir.mkdir()
    _write_png(tmp_path / "wide.png", 1200, 600)
    _write_png(tmp_path / "small.png", 80, 40)
    output_path = tmp_path / "default-image-width.docx"
    render_docx(
        markdown="![wide](wide.png)\n\n![small](small.png)\n",
        style_pack=create_style_pack(name="默认图片宽度", format_id="docx"),
        output_path=output_path,
        work_dir=work_dir,
    )

    document = Document(output_path)
    shapes = [document.inline_shapes[index] for index in range(len(document.inline_shapes))]
    assert len(shapes) == 2
    section = document.sections[0]
    body_width_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    assert abs(shapes[0].width.cm - body_width_cm) < 0.05
    assert shapes[1].width.cm < shapes[0].width.cm


def test_renderer_prefers_bundled_pandoc_when_available(tmp_path) -> None:
    bundled_root = tmp_path / "pandoc"
    bundled_root.mkdir()
    bundled_executable = bundled_root / ("pandoc.exe" if os.name == "nt" else "pandoc")
    bundled_executable.write_bytes(b"not-executed-in-this-test")

    assert resolve_pandoc(bundled_root=bundled_root) == bundled_executable


def test_renderer_writes_self_contained_dc_styles_metadata_and_plain_numbering(tmp_path) -> None:
    markdown = """---
schema: normalized-md/v1
title: 项目建设方案
subtitle: 实施说明
organization: XX公司
department: XX部门
authors:
  - 张三
  - 李四
date: 2026-09-18
abstract: 在安全、制度和生态四个方面推进智慧交易。
keywords:
  - 人工智能
  - 大模型
  - 公共资源交易
  - 数智化
---

# 一、建设背景

## 二、工作内容

正文包含 **重点内容**。

1. 第一项
2. 第二项

| 项目 | 数量 |
| --- | ---: |
| A | 10 |
"""
    output_path = tmp_path / "项目建设方案.docx"

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="测试样式", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    style_ids = {paragraph.style.style_id for paragraph in document.paragraphs}
    assert {
        "DC_DocumentTitle",
        "DC_DocumentSubtitle",
        "DC_OrganizationDepartment",
        "DC_Author",
        "DC_Date",
        "DC_AbstractLabel",
        "DC_KeywordsLabel",
    } <= style_ids
    assert any(run.style and run.style.style_id == "DC_AbstractContent" for paragraph in document.paragraphs for run in paragraph.runs)
    assert any(run.style and run.style.style_id == "DC_KeywordsContent" for paragraph in document.paragraphs for run in paragraph.runs)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "摘要：在安全、制度和生态四个方面推进智慧交易。" in text
    assert "关键词：人工智能；大模型；公共资源交易；数智化" in text
    assert round(document.sections[0].page_width.mm) == 210
    assert round(document.sections[0].page_height.mm) == 297
    assert {"DC_Heading1", "DC_Heading2"} <= style_ids
    assert next(paragraph for paragraph in document.paragraphs if paragraph.text == "一、建设背景").style.style_id == "DC_Heading1"
    assert next(paragraph for paragraph in document.paragraphs if paragraph.text == "二、工作内容").style.style_id == "DC_Heading2"
    table_style_ids = {
        paragraph.style.style_id
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    }
    assert "DC_TableBody" in table_style_ids
    assert all(paragraph.style.style_id != "Heading 2" for paragraph in document.paragraphs)
    list_items = [paragraph for paragraph in document.paragraphs if paragraph.text in {"第一项", "第二项"}]
    assert len(list_items) == 2
    assert all(paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None for paragraph in list_items)
    assert any(run.style and run.style.style_id == "DC_Point" for paragraph in document.paragraphs for run in paragraph.runs)

    with zipfile.ZipFile(output_path) as archive:
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
    for style_id in (
        "DC_DocumentTitle",
        "DC_DocumentSubtitle",
        "DC_OrganizationDepartment",
        "DC_Author",
        "DC_Date",
        "DC_AbstractLabel",
        "DC_AbstractContent",
        "DC_KeywordsLabel",
        "DC_KeywordsContent",
        "DC_Heading1",
        "DC_Heading2",
        "DC_Heading3",
        "DC_Heading4",
        "DC_Heading5",
        "DC_Heading6",
        "DC_BodyParagraph",
        "DC_TableBody",
        "DC_TableRowHeader",
        "DC_TableColumnHeader",
        "DC_Image",
        "DC_Point",
        "DC_ReferenceHeading",
        "DC_ReferenceEntry",
        "DC_Footer",
    ):
        assert f'w:styleId="{style_id}"' in styles_xml
    assert 'w:eastAsia="仿宋"' in styles_xml
    assert 'w:ascii="Times New Roman"' in styles_xml
    assert 'w:val="DC_Point"' in document_xml
    for level in range(1, 7):
        style = document.styles[f"DC_Heading{level}"]
        ppr = style._element.find(qn("w:pPr"))
        assert ppr is not None
        outline = ppr.find(qn("w:outlineLvl"))
        assert outline is not None
        assert outline.get(qn("w:val")) == str(level - 1)


@pytest.mark.parametrize(
    "value",
    ["参考文献：", "（一）参考文献", "（1）参考文献", "参考文献（References）"],
)
def test_reference_heading_variants_are_recognized(value: str) -> None:
    assert _REFERENCE_HEADING_PATTERN.fullmatch(value) is not None


def test_renderer_starts_post_document_styles_only_at_reference_heading(tmp_path) -> None:
    markdown = """# 正文标题

正文段落。

六、参考文献

[1] 参考资料条目。
"""
    output_path = tmp_path / "references.docx"

    render_docx(
        markdown=markdown,
        style_pack=create_style_pack(name="参考文献测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    paragraphs = {paragraph.text: paragraph.style.style_id for paragraph in document.paragraphs}
    assert paragraphs["正文段落。"] == "DC_BodyParagraph"
    assert paragraphs["六、参考文献"] == "DC_ReferenceHeading"
    assert paragraphs["[1] 参考资料条目。"] == "DC_ReferenceEntry"


def test_renderer_without_reference_heading_has_no_post_document_body_style(tmp_path) -> None:
    output_path = tmp_path / "no-references.docx"
    render_docx(
        markdown="# 正文标题\n\n普通正文。\n",
        style_pack=create_style_pack(name="无参考文献测试", format_id="docx"),
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    document = Document(output_path)
    assert all(
        paragraph.style.style_id not in {"DC_ReferenceHeading", "DC_ReferenceEntry"}
        for paragraph in document.paragraphs
    )


def test_renderer_keeps_selected_chinese_and_western_fonts_in_ooxml(tmp_path) -> None:
    pack = create_style_pack(name="黑体测试", format_id="docx")
    roles = {key: dict(value) for key, value in pack.roles.items()}
    roles["body-paragraph"] = {
        **roles["body-paragraph"],
        "font": {**roles["body-paragraph"]["font"], "chinese_font": "黑体", "western_font": "黑体"},
    }
    pack = pack.model_copy(update={"roles": roles})

    output_path = tmp_path / "font.docx"
    render_docx(
        markdown="正文内容。",
        style_pack=pack,
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    with zipfile.ZipFile(output_path) as archive:
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
    assert 'w:styleId="DC_BodyParagraph"' in styles_xml
    body_style = styles_xml[styles_xml.index('w:styleId="DC_BodyParagraph"') :]
    assert 'w:eastAsia="黑体"' in body_style
    assert 'w:ascii="黑体"' in body_style
    assert 'w:hAnsi="黑体"' in body_style
    assert 'w:cs="黑体"' in body_style


@pytest.mark.parametrize(
    ("rule", "expected_line", "expected_rule"),
    [
        ("maximum", "580", "auto"),
        ("minimum", "580", "atLeast"),
        ("exact", "580", "exact"),
        ("multiple", "360", "auto"),
    ],
)
def test_renderer_writes_line_spacing_rule_and_value_to_ooxml(tmp_path, rule, expected_line, expected_rule) -> None:
    base = create_style_pack(name="行距测试", format_id="docx")
    roles = {key: dict(value) for key, value in base.roles.items()}
    roles["body-paragraph"] = {
        **roles["body-paragraph"],
        "paragraph": {
            **roles["body-paragraph"]["paragraph"],
            "line_spacing_rule": rule,
            "line_spacing_multiple": 1.5 if rule == "multiple" else None,
        },
    }
    pack = base.model_copy(update={"roles": roles})
    output_path = tmp_path / "line-spacing.docx"

    render_docx(
        markdown="正文内容。",
        style_pack=pack,
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    with zipfile.ZipFile(output_path) as archive:
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
    body_style = styles_xml[styles_xml.index('w:styleId="DC_BodyParagraph"') :]
    assert f'w:line="{expected_line}"' in body_style
    assert f'w:lineRule="{expected_rule}"' in body_style


def test_renderer_writes_character_indent_attributes_to_ooxml(tmp_path) -> None:
    base = create_style_pack(name="字符缩进测试", format_id="docx")
    roles = {key: dict(value) for key, value in base.roles.items()}
    roles["body-paragraph"] = {
        **roles["body-paragraph"],
        "paragraph": {
            **roles["body-paragraph"]["paragraph"],
            "first_line_indent_pt": 32,
            "first_line_indent_unit": "character",
            "left_indent_pt": 16,
            "left_indent_unit": "character",
        },
    }
    pack = base.model_copy(update={"roles": roles})
    output_path = tmp_path / "character-indent.docx"

    render_docx(
        markdown="正文内容。",
        style_pack=pack,
        output_path=output_path,
        work_dir=tmp_path / "work",
    )

    with zipfile.ZipFile(output_path) as archive:
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
    body_style = styles_xml[styles_xml.index('w:styleId="DC_BodyParagraph"') :]
    assert 'w:firstLineChars="200"' in body_style
    assert 'w:leftChars="100"' in body_style
    assert 'w:firstLine="640"' in body_style
    assert 'w:left="320"' in body_style


def _write_png(path, width: int, height: int) -> None:
    def chunk(name: bytes, payload: bytes) -> bytes:
        checksum = zlib.crc32(name + payload) & 0xFFFFFFFF
        return struct.pack(">I", len(payload)) + name + payload + struct.pack(">I", checksum)

    scanline = b"\x00" + b"\x37\x88\xb8" * width
    image_data = zlib.compress(scanline * height)
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", image_data)
        + chunk(b"IEND", b"")
    )

