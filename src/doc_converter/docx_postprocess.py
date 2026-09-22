"""Deterministic OOXML cleanup after Pandoc has produced a DOCX."""

from __future__ import annotations

import re
from typing import Any

from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from doc_converter.docx_styles import ROLE_TO_STYLE_ID
from doc_converter.normalizer import (
    _TABLE_CAPTION_PATTERN,
    _extract_front_matter,
    _extract_leading_metadata,
    _join_split_image_syntax,
    _normalize_html_images,
    _normalize_html_sub_sup,
    _normalize_html_tables,
    _normalize_metadata,
)

_ORDERED_LIST = re.compile(r"^(\s*)(\d+)\.(\s+)")
_TABLE_BORDER_EDGES = ("top", "bottom", "left", "right", "insideH", "insideV")
_HEADING_STYLE = re.compile(r"^Heading\s*([1-6])$", re.IGNORECASE)
_FIGURE_CAPTION_PATTERN = re.compile(
    r"^\s*(?:图\s*(?:\d+(?:[.-]\d+)*|[一二三四五六七八九十百千万零〇]+)|"
    r"fig(?:ure)?\.?\s*\d+(?:[.-]\d+)*)\s*(?:[：:、.．)）]\s*)?\S",
    re.IGNORECASE,
)
_REFERENCE_HEADING_PATTERN = re.compile(
    r"^(?:(?:\d+|[一二三四五六七八九十百千万]+)[、.．)\s]*|[（(]\s*(?:\d+|[一二三四五六七八九十百千万]+)\s*[）)]\s*)?"
    r"(?:参考文献|参考资料|references?)"
    r"(?:\s*[：:])?"
    r"(?:\s*[（(]\s*references?\s*[）)])?\s*$",
    re.IGNORECASE,
)


def split_metadata(markdown: str) -> tuple[dict[str, Any], str]:
    """Remove Front Matter before Pandoc and return normalized metadata."""

    raw_metadata, lines, first_content_line = _extract_front_matter(markdown)
    metadata = _normalize_metadata(raw_metadata)
    extracted, consumed, _ = _extract_leading_metadata(lines, metadata, first_content_line)
    metadata.update(extracted)
    body_lines = [line for index, line in enumerate(lines) if index not in consumed]
    body = _join_split_image_syntax("\n".join(body_lines).strip("\n"))
    body = _normalize_html_sub_sup(_normalize_html_tables(_normalize_html_images(body)))
    return metadata, body + "\n"


def protect_ordered_list_literals(markdown: str) -> str:
    """Escape ordered-list markers so Pandoc emits ordinary paragraphs."""

    lines: list[str] = []
    in_fence = False
    for line in markdown.splitlines():
        if line.strip().startswith("```") or line.strip().startswith("~~~"):
            in_fence = not in_fence
        if not in_fence:
            line = _ORDERED_LIST.sub(r"\1\2\\.\3", line)
        lines.append(line)
    return "\n".join(lines) + "\n"


def apply_postprocess(
    document: DocxDocument,
    metadata: dict[str, Any],
    *,
    max_image_width_cm: float | None = None,
) -> None:
    """Apply semantic style IDs, metadata paragraphs, lists and footers."""

    in_references = False
    for paragraph in document.paragraphs:
        if _REFERENCE_HEADING_PATTERN.fullmatch(paragraph.text.strip()):
            paragraph.style = ROLE_TO_STYLE_ID["reference-heading"]
            _remove_numbering(paragraph)
            in_references = True
            continue
        _map_paragraph(
            paragraph,
            ROLE_TO_STYLE_ID["reference-entry"] if in_references else ROLE_TO_STYLE_ID["body-paragraph"],
        )
    for table in document.tables:
        _apply_table_borders(table)
        _style_table_cells(table)

    _style_table_captions(document)
    _style_figure_captions(document)
    _apply_image_max_width(document, max_image_width_cm)
    _insert_metadata(document, metadata)
    _add_page_footers(document)


def _style_table_cells(table: Any) -> None:
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            if row_index == 0:
                default_style = ROLE_TO_STYLE_ID["table-column-header"]
            elif column_index == 0:
                default_style = ROLE_TO_STYLE_ID["table-row-header"]
            else:
                default_style = ROLE_TO_STYLE_ID["body-table"]
            for paragraph in cell.paragraphs:
                _map_paragraph(
                    paragraph,
                    default_style,
                    map_bold_to_point=default_style == ROLE_TO_STYLE_ID["body-table"],
                )
            for nested_table in cell.tables:
                _apply_table_borders(nested_table)
                _style_table_cells(nested_table)


def _apply_table_borders(table: Any) -> None:
    """Give every generated table a consistent half-point grid."""

    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.insert_element_before(
            borders,
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
        )
    for edge_name in _TABLE_BORDER_EDGES:
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), "000000")


def _map_paragraph(
    paragraph: Paragraph,
    default_style: str,
    *,
    map_bold_to_point: bool = True,
) -> None:
    style_id = paragraph.style.style_id if paragraph.style is not None else ""
    had_numbering = _has_numbering(paragraph)
    heading_match = _HEADING_STYLE.match(style_id)
    if paragraph._p.xpath(".//w:drawing"):
        paragraph.style = ROLE_TO_STYLE_ID["image"]
    elif heading_match:
        paragraph.style = ROLE_TO_STYLE_ID[f"heading-{heading_match.group(1)}"]
    elif style_id.lower().startswith("heading"):
        # Pandoc occasionally emits ``Heading1`` without a space.  Preserve
        # its level when possible, otherwise use the first-level body heading.
        level_match = re.search(r"([1-6])$", style_id)
        level = level_match.group(1) if level_match else "1"
        paragraph.style = ROLE_TO_STYLE_ID[f"heading-{level}"]
    elif style_id not in ROLE_TO_STYLE_ID.values():
        paragraph.style = default_style
    if not had_numbering:
        _remove_numbering(paragraph)
    for run in paragraph.runs:
        if run.bold is True and map_bold_to_point:
            run.style = ROLE_TO_STYLE_ID["point"]
            _remove_direct_bold(run)


def _apply_image_max_width(document: DocxDocument, max_image_width_cm: float | None) -> None:
    """Fit pictures to the body width or a smaller custom width without upscaling."""

    section = document.sections[0]
    page_width = section.page_width
    left_margin = section.left_margin
    right_margin = section.right_margin
    if page_width is None or left_margin is None or right_margin is None:
        return
    body_width_emu = int(page_width) - int(left_margin) - int(right_margin)
    if body_width_emu <= 0:
        return
    width_limit_emu = body_width_emu
    if isinstance(max_image_width_cm, (int, float)) and max_image_width_cm > 0:
        width_limit_emu = min(width_limit_emu, int(Cm(float(max_image_width_cm))))

    for paragraph in _all_paragraphs(document):
        for drawing in paragraph._p.xpath(".//wp:inline | .//wp:anchor"):
            for extent in drawing.xpath("./wp:extent"):
                width = int(extent.get("cx", "0"))
                height = int(extent.get("cy", "0"))
                if width <= width_limit_emu or width <= 0 or height <= 0:
                    continue
                scale = width_limit_emu / width
                resized_width = str(width_limit_emu)
                resized_height = str(max(1, round(height * scale)))
                extent.set("cx", resized_width)
                extent.set("cy", resized_height)
                for picture_extent in drawing.xpath(".//a:xfrm/a:ext"):
                    picture_extent.set("cx", resized_width)
                    picture_extent.set("cy", resized_height)


def _all_paragraphs(document: DocxDocument) -> list[Paragraph]:
    paragraphs = list(document.paragraphs)

    def add_table_paragraphs(table: Any) -> None:
        for row in table.rows:
            visited_cells: set[int] = set()
            for cell in row.cells:
                cell_identity = id(cell._tc)
                if cell_identity in visited_cells:
                    continue
                visited_cells.add(cell_identity)
                paragraphs.extend(cell.paragraphs)
                for nested_table in cell.tables:
                    add_table_paragraphs(nested_table)

    for table in document.tables:
        add_table_paragraphs(table)
    return paragraphs


def _style_table_captions(document: DocxDocument) -> None:
    body = document._body._element
    children = list(body)
    for index, child in enumerate(children):
        if child.tag != qn("w:tbl") or index == 0:
            continue
        previous = children[index - 1]
        if previous.tag != qn("w:p"):
            continue
        paragraph = Paragraph(previous, document)
        if _TABLE_CAPTION_PATTERN.match(paragraph.text):
            paragraph.style = ROLE_TO_STYLE_ID["table-caption"]


def _style_figure_captions(document: DocxDocument) -> None:
    for paragraph in document.paragraphs:
        if _FIGURE_CAPTION_PATTERN.match(paragraph.text):
            paragraph.style = ROLE_TO_STYLE_ID["figure-caption"]


def _remove_numbering(paragraph: Paragraph) -> None:
    ppr = paragraph._p.pPr
    if ppr is None:
        return
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def _has_numbering(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.find(qn("w:numPr")) is not None


def _remove_direct_bold(run: Any) -> None:
    rpr = run._r.rPr
    if rpr is None:
        return
    for tag in ("w:b", "w:bCs"):
        element = rpr.find(qn(tag))
        if element is not None:
            rpr.remove(element)


def _insert_metadata(document: DocxDocument, metadata: dict[str, Any]) -> None:
    entries: list[tuple[str, str, str]] = []
    if isinstance(metadata.get("title"), str) and metadata["title"].strip():
        entries.append(("document-title", metadata["title"], ROLE_TO_STYLE_ID["document-title"]))
    if isinstance(metadata.get("subtitle"), str) and metadata["subtitle"].strip():
        entries.append(("document-subtitle", metadata["subtitle"], ROLE_TO_STYLE_ID["document-subtitle"]))
    if isinstance(metadata.get("organization"), str) and metadata["organization"].strip():
        entries.append(("organization", metadata["organization"], ROLE_TO_STYLE_ID["organization-department"]))
    if isinstance(metadata.get("department"), str) and metadata["department"].strip():
        entries.append(("department", metadata["department"], ROLE_TO_STYLE_ID["organization-department"]))
    authors = metadata.get("authors")
    if isinstance(authors, list) and authors:
        entries.append(("authors", "、".join(str(author) for author in authors), ROLE_TO_STYLE_ID["author"]))
    if isinstance(metadata.get("date"), str) and metadata["date"].strip():
        entries.append(("date", metadata["date"], ROLE_TO_STYLE_ID["date"]))
    abstract = _strip_metadata_label(metadata.get("abstract"), "摘要")
    keywords = _keyword_text(metadata.get("keywords"))
    if not entries and not abstract and not keywords:
        return

    body = document._body._element
    first = next((child for child in body if child.tag in {qn("w:p"), qn("w:tbl")}), None)
    insert_index = list(body).index(first) if first is not None else len(body)
    for _, text, style_id in entries:
        element = OxmlElement("w:p")
        body.insert(insert_index, element)
        insert_index += 1
        paragraph = Paragraph(element, document)
        paragraph.style = style_id
        paragraph.add_run(text)
    if abstract:
        _insert_labeled_metadata_paragraph(
            document,
            body,
            insert_index,
            label="摘要：",
            content=abstract,
            paragraph_style=ROLE_TO_STYLE_ID["abstract-label"],
            content_style=ROLE_TO_STYLE_ID["abstract-content"],
        )
        insert_index += 1
    if keywords:
        _insert_labeled_metadata_paragraph(
            document,
            body,
            insert_index,
            label="关键词：",
            content=keywords,
            paragraph_style=ROLE_TO_STYLE_ID["keywords-label"],
            content_style=ROLE_TO_STYLE_ID["keywords-content"],
        )


def _insert_labeled_metadata_paragraph(
    document: DocxDocument,
    body: Any,
    insert_index: int,
    *,
    label: str,
    content: str,
    paragraph_style: str,
    content_style: str,
) -> None:
    element = OxmlElement("w:p")
    body.insert(insert_index, element)
    paragraph = Paragraph(element, document)
    paragraph.style = paragraph_style
    paragraph.add_run(label)
    content_run = paragraph.add_run(content)
    content_run.style = content_style


def _strip_metadata_label(value: Any, label: str) -> str:
    if not isinstance(value, str):
        return ""
    text = value.strip()
    return re.sub(rf"^{re.escape(label)}\s*[:：]\s*", "", text).strip()


def _keyword_text(value: Any) -> str:
    if isinstance(value, list):
        items = [str(item).strip() for item in value if str(item).strip()]
        return "；".join(items)
    if isinstance(value, str):
        return _strip_metadata_label(value, "关键词")
    return ""


def _add_page_footers(document: DocxDocument) -> None:
    settings = document.settings._element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        settings.append(OxmlElement("w:evenAndOddHeaders"))
    for section in document.sections:
        _set_footer(section.footer, "right")
        even_footer = getattr(section, "even_page_footer", None)
        if even_footer is not None:
            _set_footer(even_footer, "left")


def _set_footer(footer: Any, alignment: str) -> None:
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.style = ROLE_TO_STYLE_ID["footer"]
    paragraph.alignment = {"left": 0, "right": 2}[alignment]
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run("－ ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    paragraph.add_run(" －")
