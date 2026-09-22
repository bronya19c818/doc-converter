"""Internal Markdown extensions for constructs without a Pandoc Markdown form."""

from __future__ import annotations

import re
import tempfile
from copy import deepcopy
from dataclasses import dataclass, field
from html import escape
from html.parser import HTMLParser
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree  # type: ignore[import-untyped]

PAGE_BREAK_MARKER = "DOC_CONVERTER_PAGE_BREAK_EXTENSION"
COMPLEX_TABLE_MARKER = "DOC_CONVERTER_COMPLEX_TABLE_EXTENSION"
_TABLE_TAG_PATTERN = re.compile(r"<\s*(/?)\s*table\b[^>]*>", re.IGNORECASE)
_COMPLEX_TABLE_ATTRIBUTE_PATTERN = re.compile(r"\b(?:rowspan|colspan)\s*=", re.IGNORECASE)
_FOOTNOTE_REFERENCE_PATTERN = re.compile(r"\[\^([^\]]+)\](?!:)")

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS = {"w": _W_NS, "r": _R_NS, "ct": _CT_NS, "rel": _REL_NS}
_NOTE_STYLE_REPLACEMENTS = {
    "FootnoteReference": "EndnoteReference",
    "FootnoteText": "EndnoteText",
}


@dataclass
class HtmlTableCell:
    text: str = ""
    markup: str = ""
    rowspan: int = 1
    colspan: int = 1
    nested_tables: list["HtmlTableSpec"] = field(default_factory=list)


@dataclass
class HtmlTableSpec:
    rows: list[list[HtmlTableCell]] = field(default_factory=list)


@dataclass(frozen=True)
class _TablePlacement:
    row: int
    column: int
    rowspan: int
    colspan: int
    cell: HtmlTableCell


class _HtmlTableSpecParser(HTMLParser):
    """Read a bounded HTML table block, including one or more nested tables."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[HtmlTableSpec] = []
        self._stack: list[HtmlTableSpec] = []
        self._rows: list[list[HtmlTableCell] | None] = []
        self._cells: list[HtmlTableCell | None] = []
        self._buffers: list[list[str] | None] = []
        self._markup_buffers: list[list[str] | None] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.casefold()
        if name == "table":
            table = HtmlTableSpec()
            if self._stack and self._cells[-1] is not None:
                self._cells[-1].nested_tables.append(table)
            else:
                self.tables.append(table)
            self._stack.append(table)
            self._rows.append(None)
            self._cells.append(None)
            self._buffers.append(None)
            self._markup_buffers.append(None)
            return
        if not self._stack:
            return
        if name == "tr":
            row: list[HtmlTableCell] = []
            self._stack[-1].rows.append(row)
            self._rows[-1] = row
            return
        if name in {"td", "th"} and self._rows[-1] is not None:
            attributes = {key.casefold(): value for key, value in attrs}
            cell = HtmlTableCell(
                rowspan=_positive_int(attributes.get("rowspan")),
                colspan=_positive_int(attributes.get("colspan")),
            )
            self._rows[-1].append(cell)
            self._cells[-1] = cell
            self._buffers[-1] = []
            self._markup_buffers[-1] = []
            return
        if self._buffers[-1] is None:
            return
        markup_buffer = self._markup_buffers[-1]
        if markup_buffer is None:
            return
        if name == "br":
            self._buffers[-1].append("\ue000")
            markup_buffer.append("<br>")
        elif name not in {"tbody", "thead", "tfoot", "tr", "colgroup", "col"}:
            markup_buffer.append(self.get_starttag_text() or f"<{name}>")

    def handle_endtag(self, tag: str) -> None:
        name = tag.casefold()
        if not self._stack:
            return
        if name in {"td", "th"} and self._cells[-1] is not None and self._buffers[-1] is not None:
            self._cells[-1].text = _normalize_cell_text("".join(self._buffers[-1]))
            self._cells[-1].markup = "".join(self._markup_buffers[-1] or [])
            self._cells[-1] = None
            self._buffers[-1] = None
            self._markup_buffers[-1] = None
            return
        if name == "table":
            self._stack.pop()
            self._rows.pop()
            self._cells.pop()
            self._buffers.pop()
            self._markup_buffers.pop()
            return
        if self._markup_buffers[-1] is not None and name not in {
            "tbody",
            "thead",
            "tfoot",
            "tr",
            "colgroup",
            "col",
        }:
            self._markup_buffers[-1].append(f"</{name}>")

    def handle_data(self, data: str) -> None:
        if not self._buffers or self._buffers[-1] is None:
            return
        markup_buffer = self._markup_buffers[-1]
        if markup_buffer is None:
            return
        self._buffers[-1].append(data)
        markup_buffer.append(data)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.casefold()
        if name == "br":
            self.handle_starttag(tag, attrs)
            return
        if name == "img" and self._markup_buffers and self._markup_buffers[-1] is not None:
            self._markup_buffers[-1].append(self.get_starttag_text() or "<img>")
            return
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)


def _positive_int(value: str | None) -> int:
    try:
        return max(1, int(value or "1"))
    except ValueError:
        return 1


def _normalize_cell_text(value: str) -> str:
    lines = re.split("\ue000", value)
    normalized = [re.sub(r"\s+", " ", line).strip() for line in lines]
    return "\n".join(line for line in normalized if line)


def is_complex_html_table(block: str) -> bool:
    return (
        len(_TABLE_TAG_PATTERN.findall(block)) > 2
        or _COMPLEX_TABLE_ATTRIBUTE_PATTERN.search(block) is not None
    )


def extract_complex_tables(markdown: str) -> tuple[str, list[HtmlTableSpec]]:
    """Replace complex raw HTML tables with deterministic renderer placeholders."""

    tables: list[HtmlTableSpec] = []
    replacements: list[tuple[int, int, str]] = []
    search_from = 0
    while match := re.search(r"<\s*table\b", markdown[search_from:], re.IGNORECASE):
        start = search_from + match.start()
        depth = 0
        end = None
        for tag_match in _TABLE_TAG_PATTERN.finditer(markdown, start):
            if tag_match.group(1):
                depth -= 1
            else:
                depth += 1
            if depth == 0:
                end = tag_match.end()
                break
        if end is None:
            break
        block = markdown[start:end]
        search_from = end
        if not is_complex_html_table(block):
            continue
        parser = _HtmlTableSpecParser()
        parser.feed(block)
        parser.close()
        if not parser.tables:
            continue
        table_index = len(tables)
        tables.append(parser.tables[0])
        replacements.append(
            (start, end, f"{COMPLEX_TABLE_MARKER}_{table_index}")
        )
    for start, end, replacement in reversed(replacements):
        markdown = markdown[:start] + f"\n{replacement}\n" + markdown[end:]
    return markdown, tables


def extract_page_breaks(markdown: str) -> str:
    """Turn the invisible page-break extension into a Pandoc-visible marker."""

    return re.sub(
        r"<!--\s*doc-converter:page-break\s*-->",
        PAGE_BREAK_MARKER,
        markdown,
        flags=re.IGNORECASE,
    )


def _split_page_break_run(run_xml: etree._Element, text_node: etree._Element, text: str) -> None:
    """Split a run around every marker while retaining the original run formatting."""

    parent = run_xml.getparent()
    if parent is None:
        return
    segments = text.split(PAGE_BREAK_MARKER)
    text_node.text = segments[0]
    insertion_index = parent.index(run_xml) + 1
    for segment in segments[1:]:
        break_run = etree.Element(f"{{{_W_NS}}}r")
        run_properties = run_xml.find(f"{{{_W_NS}}}rPr")
        if run_properties is not None:
            break_run.append(deepcopy(run_properties))
        etree.SubElement(break_run, f"{{{_W_NS}}}br").set(
            f"{{{_W_NS}}}type",
            "page",
        )
        parent.insert(insertion_index, break_run)
        insertion_index += 1
        if not segment:
            continue
        text_run = deepcopy(run_xml)
        text_nodes = text_run.xpath("./w:t")
        if text_nodes:
            text_nodes[0].text = segment
            for extra in text_nodes[1:]:
                extra.text = ""
        parent.insert(insertion_index, text_run)
        insertion_index += 1


def restore_page_breaks(document: DocxDocument) -> None:
    """Restore every page-break marker in body and nested table paragraphs."""

    for paragraph_xml in document.element.body.xpath(".//w:p"):
        for run_xml in list(paragraph_xml.xpath("./w:r")):
            text_nodes = run_xml.xpath("./w:t")
            for text_node in text_nodes:
                text = text_node.text or ""
                if PAGE_BREAK_MARKER not in text:
                    continue
                _split_page_break_run(run_xml, text_node, text)


def restore_complex_tables(
    document: DocxDocument,
    tables: list[HtmlTableSpec],
    *,
    resource_root: Path | None = None,
) -> None:
    if not tables:
        return
    for index, specification in enumerate(tables):
        marker = f"{COMPLEX_TABLE_MARKER}_{index}"
        paragraph = next(
            (item for item in document.paragraphs if marker in item.text),
            None,
        )
        if paragraph is None:
            continue
        table = _append_table(document, specification, resource_root=resource_root)
        paragraph._p.addprevious(table._tbl)
        paragraph._p.getparent().remove(paragraph._p)


def _append_table(
    document: DocxDocument,
    specification: HtmlTableSpec,
    *,
    resource_root: Path | None = None,
) -> Table:
    placements, column_count = _table_placements(specification)
    table = document.add_table(rows=max(1, len(specification.rows)), cols=max(1, column_count))
    for placement in placements:
        if placement.rowspan == 1 and placement.colspan == 1:
            continue
        start = table.cell(placement.row, placement.column)
        end = table.cell(
            placement.row + placement.rowspan - 1,
            placement.column + placement.colspan - 1,
        )
        start.merge(end)
    for placement in placements:
        cell = table.cell(placement.row, placement.column)
        _fill_cell(cell, placement.cell, resource_root=resource_root)
    return table


def _table_placements(specification: HtmlTableSpec) -> tuple[list[_TablePlacement], int]:
    occupied: set[tuple[int, int]] = set()
    placements: list[_TablePlacement] = []
    column_count = 0
    for row_index, row in enumerate(specification.rows):
        column = 0
        for cell in row:
            while any((row_index, column + offset) in occupied for offset in range(cell.colspan)):
                column += 1
            placement = _TablePlacement(
                row=row_index,
                column=column,
                rowspan=cell.rowspan,
                colspan=cell.colspan,
                cell=cell,
            )
            placements.append(placement)
            for row_offset in range(placement.rowspan):
                for column_offset in range(placement.colspan):
                    occupied.add((row_index + row_offset, column + column_offset))
            column += placement.colspan
            column_count = max(column_count, column)
    return placements, column_count


def _fill_cell(
    cell: _Cell,
    specification: HtmlTableCell,
    *,
    resource_root: Path | None = None,
) -> None:
    _clear_cell_paragraph(cell)
    _InlineCellRenderer(
        cell.paragraphs[0],
        resource_root=resource_root,
    ).render(specification.markup or specification.text)
    for nested in specification.nested_tables:
        _append_nested_table(cell, nested, resource_root=resource_root)


def _append_nested_table(
    cell: _Cell,
    specification: HtmlTableSpec,
    *,
    resource_root: Path | None = None,
) -> None:
    nested = cell.add_table(rows=max(1, len(specification.rows)), cols=max(1, _table_placements(specification)[1]))
    placements, _ = _table_placements(specification)
    for placement in placements:
        if placement.rowspan == 1 and placement.colspan == 1:
            continue
        nested.cell(placement.row, placement.column).merge(
            nested.cell(
                placement.row + placement.rowspan - 1,
                placement.column + placement.colspan - 1,
            )
        )
    for placement in placements:
        _fill_cell(
            nested.cell(placement.row, placement.column),
            placement.cell,
            resource_root=resource_root,
        )


def _clear_cell_paragraph(cell: _Cell) -> None:
    paragraph = cell.paragraphs[0]
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


class _InlineCellRenderer(HTMLParser):
    """Render common HTML/Markdown inline constructs into an existing DOCX cell."""

    _MARKDOWN_LINK_PATTERN = re.compile(r"(!?)\[([^\]]*)\]\(([^)]+)\)")
    _STYLE_TAGS = {
        "strong": "bold",
        "b": "bold",
        "em": "italic",
        "i": "italic",
        "u": "underline",
        "del": "strike",
        "s": "strike",
        "sub": "subscript",
        "sup": "superscript",
        "code": "code",
    }

    def __init__(self, paragraph: Paragraph, *, resource_root: Path | None) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraph = paragraph
        self.resource_root = resource_root
        self._styles: dict[str, int] = {}
        self._links: list[str | None] = []

    def render(self, markup: str) -> None:
        markup = re.sub(
            r"!\[([^\]]*)\]\(<([^>]+)>\)",
            lambda match: (
                f'<img src="{escape(match.group(2), quote=True)}" '
                f'alt="{escape(match.group(1), quote=True)}">'
            ),
            markup,
        )
        self.feed(markup)
        self.close()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.casefold()
        attributes = {key.casefold(): value for key, value in attrs}
        if name == "br":
            self._add_break()
            return
        if name == "img":
            self._add_image(attributes.get("src"), attributes.get("alt") or "")
            return
        if name == "a":
            self._links.append(attributes.get("href"))
            return
        style = self._STYLE_TAGS.get(name)
        if style:
            self._styles[style] = self._styles.get(style, 0) + 1

    def handle_endtag(self, tag: str) -> None:
        name = tag.casefold()
        if name == "a":
            if self._links:
                self._links.pop()
            return
        style = self._STYLE_TAGS.get(name)
        if style and self._styles.get(style, 0) > 0:
            self._styles[style] -= 1

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)

    def handle_data(self, data: str) -> None:
        parts = data.split(PAGE_BREAK_MARKER)
        self._emit_text(parts[0])
        for part in parts[1:]:
            self._add_break(page=True)
            self._emit_text(part)

    def _emit_text(self, data: str) -> None:
        cursor = 0
        for match in self._MARKDOWN_LINK_PATTERN.finditer(data):
            self._add_run(data[cursor : match.start()])
            is_image, label, destination = match.groups()
            if is_image:
                self._add_image(destination, label)
            else:
                self._add_hyperlink(label, destination)
            cursor = match.end()
        self._add_run(data[cursor:])

    def _add_run(self, text: str) -> None:
        if not text:
            return
        href = self._links[-1] if self._links else None
        if href:
            self._add_hyperlink(text, href)
            return
        run = self.paragraph.add_run(text)
        self._apply_run_style(run)

    def _add_break(self, *, page: bool = False) -> None:
        run = self.paragraph.add_run()
        run.add_break(WD_BREAK.PAGE if page else WD_BREAK.LINE)

    def _add_hyperlink(self, text: str, destination: str) -> None:
        if not text:
            return
        relation_id = self.paragraph.part.relate_to(destination, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relation_id)
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        style = OxmlElement("w:rStyle")
        style.set(qn("w:val"), "Hyperlink")
        run_properties.append(style)
        run.append(run_properties)
        text_element = OxmlElement("w:t")
        text_element.text = text
        run.append(text_element)
        hyperlink.append(run)
        self.paragraph._p.append(hyperlink)

    def _add_image(self, destination: str | None, alt: str) -> None:
        if not destination:
            self._add_run(alt)
            return
        raw_destination = destination.split()[0].strip("<>")
        if self.resource_root is None or "://" in raw_destination:
            self._add_run(f"{alt or '图片'} ({destination})")
            return
        root = self.resource_root.resolve()
        raw_path = Path(raw_destination)
        source = (raw_path if raw_path.is_absolute() else root / raw_path).resolve()
        try:
            source.relative_to(root)
        except ValueError:
            self._add_run(f"{alt or '图片'} ({destination})")
            return
        if source.is_file():
            self.paragraph.add_run().add_picture(str(source))
            return
        self._add_run(f"{alt or '图片'} ({destination})")

    def _apply_run_style(self, run: Run) -> None:
        run.bold = self._styles.get("bold", 0) > 0
        run.italic = self._styles.get("italic", 0) > 0
        run.underline = self._styles.get("underline", 0) > 0
        run.font.strike = self._styles.get("strike", 0) > 0
        run.font.subscript = self._styles.get("subscript", 0) > 0
        run.font.superscript = self._styles.get("superscript", 0) > 0
        if self._styles.get("code", 0) > 0:
            run.font.name = "Consolas"


def restore_endnotes(output_path: Path, markdown: str) -> None:
    """Move explicitly labelled internal endnotes from Pandoc footnotes to OOXML endnotes."""

    labels = [match.group(1) for match in _FOOTNOTE_REFERENCE_PATTERN.finditer(markdown)]
    endnote_positions = {
        index for index, label in enumerate(labels) if label.startswith("doc-converter-endnote-")
    }
    if not endnote_positions:
        return

    with ZipFile(output_path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    document_root = etree.fromstring(members["word/document.xml"])
    reference_nodes = document_root.xpath("//w:footnoteReference", namespaces=_NS)
    moved_ids: list[str] = []
    for index in sorted(endnote_positions):
        if index >= len(reference_nodes):
            continue
        reference = reference_nodes[index]
        note_id = reference.get(f"{{{_W_NS}}}id")
        if not note_id:
            continue
        reference.tag = f"{{{_W_NS}}}endnoteReference"
        for style in reference.xpath("ancestor::w:r/w:rPr/w:rStyle", namespaces=_NS):
            value = style.get(f"{{{_W_NS}}}val")
            if value in _NOTE_STYLE_REPLACEMENTS:
                style.set(f"{{{_W_NS}}}val", _NOTE_STYLE_REPLACEMENTS[value])
        moved_ids.append(note_id)
    if not moved_ids:
        return

    footnotes_root = etree.fromstring(members["word/footnotes.xml"])
    endnotes_root = etree.Element(f"{{{_W_NS}}}endnotes", nsmap=footnotes_root.nsmap)
    for child in list(footnotes_root):
        note_id = child.get(f"{{{_W_NS}}}id")
        if note_id in {"-1", "0"}:
            copied = deepcopy(child)
            copied.tag = f"{{{_W_NS}}}endnote"
            endnotes_root.append(copied)
            continue
        if note_id in moved_ids:
            copied = deepcopy(child)
            copied.tag = f"{{{_W_NS}}}endnote"
            for element in copied.iter():
                if element.tag == f"{{{_W_NS}}}footnoteRef":
                    element.tag = f"{{{_W_NS}}}endnoteRef"
                if element.tag == f"{{{_W_NS}}}rStyle":
                    value = element.get(f"{{{_W_NS}}}val")
                    if value in _NOTE_STYLE_REPLACEMENTS:
                        element.set(f"{{{_W_NS}}}val", _NOTE_STYLE_REPLACEMENTS[value])
                if element.tag == f"{{{_W_NS}}}pStyle":
                    value = element.get(f"{{{_W_NS}}}val")
                    if value in _NOTE_STYLE_REPLACEMENTS:
                        element.set(f"{{{_W_NS}}}val", _NOTE_STYLE_REPLACEMENTS[value])
            endnotes_root.append(copied)
            footnotes_root.remove(child)
    members["word/document.xml"] = etree.tostring(document_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    members["word/footnotes.xml"] = etree.tostring(footnotes_root, xml_declaration=True, encoding="UTF-8")
    members["word/endnotes.xml"] = etree.tostring(endnotes_root, xml_declaration=True, encoding="UTF-8")

    content_types = etree.fromstring(members["[Content_Types].xml"])
    if not content_types.xpath(
        "/ct:Types/ct:Override[@PartName='/word/endnotes.xml']",
        namespaces=_NS,
    ):
        override = etree.SubElement(content_types, f"{{{_CT_NS}}}Override")
        override.set("PartName", "/word/endnotes.xml")
        override.set(
            "ContentType",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        )
    members["[Content_Types].xml"] = etree.tostring(content_types, xml_declaration=True, encoding="UTF-8")

    relationships = etree.fromstring(members["word/_rels/document.xml.rels"])
    if not relationships.xpath(
        "/rel:Relationships/rel:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes']",
        namespaces=_NS,
    ):
        existing_ids = [
            int(value[3:])
            for value in relationships.xpath("/rel:Relationships/rel:Relationship/@Id", namespaces=_NS)
            if value.startswith("rId") and value[3:].isdigit()
        ]
        relationship = etree.SubElement(relationships, f"{{{_REL_NS}}}Relationship")
        relationship.set("Id", f"rId{max(existing_ids, default=0) + 1}")
        relationship.set(
            "Type",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
        )
        relationship.set("Target", "endnotes.xml")
    members["word/_rels/document.xml.rels"] = etree.tostring(relationships, xml_declaration=True, encoding="UTF-8")

    settings = etree.fromstring(members["word/settings.xml"])
    if not settings.xpath("/w:settings/w:endnotePr", namespaces=_NS):
        endnote_pr = etree.SubElement(settings, f"{{{_W_NS}}}endnotePr")
        for note_id in ("-1", "0"):
            marker = etree.SubElement(endnote_pr, f"{{{_W_NS}}}endnote")
            marker.set(f"{{{_W_NS}}}id", note_id)
    members["word/settings.xml"] = etree.tostring(settings, xml_declaration=True, encoding="UTF-8")

    with tempfile.NamedTemporaryFile(
        prefix="doc-converter-endnotes-",
        suffix=".docx",
        dir=output_path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with ZipFile(temporary_path, "w", ZIP_DEFLATED) as target:
            for name, payload in members.items():
                target.writestr(name, payload)
        temporary_path.replace(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
