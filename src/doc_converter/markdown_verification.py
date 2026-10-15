"""Inspect written Markdown DOCX styles and effective outline independently of scoring."""

from __future__ import annotations

import re
from collections import defaultdict, deque
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement

from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.ooxml_patch import parse_xml_payload, read_zip_member
from doc_converter.docx_styles import ROLE_TO_STYLE_ID, compile_styles
from doc_converter.style_library import StylePack


def verify_rendered_roles(
    output: Path, decisions: list[dict[str, Any]], pack: StylePack
) -> dict[str, Any]:
    document = Document(str(output))
    resolver = EffectiveFormatResolver(
        parse_xml_payload(read_zip_member(output, "word/styles.xml"))
    )
    expected_document = Document()
    compile_styles(expected_document, pack)
    expected_resolver = EffectiveFormatResolver(expected_document.styles.element)
    paragraphs: dict[str, deque[Any]] = defaultdict(deque)
    for item in document.paragraphs:
        paragraphs[re.sub(r"\s+", " ", item.text).strip()].append(item)
    checks: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    expanded: list[dict[str, Any]] = []
    for decision in decisions:
        if decision["role"] in {"document-title", "document-subtitle"} and decision.get("features", {}).get("front_matter"):
            from doc_converter.markdown_scoring import parse_blocks
            expanded.extend({**decision, "features": {**decision["features"], "render_text": b.text}}
                            for b in parse_blocks(str(decision["before"]).splitlines()))
        else:
            expanded.append(decision)
    for d in expanded:
        role = d["role"]
        if role in {"organization", "department"}:
            role = "organization-department"
        text = d.get("features", {}).get("render_text")
        # Metadata, tables and parser-protected spans are verified separately by
        # corpus/integrity checks; they need not correspond to one body paragraph.
        if (
            not isinstance(text, str)
            or d.get("action") == "preserved"
            or role not in ROLE_TO_STYLE_ID
            or role == "body-table"
        ):
            skipped.append(
                {"line": d["line"], "role": role, "reason": "non-paragraph-or-protected"}
            )
            continue
        if role in {"author", "date", "organization-department"}:
            text = re.sub(r"^(?:作者单位|单位|机构|部门|作者|日期|时间)\s*[:：]\s*", "", text)
        queue = paragraphs.get(re.sub(r"\s+", " ", text).strip())
        paragraph = queue.popleft() if queue else None
        sid = ROLE_TO_STYLE_ID[role]
        effective = (
            resolver.resolve_paragraph(paragraph._p, paragraph_style_id=None) if paragraph else {}
        )
        actual_style = (
            paragraph.style.style_id
            if paragraph is not None and paragraph.style is not None
            else None
        )
        actual_outline = effective.get("outline_level")
        expected_outline = int(role[-1]) - 1 if role.startswith("heading-") else None
        exemplar = OxmlElement("w:p")
        expected = expected_resolver.resolve_paragraph(exemplar, paragraph_style_id=sid)
        format_keys = ("alignment", "paragraph_indent", "paragraph_spacing")
        style_matches = actual_style == sid
        outline_matches = (
            actual_outline == expected_outline if expected_outline is not None else
            actual_outline in {None, 9} if role in {"document-title", "document-subtitle"} else True
        )
        format_matches = all(effective.get(k) == expected.get(k) for k in format_keys)
        ordinary_runs = [
            run
            for run in effective.get("run_statistics", {}).get("runs", [])
            if run.get("text_length") and not run.get("style_chain")
        ]
        font_matches = all(
            run.get("font_size_pt") == expected.get("font_size_pt")
            and run.get("font_name") == expected.get("font_name")
            for run in ordinary_runs
        )
        format_matches = format_matches and font_matches
        # Native list body paragraphs retain numbering and their list style.
        if (
            paragraph is not None
            and role == "body-paragraph"
            and paragraph._p.xpath("./w:pPr/w:numPr")
        ):
            skipped.append({"line": d["line"], "role": role, "reason": "native-list-body"})
            continue
        checks.append(
            {
                "line": d["line"],
                "text": text,
                "role": role,
                "expected_style": sid,
                "actual_style": actual_style,
                "expected_outline": expected_outline,
                "actual_outline": actual_outline,
                "style_match": style_matches,
                "outline_match": outline_matches,
                "effective_format_match": format_matches,
                "font_match": font_matches,
                "passed": style_matches and outline_matches and format_matches,
            }
        )
    failures = sum(not c["passed"] for c in checks)
    return {
        "schema": "markdown-output-verification/v1",
        "passed": failures == 0,
        "style_pack_id": pack.id,
        "checked_items": len(checks),
        "failed_checks": failures,
        "checks": checks,
        "skipped": skipped,
    }
