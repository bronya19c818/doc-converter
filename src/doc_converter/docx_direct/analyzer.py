"""Read-only structural analysis for DOCX Direct normalization."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Iterator
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from doc_converter.docx_direct.effective_format import EffectiveFormatResolver
from doc_converter.docx_direct.model import DocumentAnalysis, ElementAnalysis

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
P_TAG = f"{{{W_NS}}}p"
TBL_TAG = f"{{{W_NS}}}tbl"
TR_TAG = f"{{{W_NS}}}tr"
TC_TAG = f"{{{W_NS}}}tc"
TEXT_TAGS = {qn("w:t"), qn("w:delText")}
DIRECT_RUN_PROPERTIES = {
    "rFonts": qn("w:rFonts"),
    "size": qn("w:sz"),
    "bold": qn("w:b"),
    "italic": qn("w:i"),
    "color": qn("w:color"),
    "vertAlign": qn("w:vertAlign"),
}


def analyze_docx(source_path: Path) -> DocumentAnalysis:
    """Read a DOCX package and return structural evidence without writing it."""

    source_path = Path(source_path)
    source_sha256 = _sha256(source_path)
    try:
        with ZipFile(source_path) as package:
            entries = package.namelist()
            if "word/document.xml" not in entries:
                raise ValueError("DOCX 压缩包缺少 word/document.xml。")
            document_xml = package.read("word/document.xml")
            styles_xml = package.read("word/styles.xml") if "word/styles.xml" in entries else None
    except BadZipFile as error:
        raise ValueError("输入文件不是有效的 DOCX ZIP。") from error
    except OSError as error:
        raise ValueError(f"无法读取 DOCX 文件：{error}") from error

    try:
        root = parse_xml(document_xml)
    except Exception as error:  # lxml's XMLSyntaxError is not part of the public model.
        raise ValueError("DOCX 的 word/document.xml 不是有效 XML。") from error

    try:
        document = Document(str(source_path))
    except Exception as error:
        raise ValueError("无法读取 DOCX 的文档对象模型。") from error

    style_map = _style_map(document)
    resolver = (
        EffectiveFormatResolver.from_styles_xml(styles_xml)
        if styles_xml is not None
        else EffectiveFormatResolver.empty()
    )
    elements: list[ElementAnalysis] = []
    body = root.find(qn("w:body"))
    if body is None:
        raise ValueError("DOCX 的 word/document.xml 缺少 w:body。")

    for order, (paragraph, xml_path, table_context) in enumerate(_iter_paragraphs(body)):
        style_id = _paragraph_style_id(paragraph)
        style_name = style_map.get(style_id, style_id) if style_id is not None else None
        elements.append(
            _analyze_paragraph(
                paragraph,
                element_id=f"p-{order + 1:04d}",
                document_order=order,
                xml_path=xml_path,
                style_id=style_id,
                style_name=style_name,
                table_context=table_context,
                resolver=resolver,
            )
        )

    return DocumentAnalysis(
        source_sha256=source_sha256,
        elements=elements,
        package_entries=entries,
        main_document_stats={
            "paragraphs": len(document.paragraphs),
            "analyzed_paragraphs": len(elements),
            "tables": len(document.tables),
            "sections": len(document.sections),
            "inline_shapes": len(document.inline_shapes),
        },
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _style_map(document: Any) -> dict[str, str]:
    result: dict[str, str] = {}
    for style in document.styles:
        try:
            result[style.style_id] = style.name
        except (AttributeError, KeyError):
            continue
    return result


def _iter_paragraphs(
    container: Any,
    *,
    path: tuple[int, ...] = (),
    table_context: dict[str, Any] | None = None,
) -> Iterator[tuple[Any, tuple[int, ...], dict[str, Any] | None]]:
    for child_index, child in enumerate(container):
        child_path = path + (child_index,)
        if child.tag == P_TAG:
            yield child, child_path, table_context
        elif child.tag == TBL_TAG:
            yield from _iter_table(child, child_path, table_context)


def _iter_table(
    table: Any,
    table_path: tuple[int, ...],
    parent_context: dict[str, Any] | None,
) -> Iterator[tuple[Any, tuple[int, ...], dict[str, Any] | None]]:
    parent_depth = int((parent_context or {}).get("table_depth", 0))
    table_depth = parent_depth + 1
    rows = [child for child in table if child.tag == TR_TAG]
    for row_index, row in enumerate(rows):
        cells = [child for child in row if child.tag == TC_TAG]
        for cell_index, cell in enumerate(cells):
            context = {
                "table_depth": table_depth,
                "row": row_index,
                "column": cell_index,
                "nested": table_depth > 1,
            }
            cell_path = table_path + (row_index, cell_index)
            yield from _iter_paragraphs(cell, path=cell_path, table_context=context)


def _paragraph_style_id(paragraph: Any) -> str | None:
    ppr = paragraph.find(qn("w:pPr"))
    if ppr is None:
        return None
    style = ppr.find(qn("w:pStyle"))
    return style.get(qn("w:val")) if style is not None else None


def _analyze_paragraph(
    paragraph: Any,
    *,
    element_id: str,
    document_order: int,
    xml_path: tuple[int, ...],
    style_id: str | None,
    style_name: str | None,
    table_context: dict[str, Any] | None,
    resolver: EffectiveFormatResolver,
) -> ElementAnalysis:
    run_properties = [_run_properties(run) for run in paragraph.iter(qn("w:r"))]
    ppr = paragraph.find(qn("w:pPr"))
    paragraph_properties = _paragraph_properties(ppr)
    numbering_properties = paragraph_properties.get("numbering", {})
    resolved_format = resolver.resolve_paragraph(
        paragraph,
        paragraph_style_id=style_id,
    )
    effective_format = {
        key: value
        for key, value in resolved_format.items()
        if key not in {"style_chain", "based_on_heading"}
    }
    signals = {
        "style_chain": resolved_format.get("style_chain", []),
        "based_on_heading": resolved_format.get("based_on_heading", False),
    }
    return ElementAnalysis(
        element_id=element_id,
        element_type="paragraph",
        text=_paragraph_text(paragraph),
        document_order=document_order,
        xml_path=xml_path,
        region="body",
        original_style_id=style_id,
        original_style_name=style_name,
        paragraph_properties=paragraph_properties,
        run_properties=run_properties,
        numbering_properties=numbering_properties,
        table_context=table_context,
        effective_format=effective_format,
        signals=signals,
        has_hyperlink=paragraph.find(".//" + qn("w:hyperlink")) is not None,
        has_tab=paragraph.find(".//" + qn("w:tab")) is not None,
        has_manual_break=paragraph.find(".//" + qn("w:br")) is not None,
        has_drawing=paragraph.find(".//" + qn("w:drawing")) is not None,
        has_field=(paragraph.find(".//" + qn("w:fldSimple")) is not None
        or paragraph.find(".//" + qn("w:instrText")) is not None),
        has_embedded_object=any(
            paragraph.find(".//" + tag) is not None
            for tag in (qn("w:object"), qn("w:pict"),
                        "{urn:schemas-microsoft-com:office:office}OLEObject")
        ),
        has_math=any(paragraph.find(".//" + qn(tag)) is not None
                     for tag in ("m:oMath", "m:oMathPara")),
    )


def _paragraph_text(paragraph: Any) -> str:
    parts: list[str] = []
    for node in paragraph.iter():
        if node.tag in TEXT_TAGS and node.text:
            parts.append(node.text)
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag == qn("w:br"):
            parts.append("\n")
    return "".join(parts)


def _paragraph_properties(ppr: Any) -> dict[str, Any]:
    if ppr is None:
        return {}
    values: dict[str, Any] = {}
    for tag, key in (
        ("w:jc", "alignment"),
        ("w:ind", "indent"),
        ("w:spacing", "spacing"),
        ("w:outlineLvl", "outline_level"),
    ):
        node = ppr.find(qn(tag))
        if node is not None:
            values[key] = {str(attribute): value for attribute, value in node.attrib.items()}
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        values["numbering"] = {
            "num_id": _child_val(numpr, "w:numId"),
            "ilvl": _child_val(numpr, "w:ilvl"),
        }
    return values


def _child_val(parent: Any, tag: str) -> str | None:
    node = parent.find(qn(tag))
    return node.get(qn("w:val")) if node is not None else None


def _run_properties(run: Any) -> dict[str, Any]:
    rpr = run.find(qn("w:rPr"))
    values: dict[str, Any] = {"text": _paragraph_text(run)}
    if rpr is None:
        return values
    for key, tag in DIRECT_RUN_PROPERTIES.items():
        node = rpr.find(tag)
        if node is not None:
            values[key] = {str(attribute): value for attribute, value in node.attrib.items()}
    return values
